"""Build Spring 2015 and Fall 2014 from their official exam PDFs.

The CBT workbooks supply cells and scored-part grids. Printed question pages
control wording, complete table headings, units, figures, and corrections.
"""
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from build_spring2019 import displayed

ROOT=Path(__file__).resolve().parents[1]
CONFIG={
 'spring-2015':{'workbook':'CBT_Exam_5_S.15_v01.xlsx','pdf':'admissions_studytools_exam5_sp15-5.pdf','var':'SPRING_2015_QUESTIONS','count':25,'part_count':49,'points':57.25,
  'report_start':[31,34,37,39,41,43,45,48,53,55,58,61,63,69,72,77,81,84,86,90,94,96,99,102,104]},
 'fall-2014':{'workbook':'CBT_Exam_5_F.14_v01.xlsx','pdf':'admissions_studytools_exam5_f14-5.pdf','var':'FALL_2014_QUESTIONS','count':24,'part_count':55,'points':58.25,
  'report_start':[30,35,37,38,41,42,44,46,51,53,54,57,60,62,65,66,67,70,72,74,76,79,81,83]},
}
CHAPTERS={
 'spring-2015':[
  ['ratemaking-9'],['ratemaking-4'],['ratemaking-4'],['ratemaking-5','ratemaking-6'],
  ['ratemaking-5'],['ratemaking-6'],['ratemaking-16'],['ratemaking-6'],['ratemaking-7'],
  ['ratemaking-8'],['ratemaking-8','ratemaking-12','reserving-14'],['ratemaking-9'],
  ['ratemaking-9'],['ratemaking-11'],['reserving-6'],['reserving-6'],
  ['reserving-10'],['reserving-8'],['reserving-9','reserving-15'],['reserving-11','reserving-15'],
  ['reserving-11','reserving-13'],['reserving-13'],['reserving-13'],['reserving-14'],['reserving-17']],
 'fall-2014':[
  ['ratemaking-4'],['ratemaking-5'],['ratemaking-5'],['ratemaking-6'],['ratemaking-6'],
  ['ratemaking-7'],['ratemaking-8'],['ratemaking-8','ratemaking-12'],['ratemaking-9'],
  ['ratemaking-10'],['ratemaking-9','ratemaking-12'],['ratemaking-13'],['reserving-1'],
  ['reserving-6'],['reserving-8'],['reserving-7'],['reserving-11'],
  ['reserving-9'],['reserving-15'],['reserving-13'],['reserving-14'],['reserving-6'],
  ['reserving-9','reserving-17'],['reserving-15']],
}
# first covered row, last covered row, data rows, workbook columns,
# complete title, complete headings as printed in the question PDF.
T={
 'spring-2015':{
  1:[(6,11,range(7,12),[2,3],'Primary Classification',['Primary Classification','Factor']),
     (6,12,range(7,13),[5,6],'Secondary Classification',['Secondary Classification','Factor'])],
  3:[(6,10,range(7,11),[2,3],'',['Time Period','Written Car Years'])],
  4:[(14,17,range(15,18),[3,4,5],'Rate Change History',['Effective Date','Overall Change','Type of Change']),
     (22,34,range(23,35),[3,4,6,7],'Annual Premium Exponential Trend Fits',['Calendar Year Ending','Average Earned Premium at Current Rate and Law Level','Number of Points','Annual Exponential Trend Fit'])],
  5:[(17,20,range(17,21),[3,4],'Policy Distribution',['Quarter','Share of Policies'])],
  7:[(15,19,range(16,20),[2,3],'Ultimate Loss Costs for Report Year 2014',['Report Year Lag','Loss Cost'])],
  8:[(7,13,range(8,14),[2,3,4,5],'Historical Large Loss Information',['Accident Year','Reported Loss Including Severity Trend ($000)','Claims Greater Than $500,000','Excess Ratio at $500,000']),
     (27,31,range(28,32),[2,3],'Individual Claims Greater Than $500,000',['Claim','Reported Loss Including Severity Trend ($)'])],
  9:[(6,10,range(7,11),[2,3],'',['Expense Type','2014 Expense Amount ($)'])],
  10:[(6,8,range(7,9),[2,3,4],'',['Territory','Premium ($)','Ultimate Loss Ratio'])],
  11:[(22,24,range(23,25),[2,3,4,5,6],'',['Calendar/Accident Year','Earned Exposures','On-Level Earned Premium ($000)','Gross Reported Loss Evaluated December 31, 2014 ($000)','Ratio of Subrogation and Salvage Received to Paid Loss Evaluated December 31, 2014']),
      (27,32,range(29,33),[2,3,4,5,6],'Gross Reported Loss Development Factors',['Accident Year','12–24','24–36','36–48','48–Ultimate']),
      (34,39,range(36,40),[2,3,4,5,6],'Ratio of Subrogation and Salvage Received to Paid Loss Development Factors',['Accident Year','12–24','24–36','36–48','48–Ultimate'])],
  13:[(6,11,range(9,12),[2,3,4],'Earned Exposures by Vehicle Symbol and Driver Age',['Vehicle Symbol','Under 25','25 and Older']),
      (13,15,range(14,16),[2,3],'',['Driver Age','Incurred Loss and ALAE ($)']),
      (17,22,range(18,23),[2,3],'Current Rating Factors',['Rating Class','Factor'])],
  14:[(6,9,range(7,10),[2,3,4],'',['Size of Loss','Reported Claim Counts','Reported Ground-Up Losses ($)'])],
  15:[(6,18,list(range(7,10))+list(range(11,15))+list(range(16,19)),[2,3,4,5,6,7],'Claim Transactions Through December 31, 2014',['Claim ID','Accident Date','Report Date','Transaction Date','Incremental Payment ($)','Ending Case Reserve ($)'])],
  16:[(6,11,range(8,12),[2,3,4],'Reported Claims ($000) as of (months)',['Accident Year','12','24']),
      (6,11,range(8,12),[6,7,8],'Paid Claims ($000) as of (months)',['Accident Year','12','24']),
      (13,18,range(15,19),[2,3,4],'Reported Claim Counts as of (months)',['Accident Year','12','24']),
      (13,18,range(15,19),[6,7,8],'Closed Claim Counts as of (months)',['Accident Year','12','24'])],
  17:[(6,10,range(7,11),[2,3,4,5,6],'',['Calendar/Accident Year','Earned Car Years','Earned Premium ($000)','Reported Claims ($000)','Selected Cumulative Development Factor'])],
  18:[(6,10,range(8,11),[2,3,4,5],'Reported Claims ($000) as of (months)',['Accident Year','12','24','36']),
      (12,15,range(13,16),[2,3],'',['Calendar Year','Average Occupied Beds and Outpatient Visits'])],
  19:[(6,9,range(7,10),[2,3,4,5,6,7,8],'Unpaid Claims Analysis ($ millions)',['Accident Year','Paid Claims','Reported Claims','Paid Development Ultimate','Reported Development Ultimate','Paid Bornhuetter–Ferguson Ultimate','Selected Ultimate Claims'])],
  20:[(6,15,range(7,16),[2,5],'Calendar/Accident Year 2014 Information',['Item','Value']),
      (17,19,range(18,20),[2,5],'Additional Information',['Item','Value']),
      (21,27,range(23,28),[2,3,4,5,6,7],'Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48','60']),
      (29,35,range(31,36),[2,3,4,5,6,7],'Paid Claims ($000) as of (months)',['Accident Year','12','24','36','48','60'])],
  21:[(6,10,range(8,11),[2,3,4,5],'Cumulative Closed Claim Counts as of (months)',['Accident Year','12','24','36']),
      (12,16,range(14,17),[2,3,4,5],'Cumulative Paid Claims ($000) as of (months)',['Accident Year','12','24','36']),
      (18,21,range(19,22),[2,3],'',['Accident Year','Ultimate Claim Counts'])],
  22:[(7,11,range(9,12),[2,3,4,5],'Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36']),
      (7,11,range(9,12),[7,8,9,10],'Cumulative Reported Claim Counts as of (months)',['Accident Year','12','24','36']),
      (13,17,range(15,18),[2,3,4,5],'Cumulative Paid Claims ($000) as of (months)',['Accident Year','12','24','36']),
      (13,17,range(15,18),[7,8,9,10],'Cumulative Closed Claim Counts as of (months)',['Accident Year','12','24','36']),
      (19,23,range(21,24),[2,3,4,5],'Case Outstanding ($000) as of (months)',['Accident Year','12','24','36']),
      (19,23,range(21,24),[7,8,9,10],'Open Claim Counts as of (months)',['Accident Year','12','24','36'])],
  24:[(6,10,range(8,11),[2,3,4,5],'Paid Claims Gross of Salvage and Subrogation ($000) as of (months)',['Accident Year','12','24','36']),
      (12,16,range(14,17),[2,3,4,5],'Received Salvage and Subrogation ($000) as of (months)',['Accident Year','12','24','36'])],
  25:[(7,12,range(9,13),[2,3,4,5,6,7,8],'Calendar-Year and Accident-Year ULAE Data ($)',['Year','Calendar-Year Paid ULAE','Calendar-Year Paid Claims','Calendar-Year Reported Claims','Ultimate on Claims Reported in Calendar Year','Accident-Year IBNR as of December 31, 2014','Accident-Year Reported Claims'])],
 },
 'fall-2014':{
  2:[(6,10,range(7,11),[2,3,4,5],'',['Policy','Effective Date','Expiration Date','Initial Policy Premium ($)'])],
  4:[(6,12,range(7,13),[2,3],'',['Ratio of Wage to State Average Weekly Wage','Percentage of Workers'])],
  5:[(6,9,range(7,10),[2,3],'',['Period (months)','Reported Loss and ALAE Age-to-Age Factor']),
     (11,14,range(12,15),[2,3,4,5],'',['Calendar/Accident Year','Earned Exposures (000)','Amount of Insurance Years ($000)','Reported Non-Catastrophe Loss and ALAE ($000)'])],
  6:[(10,14,range(11,15),[2,3,4],'',['Expense Category','Selected Expense Ratio','Percent Fixed'])],
  7:[(6,9,range(7,10),[2,3,4],'',['Calendar/Accident Year','Written Policies','Ultimate Loss and LAE ($000)'])],
  8:[(28,30,range(29,31),[2,3],'',['Calendar Year Ending','Earned Premium ($000)']),
     (32,34,range(33,35),[2,3],'',['Accident Year as of December 31, 2013','Reported Losses ($000)']),
     (36,42,range(38,43),[2,3,4,5,6,7],'Reported Loss Age-to-Age Development Factors',['Accident Year','12–24','24–36','36–48','48–60','60–72'])],
  9:[(7,10,range(9,11),[2,3,4,5],'Insured Risks',['Rating Group','True Expected Cost ($)','Company A','Company B'])],
  11:[(6,9,range(7,10),[2,3,4,5,6],'',['Territory','Earned Exposures','Earned Premium ($000)','Ultimate Losses Excluding Catastrophes ($000)','Current Relativity'])],
  12:[(6,12,range(8,13),[2,3,4,5,6,7,8,9],'Asset Share Model by Policy Year',['Policy Year','Premium ($)','Present Value of Losses ($)','Variable Expenses — New ($)','Variable Expenses — Renewal ($)','Fixed Expenses — New ($)','Fixed Expenses — Renewal ($)','Income ($)']),
      (14,20,range(15,21),[2,3,4,5,6,7,8],'Persistency and Present Values',['Policy Year','Persistency','Cumulative Persistency','Profit ($)','Discount Factor','Present Value of Profits ($)','Present Value of Premiums ($)'])],
  14:[(6,17,range(8,18),[2,3,4,5,6,7,8,9,10],'Claim Transactions by Calendar Year',['Claim ID','Accident Date','Report Date','2011 Paid ($)','2011 Ending Case Outstanding ($)','2012 Paid ($)','2012 Ending Case Outstanding ($)','2013 Paid ($)','2013 Ending Case Outstanding ($)'])],
  15:[(6,9,range(7,10),[2,3,4,5],'',['Accident Year','Reported Claims ($000)','Reported Development Factor to Ultimate','On-Level Earned Premium ($000)'])],
  16:[(6,13,range(8,14),[2,3,4,5,6,7],'Reported Claim Counts as of (months)',['Accident Half-Year','6','12','18','24','30'])],
  17:[(14,19,range(16,20),[2,3,4,5,6],'Cumulative Paid Claims ($) as of (months)',['Accident Year','12','24','36','48']),
      (21,26,range(23,27),[2,3,4,5,6],'Cumulative Closed Claim Counts as of (months)',['Accident Year','12','24','36','48']),
      (28,33,range(30,34),[2,3,4,5,6],'Incremental Closed Claim Counts as of (months)',['Accident Year','12','24','36','48']),
      (35,39,range(36,40),[2,3],'',['Accident Year','Ultimate Claim Counts'])],
  18:[(6,10,range(7,11),[2,3,4,5,6],'',['Accident Year','On-Level Earned Premium ($)','Reported Claims ($)','Reported Development Factor to Ultimate','Expected Claims Ratio'])],
  20:[(6,11,range(8,12),[2,3,4,5,6],'Cumulative Paid Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (6,11,range(8,12),[8,9,10,11,12],'Case Outstanding ($000) as of (months)',['Accident Year','12','24','36','48']),
      (13,18,range(15,19),[2,3,4,5,6],'Closed Claim Counts (000) as of (months)',['Accident Year','12','24','36','48']),
      (13,18,range(15,19),[8,9,10,11,12],'Open Claim Counts (000) as of (months)',['Accident Year','12','24','36','48']),
      (20,24,range(21,25),[2,3],'',['Accident Year','Projected Ultimate Claim Counts (000)']),
      (28,37,range(29,38),[2,3],'Interpolated Cumulative Paid Claims — Accident Year 2010 ($000)',['Closed Claim Counts (000)','Paid Claims ($000)']),
      (28,36,range(29,37),[5,6],'Interpolated Cumulative Paid Claims — Accident Year 2011 ($000)',['Closed Claim Counts (000)','Paid Claims ($000)']),
      (28,34,range(29,35),[8,9],'Interpolated Cumulative Paid Claims — Accident Year 2012 ($000)',['Closed Claim Counts (000)','Paid Claims ($000)'])],
  21:[(6,11,range(8,12),[2,3,4,5,6],'Gross Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (13,18,range(15,19),[2,3,4,5,6],'Net Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48'])],
  23:[(6,11,range(8,12),[2,3,4,5,6],'Cumulative Paid Claims as of (months)',['Accident Year','12','24','36','48']),
      (13,17,range(14,18),[2,3],'',['Calendar Year','Paid ULAE'])],
 }
}

# Workbook transcription errors corrected from the printed question pages.
FIXES={
 ('spring-2015',1,'B14'):'• The policy has one 30-year-old driver who uses the vehicle to commute 10 miles to work one day per week.',
 ('spring-2015',1,'B20'):'• Only auto usage and safe driver insurance plan (SDIP) are used to determine primary and secondary classification.',
 ('spring-2015',1,'B28'):'• Uses the vehicle 3 days per week to commute 10 miles to school.',
 ('spring-2015',5,'B6'):'• On July 1, 2014, a rate change of +10% went into effect.',
 ('spring-2015',14,'B17'):'Calculate the combined loss elimination ratio for a policy with a $1,000 deductible and a policy limit of $3,000.',
 ('spring-2015',25,'E10'):'128,672',
 ('fall-2014',1,'B19'):'Briefly describe the impact the exposure base change could have on severity.',
 ('fall-2014',12,'B23'):'• Surplus equals GAAP equity.',
 ('fall-2014',16,'B4'):'The following information is available for an insurer:',
}
for row,year in zip((9,10,11),(2012,2013,2014)):
    for column in ('B','G'):
        for offset in (0,6,12):
            FIXES['spring-2015',22,f'{column}{row+offset}']=str(year)

def val(exam,n,cell):
    if (exam,n,cell.coordinate) in FIXES:return FIXES[exam,n,cell.coordinate]
    if cell.value is None:return ''
    if isinstance(cell.value,datetime):return f'{cell.value:%B} {cell.value.day}, {cell.value.year}'
    if isinstance(cell.value,int) and 2000<=cell.value<=2030:return str(cell.value)
    result=displayed(cell)
    if isinstance(cell.value,(int,float)) and cell.value>=1000 and cell.value==int(cell.value) and '%' not in cell.number_format:
        result=f'{cell.value:,.0f}'
    return re.sub(r'\s+',' ',result).strip()

def row_text(exam,n,sh,row,excluded=()):
    return ' '.join(val(exam,n,c) for c in sh[row][1:17]
        if c.column not in excluded and c.value is not None and c.data_type!='f').strip()

def blocks(exam,n,sh,first_part):
    specs=T[exam].get(n,[])
    covered={}
    for first,last,rows,cols,title,headers in specs:
        for row in range(first,last+1):covered.setdefault(row,set()).update(cols)
    result=[]
    for row in range(3,sh.max_row+1):
        for first,last,rows,cols,title,headers in specs:
            if row==first:
                result.append({'type':'table','title':title,'headers':headers,
                    'rows':[[val(exam,n,sh.cell(r,c)) for c in cols] for r in rows]})
        if row<first_part:
            line=row_text(exam,n,sh,row,covered.get(row,set()))
            if line:result.append({'type':'line','text':line})
    return result,covered

def clean(s):
    s=re.sub(r'(?m)^\s*EXAM 5 (?:SPRING 2015|FALL 2014) SAMPLE ANSWERS AND EXAMINER.S REPORT\s*$','',s)
    s=s.replace('','•').replace('','→').replace('','→').replace('','←')
    return re.sub(r'\n{3,}','\n\n',s).strip()

def report_parts(s):
    markers=list(re.finditer(r'(?im)^\s*Part\s+([a-z])\s*:?(?:\s*\.?\d+(?:\.\d+)?\s+point(?:\(s\)|s)?)?\s*$',s))
    if not markers:return {'a':s}
    return {m.group(1).lower():s[m.end():markers[i+1].start() if i+1<len(markers) else None] for i,m in enumerate(markers)}

def report(exam,n,pdf):
    starts=CONFIG[exam]['report_start']
    first=starts[n-1]
    last=starts[n] if n<CONFIG[exam]['count'] else len(pdf.pages)+1
    raw='\n'.join(pdf.pages[page-1].extract_text() for page in range(first,last))
    sample,insight=re.split(r'(?im)^\s*EXAMINER.S REPORT\s*:?\s*$',raw,maxsplit=1)
    sample=re.split(r'(?im)^\s*SAMPLE(?:/ACCEPTED)? ANSWERS?\s*:?\s*$',sample,maxsplit=1)[-1]
    answers=report_parts(sample)
    notes=report_parts(insight)
    for key,s in list(answers.items()):
        samples=re.split(r'(?im)^\s*(?:Sample|Accepted Answer|Sample Answer|Sample Response)\s*\d+(?:\s*\([^)]*\))?\s*:?\s*$',clean(s))
        candidates=samples[1:] if len(samples)>1 else samples
        answers[key]=next((clean(item) for item in candidates if clean(item)),'')
    for key,s in list(notes.items()):
        s=clean(s)
        m=re.search(r'(?i)(?:Candidates?|The candidate) (?:were|was|are|is) expected to[^.]*\.',s)
        if m:s=m.group(0)
        else:
            m=re.search(r'(?i)(?:To (?:receive|score) full credit|Candidates needed to)[^.]*\.',s)
            if m:s=m.group(0)
            else:s=re.split(r'(?<=[.!?])\s+',re.sub(r'\s+',' ',s))[0]
        s=re.sub(r'\s+',' ',s).strip()
        s=re.sub(r'(?i)^(?:Candidates?|The candidate) (?:were|was|are|is) expected to\s+','',s)
        if len(s)>270:s=s[:270].rsplit(' ',1)[0].rstrip(' ,;:')+'.'
        if s and not s.endswith(('.','!','?')):s+='.'
        notes[key]=s[0].upper()+s[1:] if s else ''
    return answers,notes

SOLUTION_FIXES={
 ('spring-2015',19,'a'):'The 2012 paid development factor is 14/7 = 2. The paid Bornhuetter–Ferguson result gives 12 = 7 + E(1 − 1/2), so the common a priori expected ultimate E is $10 million. The reported factors are 10.9/7.8 = 1.3974 for 2012, 11/5 = 2.2 for 2013, and 5.7/1.9 = 3 for 2014. Reported Bornhuetter–Ferguson ultimates are, respectively, 7.8 + 10(1 − 1/1.3974) ≈ $10.64 million, 5 + 10(1 − 1/2.2) ≈ $10.45 million, and 1.9 + 10(1 − 1/3) ≈ $8.57 million.',
 ('spring-2015',19,'b'):'For 2012, the selected $13.8 million emphasizes the paid estimates; that may be reasonable if case reserves weakened or settlements accelerated, but the reported estimates are substantially lower. For 2013, the selected $10.2 million is close to the reported Bornhuetter–Ferguson result and can be reasonable given its intermediate maturity, provided the reporting and payment patterns remain credible. For immature 2014, the selected $7.1 million is below both Bornhuetter–Ferguson results; it may be low unless there is evidence that expected claims or the development pattern has changed. Assess each year using its maturity and the suitability of the underlying methods.',
 ('spring-2015',20,'a'):'One accepted set of estimates is: frequency–severity = 36 million × 0.0042 × ($160 × 1.05²) ≈ $26.672 million; reported development = $20.3 million × 1.72 = $34.916 million; paid development = $8.7 million × 2.96 = $25.752 million; reported Bornhuetter–Ferguson = $20.3 million + $56.16 million × 44.5% × (1 − 1/1.72) ≈ $30.761 million. The report accepted other reasonable interpretations of the misstated exposure base for the Bornhuetter–Ferguson expected claims calculation.',
 ('spring-2015',20,'b'):'The reported triangle shows an unusually high 2014 value without a similar increase in paid claims, suggesting stronger case reserves; reported development may overstate ultimate. Paid development is more stable but its 12-to-ultimate factor is highly leveraged. Frequency–severity is less affected by case reserves, though its 2012 severity selection may be stale. Reported Bornhuetter–Ferguson tempers the high reported amount with expected claims but still depends on the credibility of that expectation.',
 ('spring-2015',23,'a'):'Paid development uses historical payment factors that may no longer apply after smaller claims begin settling faster. Reported development can also shift if the changed mix of open claims affects case outstanding or reported claims. The direction of either distortion depends on how settlement and case reserves change; explain the mechanism rather than assuming an automatic overstatement or understatement.',
 ('spring-2015',23,'b'):'Separate small and large claims and analyze their development patterns by size. Adjust historical paid development to reflect the new settlement mix within each group, then combine the estimates. A blanket Berquist–Sherman adjustment can mislead when paid dollars are no longer proportional to closed claim counts.',
}
INSIGHT_FIXES={
 ('spring-2015',6,'a'):'Judge calendar-year aggregation for short-tailed auto physical damage using both its practicality and the alignment of premium and losses.',
 ('spring-2015',6,'b'):'Explain how policy-year aggregation aligns homeowners premium and losses while delaying the maturity of the data.',
 ('spring-2015',6,'c'):'Evaluate report-year aggregation for long-tailed medical professional liability, including the claims-made coverage trigger.',
 ('spring-2015',10,'a'):'Calculate each territory indication and apply the 20% cap to the constrained territory.',
 ('spring-2015',10,'b'):'Explain an operational or expense action that restores balance to the fundamental insurance equation when rate changes are constrained.',
 ('spring-2015',11,'a'):'Develop gross losses and salvage recoveries, trend losses and premium to the prospective period, apply LAE and credibility, and justify factor selections.',
 ('spring-2015',7,'c'):'Explain how the shorter reporting window changes occurrence coverage while leaving mature claims-made losses from prior accident years largely unaffected.',
 ('spring-2015',14,'b'):'Calculate the loss elimination ratio for the combined deductible and policy limit without double-counting the overlapping loss layer.',
 ('spring-2015',14,'c'):'Explain two ways historical deductibles or limits censor the loss data and propose an adjustment for each.',
 ('spring-2015',17,'a'):'On-level earned premium, trend claims to a common cost level, estimate the Cape Cod expected loss ratio, and apply it to accident year 2014.',
 ('spring-2015',19,'a'):'Infer the a priori expected claims from paid development and paid Bornhuetter–Ferguson, then apply the reported Bornhuetter–Ferguson formula by accident year.',
 ('spring-2015',20,'a'):'Calculate all four ultimate estimates and state the assumption used for the Bornhuetter–Ferguson exposure base.',
 ('spring-2015',20,'b'):'Assess each method against the unusual reported emergence, stable paid history, age-to-ultimate leverage, and expected severity.',
 ('spring-2015',23,'a'):'Explain how the new settlement priority changes paid development and may affect reported development through case reserves.',
 ('spring-2015',23,'b'):'Adjust development by claim size so the changing mix of settled claims does not distort a blanket Berquist–Sherman adjustment.',
 ('fall-2014',2,'a'):'Earn calendar-year 2013 premium by policy, including audit adjustments only when known by the evaluation date.',
 ('fall-2014',2,'b'):'Include the policy premiums and audits written by December 31, 2013, using the proper policy terms.',
 ('fall-2014',2,'c'):'Earn policy-year 2013 premium through December 31, 2013 using each policy’s elapsed coverage period.',
 ('fall-2014',2,'d'):'Include the 8% final audits for applicable policy-year 2013 policies by the December 31, 2014 evaluation.',
 ('fall-2014',3,'a'):'Assess both premium on-level methods for a detailed classification plan, including accuracy, data needs, and practical limits.',
 ('fall-2014',10,'a'):'Evaluate the consistent frequency pattern across hazard classes alongside the sparse exposure in the end classes.',
 ('fall-2014',9,'a'):'Compare the short-term profits at both proposed low-risk prices, accounting for how many insureds switch.',
 ('fall-2014',10,'b'):'Describe two data mining methods and explain what each can contribute to the GLM analysis.',
 ('fall-2014',13,'a'):'Explain how a redundant unpaid estimate can mislead management, investors, and regulators.',
 ('fall-2014',13,'b'):'Explain how an inadequate unpaid estimate can mislead management, investors, and regulators.',
 ('fall-2014',15,'a'):'Adjust historical losses for the law change and trend, derive an expected claims ratio, and subtract reported claims from expected ultimate claims.',
 ('fall-2014',19,'a'):'Explain why stronger case outstanding can overstate a reported-development estimate and propose a method or adjustment that removes the distortion.',
 ('fall-2014',22,'a'):'Interpolate expected emergence between the 12- and 24-month cumulative development factors through July 31, 2014.',
 ('fall-2014',23,'a'):'Select a paid 12-to-24 factor, combine it with the selected 24-to-ultimate factor, and calculate paid Bornhuetter–Ferguson IBNR.',
 ('fall-2014',23,'b'):'Select a ULAE ratio and apply the classical formula to half of case outstanding plus IBNR.',
 ('fall-2014',23,'c'):'State the classical technique’s timing and stable-growth assumptions, and assess them for this rapidly growing insurer.',
 ('fall-2014',23,'d'):'Describe a classical-technique refinement, such as Kittel’s denominator, that accounts for rapid company growth.',
 ('fall-2014',24,'a'):'Choose a method that includes the reported storm loss while using expected claims to temper future development.',
 ('fall-2014',24,'b'):'Show why applying historical reported factors to the storm-inflated claim amount overstates ultimate.',
 ('fall-2014',24,'c'):'Show why a paid-only method misses storm claims that were reported but not yet paid.',
}

def main(exam):
    cfg=CONFIG[exam]
    book=load_workbook(ROOT/'past_exams_excel'/cfg['workbook'],data_only=False)
    grid=book['Point Grid']
    pdf=PdfReader(ROOT/'past_exams'/cfg['pdf'])
    questions=[]
    for n in range(1,cfg['count']+1):
        sh=book[str(n)]
        markers=[(c.row,c.value.strip()[1].lower()) for row in sh for c in row
            if c.column==1 and isinstance(c.value,str) and re.fullmatch(r'\([a-z]\)',c.value.strip(),re.I)]
        assert markers,(exam,n,'no parts')
        source,covered=blocks(exam,n,sh,markers[0][0])
        answers,notes=report(exam,n,pdf)
        parts=[]
        for i,(start,letter) in enumerate(markers):
            end=markers[i+1][0] if i+1<len(markers) else sh.max_row+1
            points=grid.cell(n+7,ord(letter)-ord('a')+3).value
            assert isinstance(points,(int,float)),(exam,n,letter,'points',points)
            prompt=' '.join(row_text(exam,n,sh,row,covered.get(row,set())) for row in range(start,end) if row_text(exam,n,sh,row,covered.get(row,set())))
            answer=SOLUTION_FIXES.get((exam,n,letter),answers.get(letter,''))
            insight=INSIGHT_FIXES.get((exam,n,letter),notes.get(letter,''))
            if len(markers)==1:
                answer=answer or answers.get('a','')
                insight=insight or notes.get('a','')
            assert prompt and answer and insight,(exam,n,letter,'missing prompt/answer/insight',bool(prompt),bool(answer),bool(insight),list(answers),list(notes))
            parts.append({'id':letter,'points':points,'prompt':prompt,'solution':answer,'insight':insight})
        first=cfg['report_start'][n-1]
        last=cfg['report_start'][n] if n<cfg['count'] else len(pdf.pages)+1
        q={'id':f'{exam}-{n}','number':n,'exam':exam.replace('-',' ').title(),
           'chapterIds':CHAPTERS[exam][n-1], 'points':sum(p['points'] for p in parts),
           'questionPage':n+3,'solutionPages':list(range(first,last)),
           'sourceBlocks':source,'parts':parts}
        if exam=='spring-2015' and n==20:
            q['notice']='The examiner report identifies a typo in this question’s exposure base; it accepted reasonable alternative assumptions for the Bornhuetter–Ferguson expected claims calculation.'
        if exam=='spring-2015' and n==19:
            q['notice']='The examiner report’s sample answer contains arithmetic transcription errors in the reported development ratios and the selected ultimate for 2013. The worked answer uses the figures printed in the question.'
        if exam=='fall-2014' and n==10:
            q['figure']={'src':'assets/exam-graphs/fall-2014-q10.png','title':'GLM claim frequency by hazard class',
                'alt':'Claim frequency for hazard classes A through G in 2011–2013, with total exposures'}
        questions.append(q)
    assert len(questions)==cfg['count']
    assert sum(len(q['parts']) for q in questions)==cfg['part_count']
    assert abs(sum(q['points'] for q in questions)-cfg['points'])<1e-9
    path=ROOT/'question_data'/(exam.replace('-','')+'_questions.js')
    path.write_text(f'// Transcribed from the official {exam.replace("-"," ").title()} exam PDF; point grid checked against the CBT workbook.\nwindow.{cfg["var"]} = '+json.dumps(questions,ensure_ascii=False,indent=2)+';\n')
    print(exam,len(questions),sum(len(q['parts']) for q in questions),sum(q['points'] for q in questions),path)

if __name__=='__main__':
    for exam in CONFIG:main(exam)
