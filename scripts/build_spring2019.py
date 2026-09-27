"""Build the Spring 2019 question data from the supplied CBT workbook and report.

Run with a Python environment containing openpyxl and pypdf. The workbook is
the source for prompts and point values; the official report supplies one
sample solution and the examiner's notes for each part.
"""

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "past_exams_excel/CBT_Exam_5_S.19_v2.xlsx"
REPORT = ROOT / "past_exams/admissions_studytools_exam5_sp19-5.pdf"
OUTPUT = ROOT / "question_data/spring2019_questions.js"
GRAPH = ROOT / "assets/exam-graphs/spring-2019-q9.png"

SINGLE_PROMPT_ROW = {9: 34, 10: 20, 11: 12, 14: 15, 19: 14, 22: 42, 23: 34}
CHAPTERS = {
    1: ["ratemaking-4"], 2: ["ratemaking-5"],
    3: ["ratemaking-5", "ratemaking-6"], 4: ["ratemaking-6"],
    5: ["ratemaking-6"], 6: ["ratemaking-7"],
    7: ["ratemaking-8", "reserving-9"], 8: ["ratemaking-13"],
    9: ["ratemaking-10"], 10: ["ratemaking-9"],
    11: ["ratemaking-15"], 12: ["ratemaking-8", "ratemaking-13"],
    13: ["reserving-1"], 14: ["reserving-6", "reserving-7"],
    15: ["reserving-11"], 16: ["reserving-7"],
    17: ["reserving-9"], 18: ["reserving-10"],
    19: ["reserving-7", "reserving-9"], 20: ["reserving-11"],
    21: ["reserving-15"], 22: ["reserving-7", "reserving-13"],
    23: ["reserving-14"], 24: ["reserving-17"],
}

# Table boundaries and complete headings checked against exam pages 1–24 of
# the scanned official PDF. Excel often spreads one printed heading across rows.
# Each specification is (first data row, last data row, title, columns, headings,
# optional grouped first header row). Column numbers are Excel's one-based ones.
TABLES = {
  1: [(8,12,"",[2,3,4,5],["Policy","Number of Vehicles","Effective Date","Expiration Date"])],
  2: [(6,7,"",[2,3],["Calendar Year","Earned Premium ($)"]),
      (12,13,"",[2,3],["Rate Change Effective Date","Overall Rate Change"]),
      (19,24,"",[2,3],["Quarter and Year","Average Written Premium at Current Rate Level"])],
  4: [(7,10,"",[2,3],["Claim Number","Total Limits Loss ($000s)"])],
  5: [(9,15,"Cumulative Reported Losses ($000s)",[2,3,4,5,6],["Accident Year","12","24","36","48"]),
      (9,15,"Shock Losses",[8,9,10],["Accident Year","Claim Count","Reported Ground-Up Losses ($000s)"]),
      (20,20,"Selected Age-to-Age Development Factors",[2,3,4],["12–24","24–36","36–48"]),
      (29,34,"Exponential Trend",[2,3,4],["Basis","Frequency","Total Severity"])],
  6: [(8,13,"Countrywide information",[2,6,7,8],["Item","2016 Expense Ratio","2017 Expense Ratio","2018 ($000s)"])],
  7: [(7,9,"",[2,3],["Accident Year","Reported Loss and ALAE ($000s)"]),
      (7,9,"",[6,7],["Calendar Year","Earned Premium ($000s)"]),
      (13,15,"",[2,3],["Development Age to Ultimate","Selected Cumulative Development Factors"])],
  8: [(12,13,"Strategy 1 — No Change",[2,3,5,7],["Class","Number of Risks","Losses & Expenses per Risk","Rate per Risk"]),
      (19,20,"Strategy 2 — New Rating Variable",[2,3,5,7],["Class","Number of Risks","Losses & Expenses per Risk","Rate per Risk"]),
      (26,27,"Strategy 3 — Rate Increase for All Risks",[2,3,5,7],["Class","Number of Risks","Losses & Expenses per Risk","Rate per Risk"])],
  10:[(7,9,"Exposures",[2,3,4,5],["Vehicle Class","Territory 1","Territory 2","Territory 3"]),
      (7,8,"Exposures",[7,8,9,10],["Driver Type","Territory 1","Territory 2","Territory 3"]),
      (13,15,"Loss ($000s)",[2,3,4,5],["Vehicle Class","Territory 1","Territory 2","Territory 3"]),
      (13,14,"Loss ($000s)",[7,8,9,10],["Driver Type","Territory 1","Territory 2","Territory 3"])],
  12:[(8,12,"Number of Policies Renewed",[2,3,4],["Rate Change","Territory A","Territory B"]),
      (16,18,"",[2,3,4],["Measure","Territory A","Territory B"])],
  13:[(10,12,"State A",[2,3,4,5],["Accident Year","Earned Exposure","Ultimate Claim Count — BI","Ultimate Claim Count — PD"]),
      (10,12,"State B",[7,8,9,10],["Accident Year","Earned Exposure","Ultimate Claim Count — BI","Ultimate Claim Count — PD"]),
      (16,17,"State A — Paid Age-to-Ultimate Factors",[2,3,4,5],["Coverage","12-to-Ult","24-to-Ult","36-to-Ult"]),
      (16,17,"State B — Paid Age-to-Ultimate Factors",[7,8,9,10],["Coverage","12-to-Ult","24-to-Ult","36-to-Ult"]),
      (21,22,"State A — Ultimate Severity",[2,3],["Coverage","Ultimate Severity"]),
      (21,22,"State B — Ultimate Severity",[7,8],["Coverage","Ultimate Severity"])],
  14:[(8,11,"Average Case Outstanding",[2,3,4,5,6],["Accident Year","12","24","36","48"]),
      (8,11,"Average Reported Claims",[8,9,10,11,12],["Accident Year","12","24","36","48"])],
  15:[(7,10,"Cumulative Reported Claim Counts as of (months)",[2,3,4,5,6],["Accident Year","12","24","36","48"]),
      (14,17,"Cumulative Reported Claims ($000s) as of (months)",[2,3,4,5,6],["Accident Year","12","24","36","48"]),
      (21,23,"Reported Claim Count Age-to-Age Factors",[2,3,4,5],["Accident Year","12–24","24–36","36–48"])],
  16:[(7,10,"Cumulative Reported Claims ($000s) as of (months)",[2,3,4,5,6],["Accident Year","12","24","36","48"])],
  17:[(8,10,"Reported Claims",[2,3,5,7],["Accident Year","Region 1","Region 2","Combined"]),
      (14,15,"Reported Age-to-Ultimate Factors",[2,3,5],["Age","Region 1","Region 2"]),
      (19,21,"Earned Premium",[2,3,5,7],["Calendar Year","Region 1","Region 2","Combined"]),
      (25,26,"Annual Industry Trends",[2,3,5,7],["Trend","Region 1","Region 2","Combined"])],
  18:[(8,10,"Reported Claims Data",[2,3,4,5,6],["Accident Year","Reported Claims","Age-to-Ultimate","Earned Premium","Pure Premium Trend Factors"])],
  20:[(7,9,"Incremental Closed Claim Counts as of (months)",[2,3,4,5],["Accident Year","72","84","96"]),
      (13,15,"Incremental Paid Claims ($000s) as of (months)",[2,3,4,5],["Accident Year","72","84","96"])],
  21:[(8,10,"Net of Reinsurance Reported Claims ($000s)",[2,3,4,5,6],["Accident Year","12","24","36","Ultimate Frequency"])],
  22:[(8,11,"",list(range(2,11)),["Accident Year","12","24","36","48","12","24","36","48"],
       [("",1),("Average Case Outstanding ($) on Large Claims as of (months)",4),("Average Case Outstanding ($) on Small Claims as of (months)",4)]),
      (15,18,"",list(range(2,11)),["Accident Year","12","24","36","48","12","24","36","48"],
       [("",1),("Open Large Claim Counts as of (months)",4),("Open Small Claim Counts as of (months)",4)]),
      (23,26,"",list(range(2,11)),["Accident Year","12","24","36","48","12","24","36","48"],
       [("",1),("Cumulative Paid Claims ($) on Large Claims as of (months)",4),("Cumulative Paid Claims ($) on Small Claims as of (months)",4)]),
      (32,33,"Cumulative Unadjusted Reported Claim Development Factors",[2,3,4,5,6],["Age","12-to-Ult","24-to-Ult","36-to-Ult","48-to-Ult"])],
  23:[(8,11,"",[2,3,4],["Accident Year","Per Occurrence Retention","Stop-Loss Limit"]),
      (19,22,"Summary of Claims Under Per Occurrence Retention",[2,3,4],["Accident Year","Reported Claims","Percent Reported"]),
      (27,30,"Large Claims (not included in the claims table above)",[2,3,4],["Claim","Accident Year","Reported Claims"])],
  24:[(5,8,"Cumulative Paid Claims Only as of (months)",[2,3,4,5,6],["Accident Year","12","24","36","48"]),
      (12,15,"Cumulative Paid ALAE as of (months)",[2,3,4,5,6],["Accident Year","12","24","36","48"]),
      (20,23,"",[2,3],["Accident Year","Selected Ultimate Claims Only"])],
}

HEADER_ROWS = {
  1:[6], 2:[5,10,16], 4:[5,6], 5:[5,6,7,8,17,19,27,28],
  6:[5,6,7], 7:[5,6,11,12], 8:[9,10,16,17,23,24],
  10:[5,6,11,12], 12:[5,6,7,14,15],
  13:[6,8,9,14,15,19,20], 14:[5,6,7],
  15:[5,6,12,13,19,20], 16:[5,6],
  17:[6,7,12,13,17,18,23,24], 18:[5,6,7],
  20:[5,6,11,12], 21:[5,6,7],
  22:[5,6,7,13,14,20,21,22,30,31],
  23:[5,6,7,16,17,18,24,25,26], 24:[3,4,10,11,17,18,19],
}

SOLUTION_OVERRIDES = {
  (2,"a"): "The 2017 earned premium's average historical rate level is 0.25(1.00) + 0.50(1.10) + 0.25(1.155) = 1.08875. The current rate level is 1.155, so the on-level factor is 1.155 / 1.08875 = 1.06085. On-level earned premium is about $4.084 million. Trend the average 2017 earned date to the July 1, 2019 rate period: using about 4% annual premium trend gives roughly $4.55 million in trended on-level earned premium. Reasonable trend selections and timing assumptions were credited.",
  (4,"a"): "At the $25,000 basic limit, original basic-limit losses are 15 + 21 + 24 + 25 = 85 ($000s). After 8% trend, basic-limit losses are 16.2 + 22.68 + 25 + 25 = 88.88 ($000s). Basic-limit trend = 88.88 / 85 − 1 = 4.565%.",
  (7,"b"): "Trend and on-level the latest three accident years, then compare projected loss and ALAE (including the 8% ULAE provision) with projected premium. The sample calculation produces a 66.4% loss and ALAE ratio. The indicated change is [0.664 × 1.08 + 0.06] / [1 − 0.26 − 0.05] − 1 ≈ 12.63%. Applying 70% credibility with an 8% complement gives 0.70 × 12.63% + 0.30 × 8% ≈ 11.24%.",
  (8,"a"): "Strategy 1: 5,000 × (1,050 − 900) + 10,000 × (1,050 − 1,000) = $1,250,000. Strategy 2: 5,000 × (1,000 − 900) + 10,000 × (1,100 − 1,000) − 500,000 = $1,000,000. Strategy 3: 5,000 × (1,100 − 900) + 10,000 × (1,100 − 1,000) − 500,000 = $1,500,000.",
  (9,"a"): "Model 1 closely follows actual losses in the training sample, but its hold-out predictions vary sharply and exceed actual losses in the highest deciles, suggesting overfitting. Model 2 is more stable between the two samples, but its flatter predictions miss the steep rise in actual losses for the highest-risk deciles. The graphs support a trade-off between fit and stability.",
  (10,"a"): "Compare loss per exposure by territory using driver-type data, whose exposure mix is proportional across territories. Relative to Territory 1, the indicated territory factors are 1.000, 1.7013 and 0.9821. Then balance vehicle-class and driver-type factors for their correlated exposure mix. With Vehicle A and Driver X as bases, one acceptable set is Vehicle A 1.000, B 1.2551, C 1.6020; Driver X 1.000, Y 1.7755. Other justified univariate balancing approaches were credited.",
  (11,"a"): "(i) The insurer may still adjust and administer claims below the deductible, so those expenses remain. (ii) The insurer pays claims and seeks reimbursement from the insured; bankruptcy creates credit risk. (iii) The large-deductible layer has fewer, more volatile claims and greater uncertainty, which can justify a higher profit margin.",
  (15,"b"): "Claims closed without payment, reopened claims, or corrections to duplicate/erroneous reports can cause cumulative reported claim counts to fall between development ages.",
  (17,"a"): "Region 1's 2016 expected claims ratio is 2 × 45.8% − 40% = 51.6%. Trend to 2017: 51.6% × 0.98 × 1.10 = 55.6248%. Reported Bornhuetter–Ferguson ultimate = 150,000 + (1 − 1/1.154) × 1,000,000 × 55.6248% ≈ 224,231.",
  (17,"b"): "Trend Region 1's expected claims ratio to 2018: 51.6% × (0.98 × 1.10)² ≈ 59.96% (about 60%); Region 2 remains at 40%. Combined reported Bornhuetter–Ferguson ultimate = 180,000 + 1,610,000 × 59.96% × (1 − 1/2.283) + 690,000 × 40% × (1 − 1/1.558) ≈ 821,400. Rounding Region 1's ratio to 60% gives 821,722.",
  (20,"a"): "Trend paid claims at ages 84 and 96 to 2018 cost levels, adjusting pre-2012 claims for the 20% legislative reduction. Tail severity = [(6,100 + 2,400) × 1.06⁷ × 0.80 + 3,900 × 1.06⁶] / (81 + 13 + 61) ≈ 101.66 ($000s) per closed claim.",
  (22,"a"): "Adjust the historical large-claim case outstanding amounts for the 5% severity trend, then add them to cumulative paid claims to create comparable reported-claim triangles. The adjusted large-claim 2018 reported amount is 1,200 × 10 + 525,000 = 537,000. Selected large-claim age-to-ultimate factor is about 2.605, giving $1,398,856. Small-claim reported amount is 90 × 150 + 22,000 = $35,500; at the unadjusted 9.007 factor, its ultimate is $319,741. Total estimated 2018 ultimate claims are $1,718,596. The examiner also accepted the alternate unit interpretation noted with the question.",
}

PROMPT_OVERRIDES = {
  (8,"b"): "Briefly evaluate, for strategies 2 and 3: (i) the assumption that the number of risks by class will be the same as under strategy 1; and (ii) the impact this assumption has on each strategy's expected total profit.",
}


def displayed(cell):
    value = cell.value
    if value is None or cell.data_type == "f":
        return ""
    if isinstance(value, str):
        return " ".join(value.replace("\uf0b7", "•").split())
    if isinstance(value, (int, float)):
        fmt = cell.number_format
        if "%" in fmt:
            decimals = len(re.search(r"0(?:\.(0+))?%", fmt).group(1) or "") if re.search(r"0(?:\.(0+))?%", fmt) else 0
            return f"{value * 100:,.{decimals}f}%"
        decimals = len(re.search(r"0\.(0+)", fmt).group(1)) if re.search(r"0\.(0+)", fmt) else None
        if decimals is not None:
            number = f"{value:,.{decimals}f}" if "#,##0" in fmt else f"{value:.{decimals}f}"
        elif "#,##0" in fmt:
            number = f"{value:,.0f}"
        else:
            number = f"{value:g}"
        return f"${number}" if "$" in fmt else number
    return str(value)


def cleanup(text):
    text = text.replace("\xa0", " ").replace("\uf0b7", "• ")
    text = re.sub(r"(?m)^EXAM 5 SPRING 2019 – SAMPLE ANSWERS AND EXAMINER.S REPORT\s*$", "", text)
    text = re.sub(r"(?m)^\s+$", "", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def report_sections():
    pdf = PdfReader(REPORT)
    chunks = []
    text = ""
    for page_number, page in enumerate(pdf.pages, 1):
        chunks.append((len(text), page_number))
        text += (page.extract_text() or "").replace("\xa0", " ") + "\n"
    starts = list(re.finditer(r"QUESTION\s+(\d+)\s+TOTAL\s+POINT\s+VALUE", text, re.I))
    result = {}
    for index, match in enumerate(starts):
        number = int(match.group(1))
        section = cleanup(text[match.start():starts[index + 1].start() if index + 1 < len(starts) else None])
        start_page = max(page for offset, page in chunks if offset <= match.start())
        end_page = (max(page for offset, page in chunks if offset <= starts[index + 1].start()) - 1) if index + 1 < len(starts) else len(pdf.pages)
        sample, examiner = re.split(r"(?m)^EXAMINER.S REPORT\s*$", section, maxsplit=1)
        sample = sample.split("SAMPLE ANSWERS", 1)[-1].strip()
        result[number] = {"sample": sample, "examiner": examiner.strip(), "pages": list(range(start_page, end_page + 1))}
    assert sorted(result) == list(range(1, 25)), "The report does not contain all 24 answers"
    return result


def split_parts(text, sample):
    pattern = r"(?m)^Part\s+([a-g])\s*:\s*"
    if not sample:
        pattern = r"(?m)^Part\s+([a-g])\s*:?[ \t]*$"
    matches = list(re.finditer(pattern, text, re.I))
    if not matches:
        return {"a": cleanup(text)}
    return {
        match.group(1).lower(): cleanup(text[match.end():matches[index + 1].start() if index + 1 < len(matches) else None])
        for index, match in enumerate(matches)
    }


def solution_text(text):
    text = re.sub(r"^(?:\d+(?:\.\d+)?\s+points?\s*)", "", text)
    if re.search(r"(?m)^Sample\s+1\s*$", text):
        text = re.split(r"(?m)^Sample\s+1\s*$", text, maxsplit=1)[-1]
        text = re.split(r"(?m)^Sample\s+2\s*$", text, maxsplit=1)[0]
    return cleanup(text)


def worksheet_parts(sheet, point_grid, question_number, report):
    markers = [
        (cell.row, cell.value.strip()[1].lower())
        for row in sheet for cell in row
        if cell.column == 1 and isinstance(cell.value, str) and re.fullmatch(r"\([a-g]\)", cell.value.strip())
    ]
    if not markers:
        markers = [(SINGLE_PROMPT_ROW[question_number], "a")]
    sample = split_parts(report["sample"], True)
    examiner = split_parts(report["examiner"], False)
    parts = []
    for index, (start, letter) in enumerate(markers):
        stop = markers[index + 1][0] if index + 1 < len(markers) else sheet.max_row + 1
        lines = []
        for row in sheet.iter_rows(min_row=start, max_row=stop - 1):
            values = [displayed(cell) for cell in row if cell.column > 1 and cell.value is not None and cell.data_type != "f"]
            lines.extend(value for value in values if value)
        point_value = point_grid.cell(question_number + 7, ord(letter) - ord("a") + 3).value
        assert isinstance(point_value, (int, float)), (question_number, letter, point_value)
        assert letter in sample, (question_number, letter, "missing sample answer")
        parts.append({
            "id": letter,
            "points": point_value,
            "prompt": PROMPT_OVERRIDES.get((question_number, letter), " ".join(lines)),
            "solution": SOLUTION_OVERRIDES.get((question_number, letter), solution_text(sample[letter])),
            "insight": cleanup(examiner.get(letter, examiner.get("a", "See the examiner's report for grading guidance."))),
        })
    return markers[0][0], parts


def source_blocks(sheet, before_row, number):
    specs = TABLES.get(number, [])
    blocks = []
    data_rows = {row for spec in specs for row in range(spec[0], spec[1] + 1)}
    headers = set(HEADER_ROWS.get(number, []))
    previous_line_row = None
    for row_number in range(3, before_row):
        for spec in specs:
            first, last, title, columns, labels, *group = spec
            if first != row_number:
                continue
            rows = []
            for index in range(first, last + 1):
                values = []
                for position, column in enumerate(columns):
                    cell = sheet.cell(index, column)
                    value = displayed(cell)
                    if position and isinstance(cell.value, (int, float)) and cell.value >= 1000 and cell.value == int(cell.value) and "%" not in cell.number_format:
                        value = f"{cell.value:,.0f}"
                    values.append(value)
                if number == 1 and index == 12:
                    values[2:] = ["November 1, 2018", "April 30, 2019"]
                rows.append(values)
            table = {"type": "table", "title": title, "headers": labels, "rows": rows}
            if group:
                table["groups"] = [{"label": label, "span": span} for label, span in group[0]]
            blocks.append(table)
        if row_number in data_rows or row_number in headers:
            continue
        values = [displayed(cell) for cell in sheet[row_number][1:] if cell.value is not None and cell.data_type != "f"]
        values = [value for value in values if value]
        if values:
            line = " — ".join(values) if len(values) > 1 else values[0]
            if number == 7:
                line = line.replace("July 1. 2018", "July 1, 2018")
            if number == 5 and line.endswith("for an insurer"):
                line += ":"
            if number == 8 and line.endswith("increase rates for all risks"):
                line += "."
            if number == 23 and line.startswith("***2015"):
                line = line.replace("***2015", "*** 2015", 1)
            if (blocks and blocks[-1]["type"] == "line" and previous_line_row == row_number - 1
                    and not blocks[-1]["text"].rstrip().endswith((".", ":", "!", "?"))
                    and not line.startswith(("•", "i.", "ii.", "iii."))):
                blocks[-1]["text"] += " " + line
            else:
                blocks.append({"type": "line", "text": line})
            previous_line_row = row_number
    if number == 24:
        blocks.insert(0, {"type": "line", "text": "Given the following information:"})
    return blocks


def main():
    workbook = load_workbook(WORKBOOK, data_only=False)
    report = report_sections()
    point_grid = workbook["Point Grid"]
    questions = []
    for number in range(1, 25):
        sheet = workbook[str(number)]
        before_row, parts = worksheet_parts(sheet, point_grid, number, report[number])
        total = sum(part["points"] for part in parts)
        assert abs(total - sum(point_grid.cell(number + 7, col).value or 0 for col in range(3, 10))) < 0.001
        question = {
            "id": f"spring-2019-{number}", "number": number, "exam": "Spring 2019",
            "chapterIds": CHAPTERS[number], "points": total,
            "questionPage": number + 4, "solutionPages": report[number]["pages"],
            "sourceBlocks": source_blocks(sheet, before_row, number),
            "parts": parts,
        }
        if number == 9:
            question["figure"] = {
                "src": "assets/exam-graphs/spring-2019-q9.png",
                "title": "Training and hold-out results for two GLMs",
                "alt": "Two graphs compare predictions from Models 1 and 2 with actual losses across ten equal-volume deciles in training and hold-out data.",
            }
        if number == 22:
            question["notice"] = "The examiner's report notes that the average case outstanding triangles were labeled as dollars instead of thousands of dollars. Either interpretation was credited."
        questions.append(question)
    GRAPH.parent.mkdir(parents=True, exist_ok=True)
    GRAPH.write_bytes(workbook["9"]._images[0]._data())
    OUTPUT.write_text("// Generated from the supplied Spring 2019 CBT workbook and examiner's report.\nwindow.SPRING_2019_QUESTIONS = " + json.dumps(questions, ensure_ascii=False, indent=2) + ";\n")
    print(f"Wrote {len(questions)} questions, {sum(len(q['parts']) for q in questions)} scored parts, {sum(q['points'] for q in questions)} points")


if __name__ == "__main__":
    main()
