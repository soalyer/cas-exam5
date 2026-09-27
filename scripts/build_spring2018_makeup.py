"""Rebuild the Spring 2018 makeup CBT sitting from its workbook and report.

The examiner's report gives grading expectations, not model answers. Solutions
below are independently worked. Question 21 was excluded from official scoring.
The Q12 PNG is a direct LibreOffice rendering of xl/media/image1.emf from the
workbook, cropped to the chart without changing its plotted data.
"""
import json
import re
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from build_spring2018 import cell_text, row_text

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / 'past_exams/april_may_sp18-5-makeup.xls'
REPORT = ROOT / 'past_exams/april_may_sp18-5-makeup_examiners_report.pdf'
OUTPUT = ROOT / 'question_data/spring2018_makeup_questions.js'

CHAPTERS = {
  1:['ratemaking-5'], 2:['ratemaking-6'], 3:['ratemaking-6'],
  4:['ratemaking-6'], 5:['ratemaking-7'], 6:['ratemaking-8','reserving-9'],
  7:['ratemaking-9'], 8:['ratemaking-9'], 9:['ratemaking-12'],
  10:['ratemaking-11'], 11:['ratemaking-7','ratemaking-8','ratemaking-12'],
  12:['ratemaking-15'], 13:['ratemaking-8'], 14:['reserving-1'],
  15:['reserving-7'], 16:['reserving-7'], 17:['reserving-11'],
  18:['reserving-10'], 19:['reserving-7'], 20:['reserving-6'],
  21:['reserving-13'], 22:['reserving-14'], 23:['reserving-17'],
  24:['reserving-7','reserving-9','reserving-16'], 25:['reserving-15'], 26:['reserving-6'],
}
SINGLE = {7:7, 8:21, 17:25, 18:21, 21:37, 23:27}
# (covered start, covered end, row numbers, column numbers, headings, title)
TABLES = {
  1:[(6,16,range(7,17),[2,3,4,5],['Policy','Transaction date','Action','Full-term written premium'],'')],
  2:[(6,12,range(7,13),[2,3,4,5,6,7],['Policy','Effective date','Accident date','Transaction date','Incremental payment','Ending case reserve'],'')],
  3:[(6,9,range(7,10),[2,3,4,5,6],['Year','Earned premium ($000)','On-level factor','Ultimate ground-up losses ($000)','Losses excess $750,000 ($000)'],'')],
  4:[(6,11,range(8,12),[2,3,4,5,6,7],['Accident year','Earned exposures','12','24','36','48'],'Cumulative reported claims ($) by age in months')],
  5:[(6,10,range(8,11),[2,3,4,5,6,7,8,9],['Year','Policy count','Written premium','Earned premium','Commission','Other acquisition','Taxes/fees','General expense'],'Countrywide ($000 except count)')],
  6:[(7,9,range(8,10),[2,3],['Effective date','Rate change'],'Rate changes'),
     (11,12,[12],[2,3,4,5],['Measure','2015','2016','2017'],'Calendar-year premium ($)'),
     (14,18,range(16,19),[2,3,4,5],['Accident year','12','24','36'],'Cumulative paid loss and ALAE ($)')],
  8:[(7,13,range(10,14),[2,3,4,5,6,7,8,9,10,11,12,13,14],['Class','Collected premium','On-level factor','Current premium','Reported loss','Loss ratio','Indicated change','Current relativity','Indicated relativity','Relative to base','Selected relativity','Relativity change','Final change'],'Class analysis')],
  9:[(6,18,range(7,19),[2,3,4,5,6],['State','Class','Exposure','Losses','Pure premium'],'')],
  10:[(6,10,range(8,11),[2,4,5,6,7],['Loss size','Claims, $100k limit','Losses ($000), $100k','Claims, $250k limit','Losses ($000), $250k'],'Censored losses')],
  11:[(6,9,range(7,10),[2,3,4,5,6,7,8],['Year','Written premium','Earned premium','Ultimate claims','General expense','Acquisition expense','Taxes/fees'],'All values $000'),
      (11,14,range(12,15),[2,3,4,5,6],['Territory','In-force premium ($000)','Current relativity','Loss ratio','Claim count'],'')],
  15:[(6,20,range(8,21),[2,3,4,5,6,7],['Claim','Accident date','Report date','Evaluation date','Case reserve','Cumulative paid'],'')],
  16:[(6,12,range(8,13),[2,3,4,5,6,7],['Accident year','12','24','36','48','60'],'Cumulative reported claims ($000)')],
  17:[(6,11,range(8,12),[2,3,4,5,6],['Accident year','12','24','36','48'],'Cumulative closed claim counts'),
      (6,11,range(8,12),[8,9,10,11,12],['Accident year','12','24','36','48'],'Incremental paid severities'),
      (13,18,range(15,19),[2,3],['Accident year','Ultimate claim counts'],'')],
  18:[(6,11,range(8,12),[2,3,4,5,6],['Accident year','12','24','36','48'],'Cumulative reported claims ($)'),
      (13,17,range(14,18),[2,3],['Calendar year','On-level earned premium'],'')],
  19:[(6,17,range(8,18),[2,3,4],['Calendar year','Accident year','Incremental reported claims'],'')],
  20:[(6,12,range(8,13),[2,3,4,5,6,7],['Accident year','12','24','36','48','60'],'Cumulative paid loss + ALAE ($000)'),
      (6,12,range(8,13),[9,10,11,12,13,14],['Accident year','12','24','36','48','60'],'Cumulative reported loss + ALAE ($000)'),
      (14,20,range(16,21),[2,3,4,5,6,7],['Accident year','12','24','36','48','60'],'Cumulative closed counts'),
      (14,20,range(16,21),[9,10,11,12,13,14],['Accident year','12','24','36','48','60'],'Cumulative reported counts')],
  21:[(6,11,range(8,12),[2,3,4,5,6],['Accident year','12','24','36','48'],'Cumulative closed counts'),
      (13,18,range(15,19),[2,3,4,5,6],['Accident year','12','24','36','48'],'Cumulative paid claims ($)'),
      (20,24,range(22,25),[2,3,4,5,6],['Accident year','12','24','36','48'],'Exponential parameter a'),
      (26,30,range(28,31),[2,3,4,5,6],['Accident year','12','24','36','48'],'Exponential parameter b')],
  22:[(6,12,range(9,13),[2,3,4,5,6],['Accident year','12','24','36','48'],'Cumulative paid gross claims ($000)'),
      (14,20,range(17,21),[2,3,4,5,6],['Accident year','12','24','36','48'],'Cumulative salvage/subrogation received ($000)')],
  23:[(6,13,range(10,14),[2,3,5,6,7],['Accident year','Ultimate claim counts','Age','Reported % ultimate','Closed % ultimate'],'Development pattern'),
      (16,21,range(18,22),[2,3,4,5,6],['Accident year','12','24','36','48'],'Pending claim counts')],
  24:[(6,13,range(9,14),[2,3,4,5,6,7,8],['Accident year','12','24','36','48','60','Selected ultimate claims ($)'],'Cumulative paid ALAE ($)')],
  25:[(6,12,range(9,13),[2,3,4,5,6,7],['Accident year','Paid development','Reported development','Paid BF','Reported BF','Expected'],'Ultimate claims ratios')],
  26:[(6,11,range(9,12),[2,3,4],['Accident year','Dec 31, 2016','Jun 30, 2017'],'Cumulative reported claims ($000)'),
      (13,16,range(15,17),[2,3,4,5,6],['Pattern','12–ultimate','24–ultimate','36–ultimate','48–ultimate'],'Reported claim factors')],
}

def source_blocks(sheet, number, stop):
    specs = TABLES.get(number, [])
    covered = {r for first,last,*_ in specs for r in range(first,last+1)}
    blocks = []
    previous_line_row = None
    for row in range(4, stop):
        for first,_,rows,cols,headers,title in specs:
            if row == first:
                blocks.append({'type':'table','title':title,'headers':headers,
                               'rows':[[cell_text(sheet.cell(r,c)) for c in cols] for r in rows]})
        if row in covered:
            previous_line_row = None
            continue
        line = row_text(sheet,row)
        if line:
            if (previous_line_row == row-1 and blocks[-1]['type'] == 'line'
                    and not blocks[-1]['text'].endswith(('.',':',';','?','!'))
                    and not line.startswith(('•','i.','ii.','iii.'))):
                blocks[-1]['text'] += ' '+line
            else:
                blocks.append({'type':'line','text':line})
            previous_line_row = row
        else:
            previous_line_row = None
    return blocks

def report_pages():
    reader = PdfReader(REPORT)
    starts=[]; pts={}
    for page_no,page in enumerate(reader.pages,1):
        t=page.extract_text() or ''
        m=re.search(r'(?im)^QUESTION\s*:?\s*(\d+)\s*$',t)
        if m:
            n=int(m.group(1)); starts.append((n,page_no))
            pts[n]=float(re.search(r'TOTAL POINT VALUE:\s*([\d.]+)',t).group(1))
    assert [n for n,_ in starts] == list(range(1,27))
    return {n:list(range(p, starts[i+1][1] if i+1<len(starts) else len(reader.pages)+1)) for i,(n,p) in enumerate(starts)},pts

# Each tuple is an independently worked answer and a report-grounded insight.
ANSWERS = {
1:[('In force: A $100, B $310 (the $60 full-term endorsement is not prorated for this snapshot), C $150, and D $200. Total $760.','A policy snapshot includes the full-term value of the B endorsement and excludes expired D 2014 coverage.'),
   ('Policy-year 2015 written premium: A $100 + B [$250 + $60×6/12 + $50 audit] + C [$150 − $150×7/12] + D $200 = $692.50.','Prorate the B endorsement and C cancellation over the affected term; assign B’s later audit to its 2015 policy year.')],
2:[('Calendar-year 2016 incurred is $4,000 on A. Calendar-year 2017 change on A is $4,500−$4,000=$500; B adds $6,000, for $6,500. Ending case reserves are stocks, not incremental transactions.','Compute each calendar year from changes in cumulative paid plus ending case.'),
   ('At December 31, 2017, accident year 2016 has A: $4,500 paid plus zero case. Accident year 2017 has B: $3,000 paid plus $3,000 case = $6,000.','Stop at the requested valuation date; B’s January 2018 payment is excluded.'),
   ('Both claims arise from policies effective during 2016, so 2016 policy-year incurred is $4,500+$6,000=$10,500 at December 31, 2017.','Group by policy effective year, including B’s 2017 accident.'),
   ('Advantage: calendar-year losses are current and need no loss-development projection. Disadvantage: reserve changes and payments from many accident years can make one calendar year unrepresentative for prospective pricing.','Address ratemaking, including calendar-year reserve changes, rather than a reserving use case.')],
3:[('Limit each year at $750,000: ground-up less excess. The 10-year excess load is 1 + 45,221/(559,996−45,221) = 1.08785. Trend limited losses at 4% for 4, 3, and 2 years to the July 2019 average earned date, then apply the load. Trended loaded losses are about $70,654, $66,141, $67,091 thousand; on-level premiums are $92,927.52, $98,202.30, $95,038.98 thousand. Weighted ratio = 203,885.3/286,168.8 = 71.25%.','Use the limited trend, an excess-loss load based on the 10-year limited base, and the correct prospective trend periods.'),
   ('Yes. Large claims are volatile, so limiting annual experience and replacing the excess with a longer-term load stabilizes the indication while retaining an expected large-loss cost.','Both the limiting step and the long-term load need a purpose.')],
4:[('Volume-weighted reported age factors are about 1.350, 1.150, and 1.100; applying them gives 2014–2017 ultimate loss costs about $691.41, $705.24, $719.35, $733.73 per exposure. These rise approximately 2.0% annually.','Develop losses and divide by exposures before selecting loss-cost trend.'),
   ('Trend the 2017 loss cost 1.5 years to the average accident date of 2018 annual policies, January 2019, and apply reform: $733.73×1.02^1.5×1.20 ≈ $907.03 per exposure. A consistently trended all-year selection gives the same result.','Use a prospective earned date and the 1.20 reform multiplier.'),
   ('Two drivers are claim severity inflation (repair materials or medical costs) and claim frequency changes (weather, traffic, safety, or exposure mix).','Name causes of loss cost rather than expense or premium changes.')],
5:[('Using 2015–2017 countrywide data, other acquisition/written premium is 18,433/300,500=6.13% and general/earned premium is 24,059/297,250=8.09%. Add State A commission 7.0% and taxes/fees 6.2%: all-variable underwriting expense provision ≈27.42%.','Use countrywide other acquisition and general expense with their proper bases, and State A commission and taxes.'),
   ('Treat 75% of other acquisition and general expenses as fixed. Countrywide fixed cost per policy is 75%×($18,433k+$24,059k)/(135k+138k+141k)≈$76.98. The variable provision is 7%+6.2%+25%×6.13%+25%×8.09%≈16.76%.','Split only acquisition and general expenses; retain State A commission and taxes as variable.'),
   ('Using premium-based expenses, fixed expense ratio ≈$76.98/$725=10.62%. Loss plus ULAE ratio = 67.2%×1.06=71.232%. Indicated change = (71.232%+10.62%)/(1−16.76%−3.5%)−1 ≈2.64%. The mixed method better reflects per-policy fixed costs when premium size differs from the countrywide mix.','ULAE multiplies loss and ALAE; explain the expense-method choice and solve for a rate change, not a rate level.')],
6:[('Paid age factors are 1.5 at 12–24 and 1.15 at 24–36, with no tail. Percent unpaid is 0%, 1−1/1.15=13.0435%, and 1−1/1.725=42.0290%. Paid BF ultimate: 2015 $1,725,000; 2016 $1,650,000+0.65×$3,100,000×13.0435%≈$1,912,826; 2017 $900,000+0.65×$2,100,000×42.0290%≈$1,473,696.','Use actual earned premium, not on-level premium, in BF at historical accident-year level.'),
   ('On-level 2015–2017 premiums to the 1.08×1.05 current rate level; even annual writings imply historical earned rate indices about 1, 1.01, and 1.07675. On-level premiums are about $2.835m, $3.481m, and $2.212m. Trend each premium at 2% and each BF ultimate at 5% for 4, 3, and 2 years to July 2019: totals are about $9.063m and $5.936m, so R≈65.49%. Solve 1.07=[1.04R+0.06]/[1−0.20−P], giving profit provision P≈10.74%.','Trend the full BF ultimate, not only IBNR; ULAE is 4% of loss and ALAE.')],
7:[('Causality: longer commute may indicate more time in traffic, though route and mode matter. Verifiability: commute time is self-reported and may change with job or traffic. Administrative cost: monitoring changes is burdensome. Social acceptability: a commute surcharge may be difficult to explain or may penalize workers with limited travel choices.','Evaluate four nonstatistical rating-variable criteria with commute-specific reasoning.')],
8:[('Four shortcomings: reported losses are undeveloped, so class ratios may understate ultimate costs differently by class; losses lack trend to the prospective period; the common on-level factor ignores past changes in class and territory relativities, distorting current-rate premiums by class; and class loss ratios are not adjusted for territory-mix differences or credibility before selecting relativities.','The company rates on class and territory only; explain concrete analytical biases rather than requesting unrelated variables.')],
9:[('Harwayne balances other states to State A’s class mix. A has weights 200/425 and 225/425. Reweight B and C class pure premiums: B ≈$48.82; C ≈$52.21. Then weight those state-adjusted costs by their exposure volumes, 550 and 650: complement ≈$50.65 per exposure.','Adjust each other state using State A exposure weights before combining them; do not use loss dollars as weights.'),
   ('The all-state Class 1 pure premium is $31,000/750=$41.33, a simpler complement, but it may be biased for State A because the state loss levels differ.','Evaluate the alternative and its possible bias relative to Harwayne.')],
10:[('Across both policy limits, losses capped at $50,000 total $111,490 thousand: retain the under-$50k loss dollars and cap 1,250 higher claims at $50k. Losses capped at $100,000 total $155,090 thousand. ILF($100k)=$155,090/$111,490=1.3911.','Use both policy groups, including $250k claims censored down to the desired limit, and consistent $000 units.'),
    ('Use both portfolios for ground-up limited average severity: LAS($50k)=111,490/2,550=$43.7216k and LAS($100k)=155,090/2,550=$60.8196k. From the $250k-limit portfolio, probability of a claim above $100k is 200/1,400, and its mean $100k-to-$250k layer is (35,620−200×100)/200=$78.1k. Thus LAS($250k)=60.8196+(200/1,400)×78.1=$71.9768k, giving ILF($250k)=71.9768/43.7216=1.6463.','Use both policy groups below $100k and only the $250k-limit group to estimate the probability and size of the censored upper layer.')],
11:[('Exclude the one-time $1,000k commission. Selected expense ratio = general/earned 7,900/142,000 + acquisition/written 6,950/148,500 + taxes/written 840/148,500 = 10.81%. Permissible loss ratio = 1−10.81%−3.5%=85.69%.','Remove the one-time commission and match each expense category to its appropriate premium base.'),
    ('Overall claim ratio is 129,000/142,000=90.85%. Indicated rate change = 90.85%/85.69%−1 = 6.01%.','Use the selected permissible ratio, not the raw 2014 expense anomaly.'),
    ('Full-credibility claim standard =24,000×0.03=720. Square-root credibility is min(1,sqrt(count/720)): A 1.000, B 0.791, C 0.329.','Convert the exposure standard into a claim standard before applying claim counts.'),
    ('A is the base. Blend each territory loss ratio with the overall ratio using its credibility; indicated relativity for B and C equals current relativity×(blended territory LR / blended A LR). With A credibility 1, B≈1.431 and C≈1.293. The weighted relativity change is [32,000+14,000×1.431/1.22+8,000×1.293/1.35]/54,000≈1.03855. Base rate factor = overall 1.06015/1.03855≈1.02080, an indicated base change of +2.08%.','Apply credibility to territorial loss ratios, normalize to base A, then preserve the indicated overall level through an off-balance adjustment.')],
12:[('Coinsurance requirement is 80%×$500,000=$400,000. At X, loss equals insured value $350,000. Payment before the policy limit is $350,000×$350,000/$400,000=$306,250, so the penalty is $43,750.','At X the penalty peaks just before the insured-value cap changes the payment formula.'),
    ('Y is the $400,000 loss where the insurer’s $350,000 limit is reached and the penalty falls to zero.','Use the coinsurance requirement, not the property value.'),
    ('For losses above Y, the $350,000 policy limit binds. The insured bears the portion of the loss above the limit, even though the chart shows no further coinsurance penalty.','Distinguish the disappearance of the coinsurance penalty from uninsured loss above the limit.')],
13:[('Loss+LAE per exposure = ($7,000+$2,000)/40=$225. Fixed cost per policy = $1,000/(40/1.5)=$37.50. For four exposures, indicated premium = (4×$225+$37.50)/(1−0.15−0.05)=$1,171.88.','Apply the fixed fee once per policy and the variable/profit loading to the whole premium.'),
    ('The pure premium method is attractive for a new product with credible exposure and loss data but little reliable earned premium at current rates.','Explain why an exposure-based calculation is preferable to a loss-ratio denominator in the scenario.')],
14:[('Combining the data increases volume and can stabilize development-factor estimates. The carriers also write the same monoline homeowners coverage and already use the same broad reported-development method, so a pooled analysis is operationally consistent.','Give two merger-specific arguments, including credibility and similar business.'),
    ('Different case-reserve or claim-handling practices can create incompatible reported patterns. Different geographic or catastrophe mixes can also make the pooled factors unrepresentative for either legacy book.','Identify two concrete sources of heterogeneity, not just that the companies differ.')],
15:[('Cumulative paid triangle by accident year, ages 12/24/36: 2015 $1,000/$3,500/$5,600; 2016 $5,000/$8,400; 2017 $500. Include the late-reported 2016 claim only at its 2017 evaluation.','Use cumulative payments at each year-end; do not add ending reserves.'),
    ('Cumulative reported = paid + ending case. The triangle is 2015 $4,000/$5,400/$7,800; 2016 $7,000/$10,600; 2017 $5,000.','The claim 6 reserve falls to zero in 2015; its reported value is then paid $100.'),
    ('Volume-weighted reported factors are 2015+2016 age 12–24: ($5,400+$10,600)/($4,000+$7,000)=1.45455, and 2015 age 24–36: $7,800/$5,400=1.44444. With the 1.1 tail, AY 2017 ultimate = $5,000×1.45455×1.44444×1.1 ≈ $11,556.','Apply both age factors and the supplied 36-to-ultimate tail.')],
16:[('Selected reported factors are 2.0, 1.5, 1.2, and 1.0 (48–60). Latest IBNR ($000): 2013 $0; 2014 $0; 2015 $9,075×(1.2−1)=$1,815; 2016 $5,500×(1.5×1.2−1)=$4,400; 2017 $9,000×(2×1.5×1.2−1)=$23,400.','Show IBNR for every year and subtract the latest reported amount from ultimate.'),
    ('The 2017 large loss arrived unusually early. Multiplying that exceptional $9,000 by an ordinary 12-to-ultimate factor greatly overstates its future development, so an unadjusted chain ladder is inappropriate; separate the large claim or use an expected-loss method.','Explain direction and mechanism of the large-loss distortion.'),
    ('Development works well when claim reporting, settlement, case reserving, and mix remain reasonably stable, so historical age patterns are predictive of current claims.','Name underlying stable operations, not just a stable triangle.')],
17:[('Using a disposal-rate approach, 185 of 200 claims are closed at 12 months, leaving 15. Average historical conditional disposal rates are (14/26+11/20+12/20)/3=56.28% at 12–24 and (6/12+8/9)/2=69.44% at 24–36. Project incremental closures of 8.44, 4.55, and 2.00 at ages 24, 36, and 48. Trend historical incremental severities to 2017 and average by age: about $13,877, $19,648, and $27,318. Unpaid ≈8.44×13,877+4.55×19,648+2.00×27,318=$261,371.','Use future closed counts and incremental paid severities, with 3% severity trend; do not treat severity entries as loss totals.')],
18:[('Volume-weighted reported factors are 12–24 = 9,250/2,600=3.55769, 24–36 = 12,000/5,750=2.08696, and 36–48 = 8,200/5,300=1.54717; include the 1.447 tail. Age-to-ultimate factors for 2014–2017 are 1.447, 2.23875, 4.67218, and 16.62219. Cape Cod expected claims ratio = sum latest reported $19,600 / sum(premium/CDF) ≈52.98%. AY 2017 IBNR = $38,250×52.98%×(1−1/16.62219) ≈ $19,044.','Use reported claims over used-up premium, not developed claims over total premium; calculate IBNR rather than ultimate.')],
19:[('Cumulative reported triangle: 2014 950/3,100/3,800/4,500; 2015 1,100/3,500/4,700; 2016 900/2,800; 2017 700. Volume-weighted 12–24=(3,100+3,500+2,800)/(950+1,100+900)=3.18644; 24–36=(3,800+4,700)/(3,100+3,500)=1.28788; 36–48=4,500/3,800=1.18421. AY 2017 ultimate ≈700×3.18644×1.28788×1.18421×1.03 = $3,504.','Build the cumulative triangle, select age factors, and include the 1.03 tail.'),
    ('An earlier average effective date moves accidents earlier, making 2017 losses older at December 31 than assumed. Applying the old, larger 12-month factor tends to overstate ultimate.','State why maturity shifts and the direction of bias.'),
    ('Align triangles by effective/accident maturity, for example use quarterly evaluations or interpolate percent reported to a comparable effective age before selecting factors.','Adjust the development age rather than using a case-reserve correction.')],
20:[('Reported development is unreliable: average case outstanding per open claim rises sharply in calendar 2017, so past reported factors applied to stronger case values can overstate ultimate. Paid development is also unreliable because closure/disposal rates jump in 2017, accelerating payments relative to history and causing overdevelopment.','Diagnose average case adequacy and claim disposal, not only triangle factor volatility.'),
    ('Ask whether the claims unit changed case-reserve guidelines in 2017 and what changes were made. Separately ask whether staffing, settlement incentives, or closure procedures changed in 2017 and how those changes affected timing.','Ask two distinct operational questions tied to the observed reserve and disposal changes.')],
21:[('The workbook omits ultimate counts, so a numeric result needs an assumption. One examiner-accepted approach uses volume-weighted closed-count factors 108/38, 130/70, and 78/61, with 2014 age-48 count 78 at ultimate. This gives selected ultimate counts about 78, 88.23, 90.24, and 121.49. Latest-diagonal disposal rates are about 14.82%, 42.11%, 78.21%, and 100%; apply these to restate older closed counts. Use each supplied a×exp(b×adjusted closed counts) for older paid entries while retaining the latest diagonal. The adjusted paid age factors are about 2.5634, 1.5521, and 1.2347, giving 2017 ultimate ≈$810×2.5634×1.5521×1.2347=$3,979.','The report explicitly declared this question defective; unadjusted developed ultimate counts received credit, although that assumption is ordinarily questionable.')],
22:[('Salvage/subrogation age factors from simple all-year averages: 12–24 ≈3.7905, 24–36≈1.3936, 36–48=1. Ultimate 2017 recovery ≈$15k×3.7905×1.3936=$79.23k, so recoverable ≈$64.23k.','Subtract the $15k already received; ultimate recovery alone is not recoverable.'),
    ('Develop gross claims with simple all-year age factors 2.2671 and 1.1824, and salvage/subrogation with 3.7905 and 1.3936. The resulting ultimate recovery/gross claim ratios for 2014–2017 are about 19.873%, 20.143%, 20.001%, and 19.706%; their simple all-year average is 19.9305%. Apply it to supplied 2017 gross ultimate $450k: ultimate recovery ≈$89.69k, less $15k already received = $74.69k recoverable.','Match the ratio denominator to the supplied gross ultimate and subtract recoveries received.'),
    ('Salvage and subrogation receipts can be sparse and irregular, making their direct development factors volatile; a ratio to more stable gross ultimate claims can be more credible.','Explain the comparative advantage of the ratio method.')],
23:[('Future opened+closed+pending workloads are 94,040 in calendar 2018, 50,440 in 2019, and 4,580 in 2020. For example, 2018 includes 4,130 AY2015 closures; AY2016 has 16,240 opened, 25,520 closed, 4,640 pending; AY2017 has 16,030 opened, 13,740 closed, 13,740 pending. Divide each workload by 310 and multiply by trended salary: 2018 $60,000×1.03=$61,800, 2019 $63,654, 2020 $65,563.62. Unpaid ULAE ≈$18,747,329+$10,357,122+$968,650=$30,073,100.','Mango–Allen uses opened, closed, and pending workload by calendar year with nominal trended salary.')],
24:[('Volume-weighted paid ALAE factors are about 4.9994, 1.4999, 1.2002, and 1.0501; include the 1.015 tail. The 12-to-ultimate factor is ≈9.5931 and AY 2017 paid development ultimate is $1,250×9.5931 ≈$11,991.','Include the 60-to-ultimate factor; the $1,250 first-year amount is an anomaly.'),
    ('Expected ultimate ALAE = 10%×$50,000=$5,000. Paid BF = paid to date + expected unpaid share = $1,250+$5,000×(1−1/9.5931)≈$5,729.','Use the age-to-ultimate factor to compute the unpaid percentage.'),
    ('Select about $5,729 from BF. The 2017 first-year paid ALAE is unusually high, so direct development magnifies the anomaly to nearly $12,000; BF adds only expected future emergence.','Discuss the specific 2017 anomaly when choosing between estimates.')],
25:[('Possible changes: the insurer’s underlying claim cost ratio is increasing, and claims are being paid more quickly. The rising development estimates relative to expected and BF estimates, without large losses, support investigating these conditions.','Identify two distinct changing conditions and the direction of each.'),
    ('Expected claims based on industry ratios may understate the insurer’s rising cost; faster payment does not affect that prior. Paid BF incorporates higher paid claims but may understate unreported cost from an outdated expected ratio; quicker payments can make it overstate if historical paid factors imply too much still unpaid. Reported Cape Cod responds to rising reported claims through its selected expected ratio, though old periods can dilute the increase; payment speed alone has little direct effect if case reserves fall as payments rise. Reported Benktander also responds to current reported experience but can lag the cost increase through its prior; payment acceleration alone should have little direct effect on reported claims when case estimates are otherwise stable.','For each named technique, discuss both the cost-level change and quicker payments, including when a reported measure is unaffected.')],
26:[('Actual six-month emergence ($000) is 2014 2,000; 2015 6,000; 2016 8,000. Interpolate percent reported, not CDFs. Scenario 1 expected is 1,775, 4,429, 5,000; scenario 2 expected is 2,485, 6,664, 7,231, respectively.','Compare all three years over six months using reciprocals of age-to-ultimate factors.'),
    ('Scenario 2 is closer for 2015 and 2016 and has a smaller total absolute deviation, so it better reflects observed interim emergence.','Justify with actual-versus-expected closeness rather than whether a factor is high or low.')],
}

def main():
    wb=load_workbook(BytesIO(WORKBOOK.read_bytes()), data_only=True)
    grid=wb['Point Grid']
    pages,report_points=report_pages()
    questions=[]
    for n in range(1,27):
        sh=wb[str(n)]
        markers=[(c.row,c.value.strip()[0].lower()) for row in sh for c in row
                 if c.column==2 and isinstance(c.value,str) and re.fullmatch(r'[a-d]\.\s*',c.value)]
        if not markers: markers=[(SINGLE[n],'a')]
        parts=[]
        for i,(start,letter) in enumerate(markers):
            stop=markers[i+1][0] if i+1<len(markers) else sh.max_row+1
            prompt=' '.join(row_text(sh,r) for r in range(start+(n not in SINGLE),stop)).strip()
            prompt=re.sub(r'\s+',' ',prompt)
            points=grid.cell(n+2,i+3).value
            assert isinstance(points,(int,float)) and points>0,(n,letter,points)
            sheet_points=sh.cell(start,1).value
            assert n in SINGLE or sheet_points==points or (n==12 and letter=='c' and sheet_points==0.5 and points==0.25), (n,letter,sheet_points,points)
            answer,insight=ANSWERS[n][i]
            parts.append({'id':letter,'points':points,'prompt':prompt,'solution':answer,'insight':insight})
        assert len(parts)==len(ANSWERS[n]),n
        total=sum(p['points'] for p in parts)
        assert abs(total-grid.cell(n+2,2).value)<1e-9,n
        assert (abs(total-report_points[n])<1e-9 or
                (n==5 and total==2.5 and report_points[n]==2.0)),(n,total,report_points[n])
        q={'id':f'spring-2018-makeup-{n}','number':n,'exam':'Spring 2018 Makeup',
           'chapterIds':CHAPTERS[n],'points':total,'solutionPages':pages[n],
           'sourceBlocks':source_blocks(sh,n,markers[0][0]),'parts':parts}
        if n==5:
            q['notice']='The workbook point grid assigns 2.5 points; the examiner report heading says 2 points. The three part values total 2.5, so this library uses the workbook total.'
        if n==12:
            q['notice']='The Question 12 worksheet labels part c as 0.5 point, but the point grid and examiner report assign 0.25 point; the question total is 1.5 points. This library uses the official 0.25-point value.'
            q['figure']={'src':'assets/exam-graphs/spring-2018-makeup-q12.png',
                         'title':'Coinsurance penalty',
                         'alt':'Original exam chart: coinsurance penalty rises with loss amount to X, then falls to zero at Y and stays at zero.'}
        if n==21:
            q['notice']='The examiner report deemed this question defective because ultimate claim counts were omitted. It was excluded from the official 53.5-point score and pass mark. It remains here for optional practice; any result requires an explicit assumption about ultimate counts.'
            q['excludedFromOfficialScore']=True
        questions.append(q)
    assert sum(q['points'] for q in questions)==56.25
    assert sum(q['points'] for q in questions if not q.get('excludedFromOfficialScore'))==53.5
    OUTPUT.write_text('// Spring 2018 makeup CBT exam; rebuild with scripts/build_spring2018_makeup.py.\n'
                      'window.SPRING_2018_MAKEUP_QUESTIONS = '+json.dumps(questions,ensure_ascii=False,indent=2)+';\n')
    print(f'Wrote {len(questions)} questions, {sum(len(q["parts"]) for q in questions)} parts; 53.5 official points plus 2.75 excluded points.')

if __name__=='__main__': main()
