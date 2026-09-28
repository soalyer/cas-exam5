"""Build the Spring 2014, Fall 2013, and Spring 2013 Exam 5 data.

The official PDFs control wording, table headers, figures and answers. The
companion CBT workbooks provide cells and the scored-part point grids.
"""
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from build_spring2019 import displayed

ROOT = Path(__file__).resolve().parents[1]
CFG = {
    'spring-2014': dict(tag='sp14', workbook='CBT_Exam_5_S.14_v01.xlsx', pdf='admissions_studytools_exam5_sp14-5.pdf', count=23, parts=48, points=59.75, var='SPRING_2014_QUESTIONS', report=[28,31,33,36,40,43,45,48,50,52,56,60,62,65,67,69,71,72,76,78,80,83,88]),
    'fall-2013': dict(tag='f13', workbook='CBT_Exam_5_F.13_v01.xlsx', pdf='admissions_studytools_exam5_f13-5.pdf', count=24, parts=53, points=58.5, var='FALL_2013_QUESTIONS', report=[41,43,46,47,54,55,56,57,58,60,61,63,65,67,68,75,78,80,84,85,86,87,88,89]),
    'spring-2013': dict(tag='sp13', workbook='CBT_Exam_5_S.13_v01.xlsx', pdf='admissions_studytools_exam5_sp13-5.pdf', count=26, parts=61, points=63.5, var='SPRING_2013_QUESTIONS', report=[31,32,34,35,36,39,42,43,44,45,47,49,50,51,52,53,54,55,57,61,63,66,67,68,69,71]),
}

CHAPTERS = {
 'spring-2014':[
  ['ratemaking-5'],['ratemaking-6'],['ratemaking-4'],['ratemaking-6'],['ratemaking-8','ratemaking-12'],
  ['ratemaking-9'],['ratemaking-12'],['ratemaking-9'],['ratemaking-10'],['ratemaking-9'],
  ['ratemaking-11'],['reserving-1','reserving-7'],['reserving-7'],['reserving-9','reserving-15'],
  ['reserving-10'],['reserving-11','reserving-14'],['reserving-15'],['reserving-6','reserving-11'],
  ['reserving-8'],['reserving-14'],['reserving-16'],['reserving-6','reserving-15'],['reserving-6','reserving-15']],
 'fall-2013':[
  ['ratemaking-4'],['ratemaking-5'],['ratemaking-6'],['ratemaking-8'],['ratemaking-6'],
  ['ratemaking-8','reserving-13'],['ratemaking-7'],['ratemaking-8','ratemaking-12'],
  ['ratemaking-9'],['ratemaking-9'],['ratemaking-11'],['ratemaking-13'],['ratemaking-9'],
  ['reserving-6'],['reserving-15'],['reserving-11'],['reserving-9'],['reserving-15'],
  ['reserving-15','reserving-16','reserving-17'],['reserving-10'],['reserving-11','reserving-13'],
  ['reserving-14'],['reserving-6','reserving-15'],['reserving-6','reserving-15']],
 'spring-2013':[
  ['ratemaking-4'],['ratemaking-5'],['ratemaking-8'],['ratemaking-5','reserving-9'],
  ['ratemaking-8'],['ratemaking-16'],['ratemaking-6'],['ratemaking-8'],
  ['ratemaking-7'],['ratemaking-13'],['ratemaking-9','ratemaking-12'],['ratemaking-10'],
  ['ratemaking-11'],['ratemaking-13'],['ratemaking-11'],['reserving-7'],
  ['reserving-12'],['reserving-9','reserving-10'],['reserving-11'],['reserving-6','reserving-13'],
  ['reserving-8'],['reserving-15'],['reserving-13'],['reserving-14'],
  ['reserving-17'],['reserving-15']],
}

# Printed-page corrections to workbook cells, keyed by exam, question, and cell.
FIXES = {}
FIXES['spring-2014',12,'B4']='An insurance company has the following information available for the four different geographic regions within the same line of business:'
FIXES['spring-2014',23,'C9']='Selected Cumulative Percent Reported'
FIXES['spring-2013',16,'C16']='6-12'
FIXES['spring-2013',23,'C6']='Unadjusted Case Outstanding Claims ($000s)'
FIXES['spring-2013',23,'C18']='Unadjusted Cumulative Paid Claims ($000s)'
FIXES['spring-2013',2,'C13']=113800500
FIXES['spring-2013',8,'D25']=21388
FIXES['fall-2013',6,'C27']='Outstanding Case Loss and ALAE Reserves ($000)'
FIXES['spring-2014',13,'B16']='Calculate the ultimate claims for accident year 2013 using the reported claim development technique.'
INSIGHT_FIXES = {
 ('spring-2014',1,'c'):'Longer policy terms change the historical average rate level and therefore the on-level factor; explain the direction without recalculating.',
 ('spring-2014',3,'c'):'Support the date ordering with the policy effective date, observation date, cancellation date, and their constraints.',
 ('spring-2014',10,'a'):'Give distinct business reasons for proposing less than the indicated increase, such as competitive position, retention, or implementation limits.',
 ('spring-2014',10,'b'):'Identify two practical ways to address the pricing shortfall outside the proposed rate change.',
 ('spring-2014',10,'c'):'Use the exposure-weighted rating factors to solve for the base rate that yields an 8% overall increase.',
 ('spring-2014',15,'c'):'Choose between Bornhuetter-Ferguson and Cape Cod under each scenario and explain how the expected claim assumption affects the choice.',
 ('spring-2014',11,'a'):'Calculate the retrospective premium and then apply the policy minimum and maximum; avoid double counting expenses already in the loss conversion factor.',
 ('spring-2014',11,'b'):'The basic premium covers target underwriting profit and expenses outside the loss conversion factor and tax multiplier, plus the cost of the minimum and maximum limits.',
 ('spring-2014',14,'a'):'Use the reported development factor to derive the unreported percentage; do not substitute the paid percentage.',
 ('spring-2014',17,'a'):'Explain how increasing case reserve adequacy affects reported development, paid development, and the resulting unpaid estimates.',
 ('spring-2014',18,'a'):'Recognize the rising case reserve adequacy in the diagnostic data and select reserve techniques consistent with that pattern.',
 ('spring-2014',21,'a'):'Paid ALAE development can be distorted when the pace of legal expense payments differs from the historical pattern.',
 ('spring-2014',21,'b'):'The current paid ALAE-to-paid claim ratio differs from the historical ratio, so a historical ratio can misstate future ALAE.',
 ('spring-2014',21,'c'):'A frequency-severity method can be applied to legal expenses if claim counts and average legal cost are suitable for projection.',
 ('spring-2014',22,'a'):'Interpolate reported development factors to May 31 and compare actual reported emergence with the expected amount.',
 ('spring-2014',22,'b'):'Interpolate paid development factors to May 31 and compare actual paid emergence with the expected amount.',
 ('spring-2014',22,'c'):'Explain a plausible change in claim experience or handling that makes the actual-versus-expected result warrant revising ultimate claims.',
 ('spring-2014',22,'d'):'Explain why the observed variance could be temporary or timing-related, leaving the prior ultimate estimate appropriate.',
 ('spring-2014',23,'a'):'Use expected emergence between 12 and 24 months to solve for the 12-month reported percentage.',
 ('fall-2013',3,'a'):'Calendar-year aggregation aligns premium and loss by accounting period, but mixes policy and accident periods.',
 ('fall-2013',3,'b'):'Explain how calendar/accident-year aggregation matches earned premium to accident-year losses.',
 ('fall-2013',3,'c'):'Policy-year aggregation follows policies written in a period and requires development as those policies earn and losses emerge.',
 ('fall-2013',6,'a'):'Adjust historical case reserves for changes in reserve adequacy before developing losses and calculating projected pure premium.',
 ('fall-2013',7,'a'):'Distinguish how the pure-premium and loss-ratio methods treat fixed versus variable expenses.',
 ('fall-2013',10,'b'):'Adjust for amount-of-insurance mix before comparing territorial pure premiums; an unadjusted comparison confounds the variables.',
 ('fall-2013',8,'a'):'Use total credibility for the catastrophe provision and apply the modeled and non-modeled pieces without double counting.',
 ('fall-2013',11,'a'):'Calculate the traditional increased-limits factor from loss experience at each policy limit relative to the basic limit.',
 ('fall-2013',11,'c'):'Justify the selected increased-limits factor; the GLM result should be challenged if it is inconsistent with neighboring limits.',
 ('fall-2013',11,'b'):'Explain that the GLM can reflect limit-related frequency or behavioral effects that the traditional increased-limits method omits.',
 ('fall-2013',14,'a'):'Allocate paid claims to accident year and development interval, then show incremental rather than cumulative amounts.',
 ('fall-2013',14,'b'):'Allocate paid claims to report year and cumulate payments by development age.',
 ('fall-2013',17,'a'):'Apply the Benktander update for each accident year and sum the resulting unpaid amounts.',
 ('fall-2013',18,'a'):'For excess layers, consider claim count, severity, and attachment effects when choosing a reserve technique.',
 ('fall-2013',22,'a'):'Check whether gross and net reported losses move consistently with both reinsurance treaties and their attachment points.',
 ('fall-2013',22,'b'):'Choose a coherent gross, ceded, and net projection approach, accounting for treaty limits and loss development.',
 ('fall-2013',23,'b'):'Ask three specific claims-department questions that reveal how the operational change affects claim reporting and settlement.',
 ('fall-2013',24,'b'):'Compare actual with expected emergence before changing the selected reported development factors.',
 ('spring-2013',1,'a'):'Treat vehicle counts as the number of covered vehicles, then calculate the fraction of 2011 for which each was insured.',
 ('spring-2013',1,'b'):'Develop written policy-year exposures at each evaluation date rather than treating all vehicles as a full year immediately.',
 ('spring-2013',2,'d'):'Explain how moving renewals to the higher deductible changes the expected loss level and rate indication.',
 ('spring-2013',3,'a'):'Recommend five concrete improvements to the indication and explain why each adjustment is needed.',
 ('spring-2013',5,'a'):'Develop and trend the supplied experience, apply the selected expense provision, and show the final indicated rate change.',
 ('spring-2013',6,'b'):'A claims-made policy generally costs less during a period of rising claim costs because its covered reporting window is shorter.',
 ('spring-2013',6,'c'):'Occurrence coverage remains exposed to future reporting and settlement trends for longer than claims-made coverage.',
 ('spring-2013',6,'d'):'Identify the retroactive-date provision and explain how it prevents overlapping coverage after a switch to claims-made.',
 ('spring-2013',6,'e'):'Explain the claims reporting gap after switching to occurrence coverage and the role of extended reporting coverage.',
 ('spring-2013',8,'a'):'Select and justify frequency and severity trends separately before projecting 2010 ultimate loss and ALAE.',
 ('spring-2013',7,'a'):'Apply both benefit changes to every accident year because the supplied losses are at pre-July 2011 benefit levels.',
 ('spring-2013',10,'a'):'Include expected premium in the denominator when expressing lifetime profit as a percentage of premium.',
 ('spring-2013',10,'b'):'Identify two ways the lifetime-value analysis differs from a standard one-period rate indication.',
 ('spring-2013',11,'c'):'Calculate the credibility standard and square-root weights, blend indicated factors with competitor complements, then normalize to the base class.',
 ('spring-2013',18,'d'):'Compare paid and reported Bornhuetter-Ferguson responsiveness when claim ratios increase; link the conclusion to observed emergence.',
 ('spring-2013',18,'e'):'State one shared expected-loss feature and one distinct feature of Cape Cod and Bornhuetter-Ferguson.',
 ('spring-2013',19,'a'):'Develop frequency and severity separately, select appropriate trend and development assumptions, and subtract reported claims to get IBNR.',
 ('spring-2013',17,'c'):'Describe a setting where case outstanding development is useful; merely saying the insurer has limited data is insufficient.',
 ('spring-2013',21,'a'):'Calculate expected paid emergence for each accident year using the prior valuation and paid development factors.',
 ('spring-2013',21,'b'):'Calculate expected reported emergence for each accident year using the prior valuation and reported development factors.',
 ('spring-2013',21,'c'):'Use the difference between actual and expected paid and reported emergence to identify a plausible claim-process change.',
 ('spring-2013',22,'a'):'For a self-insured client, adjust claims and exposure data directly; insurer premium on-leveling is unavailable.',
 ('spring-2013',23,'a'):'Adjust the case outstanding triangle before developing reported losses; show the revised cumulative reported amounts.',
 ('spring-2013',25,'c'):'Identify a weakness of the classical ULAE method that the Kittel growth adjustment does not correct.',
 ('spring-2013',25,'a'):'Apply the Kittel ratio to the appropriate year-end case and IBNR components, then calculate the total ULAE provision.',
}
SOLUTION_FIXES = {
 ('spring-2014',1,'a'):'The on-level factor for calendar-year 2011 earned premium is current rate level 1.03763 divided by the historical average rate level 1.04309, or 0.99477. The +4.2% change may be included in both numerator and denominator or cancelled consistently.',
 ('spring-2014',1,'b'):'For policies in force on February 1, 2012, account for 25% written on January 1 and the remaining 75% written uniformly. The weighted historical average rate level is 1.03938; current rate level is 1.03763. The on-level factor is 1.03763 / 1.03938 = 0.99832.',
 ('spring-2014',1,'c'):'A two-year policy term delays the earning of prior rate levels. The weight on the older 1.04513 rate level falls and the weight on 1.000 rises, so the historical average rate level falls and the on-level factor increases.',
 ('spring-2014',23,'a'):'Original IBNR is 1,150 − 500 = 650. Let A be the selected cumulative percent reported at 12 months. Expected emergence of 433 from 12 to 24 months satisfies 433 = 650 × (0.80 − A)/(1 − A). Solving gives A = 0.40, or 40%.',
 ('spring-2014',23,'b'):'At December 31, 2014, reported claims are 500 + 433 = 933 and paid claims are 250 + 394 = 644, so case outstanding is 289. At 24 months, selected reported and paid percentages are 0.80 and 0.55. The case outstanding technique gives unpaid = 289 × [1 + (1 − 0.80)/(0.80 − 0.55)] = 289 × 1.8 = $520.20.',
 ('fall-2013',2,'a'):'At current January 2013 rates, premium per exposure is 500 × 1.00 + 55 = $555 for class A and 500 × 0.80 + 55 = $455 for class B. After applying half-year earnings to the six-month policies, 2012 earned exposures are 350 thousand for A and 300 thousand for B. Calendar-year 2012 earned premium at current rate level is 350 × $555 + 300 × $455 = $330,750 thousand, or $330.75 million.',
 ('fall-2013',10,'a'):'A univariate indication assumes the other rating variables are distributed uniformly across territories. Here amount-of-insurance mix differs by territory, so the univariate pure premiums reflect both territory and amount of insurance; using them directly would double count part of the latter effect.',
 ('fall-2013',10,'b'):'Adjust each territory’s exposures for the amount-of-insurance factors. Territory A: 50,000×0.75 + 30,000×1.00 + 20,000×1.50 = 97,500. Territory B: 25,000×0.75 + 75,000×1.00 + 150,000×1.50 = 318,750. Adjusted pure premiums are 60,000,000/97,500 = 615.38 and 300,000,000/318,750 = 941.18. Divide by the base Territory A value: indicated relativity A = 1.000, B ≈ 1.529.',
 ('spring-2013',17,'a'):'For accident year 2012, infer the case outstanding development factor from the industry factors: 1 + [(1.120 − 1) × 1.560/(1.560 − 1.120)] = 1.425. Apply it to $110 of case outstanding: unpaid claims = $110 × 1.425 = $156.75 (in the question’s units).',
 ('spring-2013',11,'a'):'Credit score can raise privacy or social concerns and may lack an obvious causal relationship to loss, although it separates pure premiums. Age is outside the insured’s control, and the insurer’s indications differ markedly from the competitor’s factors. Loss prevention devices can be costly to verify and may be misreported; the classification also puts a home with both devices into the smoke-detector class, obscuring the fire-extinguisher effect.',
 ('spring-2013',11,'b'):'Choose credit score. It separates pure premiums across the three groups (116.67, 128, and 155), has an objective definition, and is comparatively easy to verify and administer. Address applicable legal and privacy concerns before implementation.',
 ('spring-2013',11,'c'):'For Excellent, Good, and Fair credit, pure-premium relativities to the $130 overall average are 0.8975, 0.9846, and 1.1923. Expected claim counts are 150, 250, and 100, giving square-root credibilities 61.24%, 79.06%, and 50.00% against a 400-claim standard. Normalize competitor factors by their exposure-weighted average 1.015, giving 0.8374, 0.9852, and 1.2808. Blend each indicated and complement relativity: 0.8742, 0.9847, and 1.2365. Divide by the Good-class result to set the base class to 1.000: Excellent 0.888, Good 1.000, Fair 1.256.',
}
# Table overrides specify one or more complete tables for groups with split
# workbook headings or side-by-side layouts. Each tuple is
# (first row, last row, data rows, workbook columns, title, printed headers).
TABLES = {key:{} for key in CFG}
IGNORE_GROUPS={'spring-2014':{10:{30}},'fall-2013':{},'spring-2013':{}}
def table(exam,n,first,last,data,cols,title,headers):
    TABLES[exam].setdefault(n,[]).append((first,last,list(data),[ord(c)-64 for c in cols],title,headers))

# Tables with paired sections or a heading split across several workbook rows.
table('spring-2014',5,37,40,range(37,41),'BCDEFG','Reported Loss & ALAE Development Factor Averages',['Selection','12–24','24–36','36–48','48–60','60–Ultimate'])
table('spring-2014',8,8,12,range(10,13),'BCD','Earned Exposures',['Territory','Male','Female'])
table('spring-2014',8,8,12,range(10,13),'FG','Loss and LAE ($) by Territory',['Territory','Loss and LAE ($)'])
table('spring-2014',8,14,17,range(15,18),'BC','Current Territory Relativities',['Territory','Relativity'])
table('spring-2014',8,14,16,range(15,17),'FG','Current Gender Relativities',['Gender','Relativity'])
table('spring-2014',10,19,23,range(20,24),'BCD','Multiplicative Rating Factor 1',['Rating factor','Exposures (000)','Rate differential'])
table('spring-2014',10,19,23,range(20,24),'FGH','Multiplicative Rating Factor 2',['Rating factor','Exposures (000)','Rate differential'])
table('spring-2014',12,8,12,range(9,13),'BCDE','Earned Exposures by Accident Year',['Region','2011','2012','2013'])
table('spring-2014',12,8,12,range(9,13),'BGHI','Ultimate Claim Counts by Accident Year',['Region','2011','2012','2013'])
table('fall-2013',2,12,18,range(14,19),'BCDEF','Rate History',['Effective date','Base rate per exposure','Class A factor','Class B factor','Expense fee'])
table('fall-2013',6,33,35,range(34,36),'BC','Earned Exposures (000)',['Calendar year','Earned exposures (000)'])
table('fall-2013',11,6,10,range(8,11),'BCDEFGH','Policy-Limit Loss Experience',['Size of loss','Claims at $100,000 limit','Losses at $100,000 limit ($)','Claims at $250,000 limit','Losses at $250,000 limit ($)','Claims at $500,000 limit','Losses at $500,000 limit ($)'])
for start,title in [(8,'Paid Losses ($000)'),(15,'Reported Losses ($000)'),(22,'Reported Claim Counts'),(29,'Closed Claim Counts')]:
    for cols,company in [('BCDEF','Company A'),('HIJKL','Company B')]:
        table('fall-2013',15,start,start+4,range(start+1,start+5),cols,f'{company} — {title}',['Accident year','12','24','36','48'])
table('fall-2013',22,7,11,range(9,12),'BCDE','Reported Losses Gross of Reinsurance ($ millions)',['Accident year','12','24','36'])
table('fall-2013',22,14,18,range(16,19),'BCDE','Reported Losses Net of Reinsurance ($ millions)',['Accident year','12','24','36'])
table('spring-2013',4,14,17,range(15,18),'BC','Reported Loss Emergence',['Age','Percentage of loss reported'])
table('spring-2013',4,14,17,range(15,18),'EF','Selected Ultimate Loss Ratios',['Accident year','Ultimate loss ratio'])
table('spring-2013',5,19,21,range(20,22),'BC','Earned Premium ($000)',['Calendar year ending','Earned premium ($000)'])
table('spring-2013',7,16,19,range(17,20),'BC','Ultimate Losses at Pre-July 2011 Benefit Levels ($000)',['Accident year','Ultimate losses ($000)'])
table('spring-2013',8,10,13,range(11,14),'BC','Reported Loss and ALAE as of June 30, 2012',['Accident year','Reported loss and ALAE ($)'])
table('spring-2013',8,19,32,range(21,33),'BCDE','Reported Loss & ALAE',['Calendar year ending','Frequency','Severity ($)','Pure premium ($)'])
table('spring-2013',8,19,24,range(21,25),'GHIJ','Annual Exponential Fits',['# of points','Annual frequency exponential fit','Annual severity exponential fit','Annual pure premium exponential fit'])
table('spring-2013',19,6,10,range(8,11),'BCDE','Reported Claim Counts and Severities as of December 31, 2012',['Accident year','Claim counts','Severity ($)','Payroll ($000)'])
table('spring-2013',19,12,16,range(14,17),'BCD','Reporting Patterns',['As of months','Claim counts','Severities'])
for start,title in [(11,'Cumulative Paid Claims ($000s)'),(17,'Cumulative Reported Claims ($000s)'),(23,'Outstanding Claims ($000s)')]:
    table('spring-2013',20,start,start+4,range(start+2,start+5),'BCDE',title,['Accident year','12 months','24 months','36 months'])
for start,title in [(11,'Cumulative Closed Claim Counts'),(17,'Cumulative Reported Claim Counts'),(23,'Outstanding Claim Counts')]:
    table('spring-2013',20,start,start+4,range(start+2,start+5),'GHIJ',title,['Accident year','12 months','24 months','36 months'])
table('spring-2013',26,6,12,range(9,13),'BCDEFGHIJ','Claims and Projected Ultimate Claims',['Accident year','Reported claims as of December 31, 2012','Paid claims as of December 31, 2012','Reported development ultimate','Paid development ultimate','Reported Bornhuetter–Ferguson ultimate','Paid Bornhuetter–Ferguson ultimate','Claim count and severity ultimate','Disposal rate ultimate'])

def val(exam, n, cell):
    if (exam,n,cell.coordinate) in FIXES:
        return FIXES[exam,n,cell.coordinate]
    if cell.value is None or cell.data_type == 'f':
        return ''
    if isinstance(cell.value,datetime):
        return f'{cell.value:%B} {cell.value.day}, {cell.value.year}'
    if isinstance(cell.value,int) and 2000 <= cell.value <= 2030:
        return str(cell.value)
    result = displayed(cell)
    if isinstance(cell.value,(int,float)) and cell.value>=1000 and cell.value==int(cell.value) and '%' not in cell.number_format:
        result=f'{cell.value:,.0f}'
    return re.sub(r'\s+',' ',result).strip()

def significant(sh, row):
    return [c for c in sh[row][1:16] if c.value is not None and c.data_type!='f']

def automatic_tables(exam,n,sh):
    rows=[r for r in range(4,sh.max_row+1) if len(significant(sh,r))>=2]
    groups=[]
    for r in rows:
        if groups and r==groups[-1][-1]+1:groups[-1].append(r)
        else:groups.append([r])
    specs=[]
    for group in groups:
        if len(group)<2:continue
        first,last=group[0],group[-1]
        if first in IGNORE_GROUPS.get(exam,{}).get(n,set()):
            continue
        if any(first<=spec[0]<=last or first<=spec[1]<=last for spec in TABLES[exam].get(n,[])):
            continue
        cols=sorted({c.column for r in group for c in significant(sh,r)})
        header=[val(exam,n,sh.cell(first,c)) for c in cols]
        # Two-column item/value displays have no header row in the original.
        if len(cols)==2 and not header[1].isalpha() and re.search(r'\d|%|\$',header[1]):
            specs.append((first,last,range(first,last+1),cols,'',['Item','Value']))
        else:
            title=''
            prior=significant(sh,first-1)
            if len(prior)==1 and len(val(exam,n,prior[0]))<95 and not val(exam,n,prior[0]).startswith(('•','o ')):
                title=val(exam,n,prior[0])
            specs.append((first,last,range(first+1,last+1),cols,title,header))
    return sorted(specs+TABLES[exam].get(n,[]),key=lambda s:(s[0],s[3][0]))

def source_blocks(exam,n,sh,markers):
    specs=automatic_tables(exam,n,sh)
    covered={}
    for first,last,rows,cols,*_ in specs:
        for r in range(first,last+1):covered.setdefault(r,set()).update(cols)
    blocks=[]
    for r in range(4,sh.max_row+1):
        for first,last,rows,cols,title,headers in specs:
            if r==first:
                blocks.append({'type':'table','title':title,'headers':headers,'rows':[[val(exam,n,sh.cell(rr,c)) for c in cols] for rr in rows]})
        if r<markers[0][0]:
            line=' '.join(val(exam,n,c) for c in significant(sh,r) if c.column not in covered.get(r,set()))
            repeated_heading = any(line == spec[4] for spec in specs)
            if exam=='fall-2013' and n==15 and re.fullmatch(r'(Paid Losses \(\$000\)|Reported Losses \(\$000\)|Reported Counts|Closed Counts)(?: \1)',line):
                repeated_heading=True
            if line and not repeated_heading:blocks.append({'type':'line','text':line})
    return blocks,covered

def normalize(text):
    text=text.replace('\xa0',' ').replace('\uf0b7','•').replace('‐','-').replace('–','–')
    text=re.sub(r'(?m)^EXAM 5 (?:SPRING 2014) SAMPLE ANSWERS AND EXAMINER.S REPORT\s*$','',text,flags=re.I)
    return re.sub(r'\n{3,}','\n\n',text).strip()

def extract_pdf_pages(pdf,first,last):
    return normalize('\n'.join(pdf.pages[i-1].extract_text() for i in range(first,last)))

def split_parts(text):
    pats=list(re.finditer(r'(?im)^[ \t]*(?:(?:Part[ \t]+([a-g])[ \t]*(?:[.):][ \t]*|(?=\n|$)))|(?:\d+\.[ \t]*)?([a-g])[ \t]*[.):][ \t]*)',text))
    if not pats:return {'a':text}
    return {(m.group(1) or m.group(2)).lower():text[m.end():pats[i+1].start() if i+1<len(pats) else None].strip() for i,m in enumerate(pats)}

def sample_clean(text):
    text=normalize(text)
    text=re.sub(r'(?im)^\s*Exam 5\s*[–-]?\s*Question\s*#?\s*\d+[^\n]*\n','',text)
    text=re.sub(r'(?im)^\s*QUESTION\s*:?\s*\d+\s*$','',text)
    text=re.sub(r'(?im)^\s*TOTAL POINT VALUE:.*$|^\s*LEARNING OBJECTIVE\(S\):.*$','',text)
    return text.strip()

def insight_sentence(text):
    text=re.sub(r'\s+',' ',normalize(text)).strip()
    text=re.sub(r'(?i)^\s*(?:General Commentary|Part\s+[a-g])\s*','',text)
    text=text.lstrip('• ')
    sentences=re.split(r'(?<=[.!?])\s+',text)
    substantive=[s for s in sentences if len(s)>35 and not re.search(r'(?i)^(?:most|many|about|the majority|overall|roughly|approximately|few|candidates generally)\b',s)]
    useful=[s for s in substantive if re.search(r'(?i)expected|needed|should|required|common (?:error|mistake)|full credit',s)]
    chosen=(useful or substantive or sentences or ['Review the method and reasoning in the examiner report.'])[0].strip()
    chosen=re.sub(r'(?i)^(?:Candidates?|The candidate) (?:were|was|are|is) expected to\s+','',chosen)
    chosen=re.sub(r'(?i)^For full credit,\s*','',chosen)
    if len(chosen)>260:chosen=chosen[:260].rsplit(' ',1)[0].rstrip(' ,;:')+'.'
    if chosen and not chosen.endswith(('.','!','?')):chosen+='.'
    return chosen[0].upper()+chosen[1:] if chosen else 'Review the method and reasoning in the examiner report.'

def commentary_sections(exam,pdf):
    first,last=(30,41) if exam=='fall-2013' else (72,len(pdf.pages)+1)
    body=extract_pdf_pages(pdf,first,last)
    result={};cursor=0
    for n in range(1,CFG[exam]['count']+1):
        m=re.search(r'(?m)^\s*'+str(n)+r'\.\s+',body[cursor:])
        if m:
            start=cursor+m.start();result[n]=start;cursor=start+m.end()-m.start()
    return {n:body[start:result.get(n+1,len(body))] for n,start in result.items()}

def reports(exam,n,pdf,single=False):
    cfg=CFG[exam];starts=cfg['report'];first=starts[n-1]
    end=starts[n] if n<cfg['count'] else (len(pdf.pages)+1 if exam!='spring-2013' else 72)
    if exam=='spring-2014':
        raw=extract_pdf_pages(pdf,first,end)
        sample,comment=re.split(r'(?im)^\s*EXAMINER.S REPORT\s*:?',raw,maxsplit=1)
        sample=re.split(r'(?im)^\s*SAMPLE/ACCEPTED ANSWERS\s*:',sample,maxsplit=1)[-1]
        answers=split_parts(sample_clean(sample));comments=split_parts(comment)
    else:
        pages=[]
        for page in range(first,end):
            piece=pdf.pages[page-1].extract_text()
            if exam=='fall-2013' and re.search(r'(?i)\(example\s*[2-9]\)',piece.splitlines()[0] if piece.splitlines() else ''):
                continue
            pages.append(piece)
        sample=sample_clean('\n'.join(pages))
        answers={'a':sample} if single else split_parts(sample)
        comments=split_parts(commentary_sections(exam,pdf).get(n,''))
    for key,s in list(answers.items()):
        s=re.split(r'(?im)^\s*(?:Sample|Accepted Answer)\s*[2-9]\s*:?',s,maxsplit=1)[0]
        s=re.sub(r'(?im)^\s*(?:Sample|Accepted Answer)\s*1\s*:?\s*$','',s).strip()
        s=re.sub(r'^\d+(?:\.\d+)?\s+points?\s*','',s,flags=re.I)
        answers[key]=s[:8000].strip()
    return answers,{k:insight_sentence(v) for k,v in comments.items()},list(range(first,end))

def question_page(exam,n):
    if exam=='fall-2013' and n>=16:return n+4
    if exam=='fall-2013' and n>=5:return n+3
    if exam=='spring-2013' and n>=13:return n+3
    return n+2

def build(exam):
    cfg=CFG[exam];book=load_workbook(ROOT/'past_exams_excel'/cfg['workbook']);grid=book['Point Grid']
    pdf=PdfReader(ROOT/'past_exams'/cfg['pdf']);out=[]
    for n in range(1,cfg['count']+1):
        sh=book[str(n)]
        markers=[(c.row,c.value.strip()[1].lower()) for row in sh for c in row if c.column==1 and isinstance(c.value,str) and re.fullmatch(r'\([a-z]\)',c.value.strip(),re.I)]
        assert markers,(exam,n,'no scored parts')
        blocks,covered=source_blocks(exam,n,sh,markers)
        answers,notes,report_pages=reports(exam,n,pdf,len(markers)==1)
        parts=[]
        for i,(start,letter) in enumerate(markers):
            end=markers[i+1][0] if i+1<len(markers) else sh.max_row+1
            points=grid.cell(n+7,ord(letter)-ord('a')+3).value
            prompt=' '.join(' '.join(val(exam,n,c) for c in significant(sh,r) if c.column not in covered.get(r,set())) for r in range(start,end)).strip()
            prompt=re.sub(r'\s+',' ',prompt)
            if not prompt and len(markers)==1:
                prompt=' '.join(block['text'] for block in blocks if block['type']=='line')
            answer=SOLUTION_FIXES.get((exam,n,letter)) or answers.get(letter) or (answers.get('a') if len(markers)==1 else '')
            insight=INSIGHT_FIXES.get((exam,n,letter)) or notes.get(letter) or notes.get('a') or 'Review the method and reasoning in the examiner report.'
            prompt=prompt.replace('development approacn','development approach')
            assert isinstance(points,(int,float)) and prompt and answer,(exam,n,letter,points,bool(prompt),bool(answer),list(answers))
            parts.append({'id':letter,'points':points,'prompt':prompt,'solution':answer,'insight':insight})
        q={'id':f'{exam}-{n}','number':n,'exam':exam.replace('-',' ').title(),'chapterIds':CHAPTERS[exam][n-1],
           'points':sum(p['points'] for p in parts),'questionPage':question_page(exam,n),
           'solutionPages':report_pages,'sourceBlocks':blocks,'parts':parts}
        if exam=='spring-2014' and n==9:q['figure']={'src':'assets/exam-graphs/spring-2014-q9.png','title':'Credit score GLM diagnostic','alt':'Wind frequency relativity and exposures by credit score, with two-standard-error bounds'}
        if exam=='spring-2013' and n==12:q['figure']={'src':'assets/exam-graphs/spring-2013-q12.png','title':'Burglar alarm and deductible GLM diagnostics','alt':'GLM relativities, standard-error ranges, and policy counts for burglar alarms and deductibles'}
        out.append(q)
    assert len(out)==cfg['count']
    assert sum(len(q['parts']) for q in out)==cfg['parts']
    assert abs(sum(q['points'] for q in out)-cfg['points'])<1e-9
    target=ROOT/'question_data'/(exam.replace('-','')+'_questions.js')
    target.write_text(f'// Official {q["exam"]} PDF; CBT point grid.\nwindow.{cfg["var"]} = '+json.dumps(out,ensure_ascii=False,indent=2)+';\n')
    print(exam,len(out),sum(len(q['parts']) for q in out),sum(q['points'] for q in out),target)

if __name__=='__main__':
    for exam in CFG:build(exam)
