"""Rebuild the regular Spring 2018 sitting from its original CBT workbook.

The examiner's report supplies grading guidance, not sample solutions. Worked
answers below are independently calculated from the workbook and checked
against the report's expectations. The separately supplied makeup sitting is
intentionally excluded.
"""

import json
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from build_spring2019 import displayed


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "past_exams/april_may_sp18-5.xls"
REPORT = ROOT / "past_exams/april_may_sp18-5-examiners_report.pdf"
OUTPUT = ROOT / "question_data/spring2018_questions.js"
# The checked-in Q13 PNG is a high-resolution rendering of the workbook's
# xl/media/image1.emf, cropped to the chart without changing its data.

CHAPTERS = {
    1: ["ratemaking-4"], 2: ["ratemaking-12"], 3: ["ratemaking-6"],
    4: ["ratemaking-12"], 5: ["ratemaking-6"], 6: ["ratemaking-16", "reserving-7"],
    7: ["ratemaking-7"], 8: ["ratemaking-8", "reserving-10"],
    9: ["ratemaking-13"], 10: ["ratemaking-9"], 11: ["ratemaking-13"],
    12: ["ratemaking-9"], 13: ["ratemaking-10"], 14: ["ratemaking-11"],
    15: ["ratemaking-15"], 16: ["reserving-1"], 17: ["reserving-9"],
    18: ["reserving-11"], 19: ["reserving-7", "reserving-9", "reserving-15"],
    20: ["reserving-12"], 21: ["reserving-13"], 22: ["reserving-7", "reserving-14"],
    23: ["reserving-17"], 24: ["reserving-6"], 25: ["reserving-16"],
    26: ["reserving-6", "reserving-13", "reserving-15"],
}

SINGLE = {2: 23, 7: 20, 9: 25, 13: 35, 18: 26, 21: 47}

# Each table is (first covered row, last covered row, data rows, columns,
# headings, title). Wording and units follow the original workbook.
TABLES = {
    2: [(6, 13, range(7, 14), [2, 3], ["Current rate review", "Value"], ""),
        (15, 19, range(16, 20), [2, 3], ["Prior rate review", "Value"], "")],
    4: [(7, 10, range(9, 11), [2, 3, 4, 5, 6, 7, 8],
         ["Coverage", "2015 exposures", "2016 exposures", "2017 exposures", "2015 claims", "2016 claims", "2017 claims"], "")],
    5: [(6, 14, range(9, 15), [2, 3, 4, 5], ["Accident year", "Loss size", "Claim counts", "Ground-up losses ($)"], "Fire"),
        (6, 14, range(9, 15), [2, 3, 6, 7], ["Accident year", "Loss size", "Claim counts", "Ground-up losses ($)"], "Flood"),
        (6, 14, range(9, 15), [2, 3, 8, 9], ["Accident year", "Loss size", "Claim counts", "Ground-up losses ($)"], "Other")],
    6: [(6, 15, range(8, 16), [2, 3, 4, 5, 6, 7, 8],
         ["Accident year", "12", "24", "36", "48", "60", "72"], "Cumulative reported claim counts (months)"),
        (17, 23, range(18, 24), [2, 3], ["Claims-made year", "Step factor"], "")],
    7: [(8, 12, range(9, 13), [2, 3, 4], ["Expense", "Amount ($000)", "% fixed"], ""),
        (14, 18, range(15, 19), [2, 3, 4], ["Measure", "Countrywide", "State"], "")],
    8: [(6, 10, range(8, 11), [2, 3, 4, 5], ["Accident year", "12", "24", "36"], "Cumulative reported loss + ALAE ($000)"),
        (12, 15, range(13, 16), [2, 3], ["Calendar year", "Earned premium ($000)"], ""),
        (17, 25, [17, 18, 19, 20, 21, 22, 23, 25], [3, 2], ["Assumption", "Value"], "")],
    9: [(6, 11, range(8, 12), [2, 3, 4, 5, 6, 7], ["Year", "Age", "Premium ($)", "Loss ($)", "Expense ($)", "Renewal probability"], "21-year-old driver"),
        (13, 18, range(15, 19), [2, 3, 4, 5, 6, 7], ["Year", "Age", "Premium ($)", "Loss ($)", "Expense ($)", "Renewal probability"], "65-year-old driver")],
    11: [(6, 11, range(9, 12), [2, 3, 4, 5, 6, 7],
          ["Risk", "True cost", "A risks", "A rate", "B risks", "B rate"], "")],
    12: [(6, 10, range(8, 11), [2, 3, 4, 5, 6, 7],
          ["Class", "Earned exposures", "Reported loss + ALAE", "Claims", "Current relativity", "True relativity"], "")],
    13: [(9, 12, range(10, 13), [2, 3, 4, 5],
          ["Population density", "Internal GLM", "Competitor", "Industry"], "Rating factors")],
    15: [(6, 10, range(7, 11), [2, 3, 4, 5],
          ["Policy year", "Primary losses ($)", "Excess losses ($)", "Payroll ($)"], ""),
         (12, 26, range(12, 27), [3, 2], ["Assumption", "Value"], "")],
    16: [(6, 16, range(7, 17), [2, 3, 4, 5, 6],
          ["Claim ID", "Accident date", "Transaction date", "Paid ($)", "Ending case ($)"], "")],
    17: [(6, 12, range(8, 13), [2, 3, 4, 5, 6, 7],
          ["Accident year", "12", "24", "36", "48", "60"], "Cumulative reported claims ($000)"),
         (14, 19, range(15, 20), [2, 3], ["Calendar year", "Earned premium ($000)"], "")],
    18: [(6, 11, range(8, 12), [2, 3, 4, 5, 6], ["Accident year", "12", "24", "36", "48"], "Cumulative reported claims ($000)"),
         (6, 11, range(8, 12), [8, 9, 10, 11, 12], ["Accident year", "12", "24", "36", "48"], "Cumulative reported counts"),
         (14, 18, range(16, 19), [2, 3, 4, 5], ["Accident year", "12–24", "24–36", "36–48"], "Claim development factors"),
         (14, 18, range(16, 19), [8, 9, 10, 11], ["Accident year", "12–24", "24–36", "36–48"], "Count development factors")],
    21: [(6, 12, range(8, 13), [2, 3, 4, 5, 6], ["Accident year", "12", "24", "36", "48"], "Cumulative paid loss + ALAE ($)"),
         (6, 12, range(8, 13), [8, 9, 10, 11, 12], ["Accident year", "12", "24", "36", "48"], "Average case outstanding ($)"),
         (14, 20, range(16, 21), [2, 3, 4, 5, 6], ["Accident year", "12", "24", "36", "48"], "Open claim counts"),
         (14, 20, range(16, 21), [8, 9, 10, 11, 12], ["Accident year", "12", "24", "36", "48"], "Cumulative closed counts"),
         (22, 29, range(25, 30), [2, 3, 4, 5, 6], ["Accident year", "12", "24", "36", "48"], "Cumulative reported counts"),
         (22, 29, range(25, 30), [8, 9], ["Accident year", "Ultimate counts"], ""),
         (31, 38, range(34, 39), [2, 3, 4, 5, 6], ["Accident year", "12", "24", "36", "48"], "Berquist-Sherman adjusted paid claims ($)")],
    22: [(6, 12, range(9, 13), [2, 3, 4, 5, 6], ["Accident year", "12", "24", "36", "48"], "Cumulative gross reported claims ($000)"),
         (17, 23, range(20, 24), [2, 3], ["Accident year", "Stop-loss attachment ($000)"], "")],
    23: [(6, 11, range(8, 12), [2, 3, 4, 5], ["Calendar year", "Paid ULAE ($)", "Paid claims ($)", "Incurred claims ($)"], "")],
    24: [(6, 13, range(10, 14), [2, 3, 4, 5],
          ["Accident year", "Ultimate at Jun 30", "Reported at Jun 30", "Reported at Dec 31"], "Claims ($000)"),
         (15, 17, [17], list(range(2, 10)),
          ["6–12", "12–18", "18–24", "24–30", "30–36", "36–42", "42–48", "48–ultimate"], "Selected reported age-to-age factors")],
    25: [(6, 11, range(8, 12), [2, 3, 4, 5, 6], ["Accident year", "12", "24", "36", "48"], "Cumulative paid ALAE ($000)"),
         (6, 11, range(8, 12), [8, 9, 10, 11, 12], ["Accident year", "12", "24", "36", "48"], "Cumulative paid claims ($000)"),
         (13, 18, range(15, 19), [2, 3, 4, 5, 6], ["Accident year", "12", "24", "36", "48"], "Paid ALAE / paid claims"),
         (20, 28, range(25, 29), [2, 3, 4, 5], ["Accident year", "Reported development ultimate", "Selected ultimate ($000)", "Selected ultimate count"], "")],
    26: [(7, 14, range(8, 15), [2, 3], ["Projection technique", "Ultimate claims estimate"], "")],
}


def cell_text(cell):
    if isinstance(cell.value, datetime):
        return f"{cell.value:%B} {cell.value.day}, {cell.value.year}"
    value = displayed(cell)
    if isinstance(cell.value, (int, float)) and "%" not in cell.number_format:
        try:
            shown = float(value.replace(",", "").replace("$", ""))
        except ValueError:
            shown = cell.value
        if abs(shown - cell.value) > 0.00005:
            decimals = 2 if abs(cell.value) >= 1000 else 4
            precise = f"{cell.value:,.{decimals}f}".rstrip("0").rstrip(".")
            value = f"${precise}" if "$" in cell.number_format else precise
    return re.sub(r"\s+", " ", value).strip()


def row_text(sheet, row):
    return " ".join(cell_text(cell) for cell in sheet[row][1:18] if cell.value is not None and cell.data_type != "f").strip()


def source_blocks(sheet, number, stop):
    specs = TABLES.get(number, [])
    covered = {row for first, last, *_ in specs for row in range(first, last + 1)}
    blocks = []
    previous_line_row = None
    for row in range(4, stop):
        for first, _, data_rows, columns, headings, title in specs:
            if first != row:
                continue
            rows = [[cell_text(sheet.cell(r, c)) for c in columns] for r in data_rows]
            blocks.append({"type": "table", "title": title, "headers": headings, "rows": rows})
        if row in covered:
            previous_line_row = None
            continue
        line = row_text(sheet, row)
        if line:
            if (previous_line_row == row - 1 and blocks[-1]["type"] == "line"
                    and not blocks[-1]["text"].endswith((".", ":", ";", "?", "!"))
                    and not line.startswith(("•", "i.", "ii.", "iii."))):
                blocks[-1]["text"] += " " + line
            else:
                blocks.append({"type": "line", "text": line})
            previous_line_row = row
        else:
            previous_line_row = None
    return blocks


def report_pages():
    reader = PdfReader(REPORT)
    starts = []
    reported_points = {}
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        match = re.search(r"(?im)^QUESTION\s+(\d+)\s*$", text)
        if match:
            number = int(match.group(1))
            starts.append((number, page_number))
            score = re.search(r"TOTAL POINT VALUE:\s*([\d.]+)", text)
            assert score, (number, "missing report point value")
            reported_points[number] = float(score.group(1))
    assert [number for number, _ in starts] == list(range(1, 27)), starts
    pages = {
        number: list(range(page, starts[i + 1][1] if i + 1 < len(starts) else len(reader.pages) + 1))
        for i, (number, page) in enumerate(starts)
    }
    return pages, reported_points


# One accepted answer and one report-grounded grading insight per scored part.
ANSWERS = {
    1: [
        ("Written exposures in 2017 are 4 at issue, less 2 occupants for the remaining half-year: 4 - 2 × 0.5 = 3 occupant-years.",
         "Reflect the October endorsement; counting only the original four occupants overstates written exposure."),
        ("Earned exposures in 2017 are 4 × 6/12 from April through September plus 2 × 3/12 from October through December = 2.5 occupant-years.",
         "Earn only the portion of each exposure that falls in calendar year 2017."),
        ("At September 30, 2017, the October change has not occurred. Policy year 2017 written exposure is 4 occupant-years.",
         "Use the policy-year valuation date rather than the later year-end information."),
        ("Proportionality: occupant count may track household liability exposure, but home-property damage may depend more on the building. Practicality: count is definable but changes must be collected and verified. Historical precedence: switching from house-years would disrupt systems, create premium changes, and reduce comparability with prior and industry data.",
         "Evaluate all three criteria with a reason; naming them alone is insufficient."),
    ],
    2: [("Bühlmann K = EPV/VHM = 7.5/0.45 = 16.667 and Z = 100/(100 + K) = 0.8571. Trended present-rate complement = $600 × (1.20/1.09) × 1.03² = $700.78. Credibility-weighted indicated premium = 0.8571 × $750 + 0.1429 × $700.78 = $742.97.",
         "Account for the unimplemented portion of the prior indication and the two-year trend period." )],
    3: [
        ("Competitor filings offer more years of state experience but can reflect a different risk and coverage mix. Internal data matches the company's own business but one year in the new state is thin and volatile; data from its other states may have different loss behavior.",
         "Discuss a concrete advantage and disadvantage for both sources, including credibility and comparability."),
        ("The overlap fallacy says loss development and trend duplicate one another. They do not: development estimates later emergence on claims from an accident period, while trend moves the expected cost level to a future period. Both adjustments can be needed.",
         "Explicitly reject overlap and describe what each adjustment measures."),
    ],
    4: [
        ("Select pooled frequency by coverage. Bodily injury: (1,020 + 1,100 + 950)/(20,000 + 24,000 + 22,000) = 0.04652; full-credibility exposure standard = 1,082/0.04652 ≈ 23,261. Collision: 5,775/58,000 = 0.09957; standard ≈ 10,867 exposures.",
         "Calculate separate expected frequencies for the two coverages before converting the claim standard."),
        ("Exposures may be more stable and available earlier than claim counts, so an exposure standard can be used to assess credibility before claims have fully emerged.",
         "Simply observing that exposures outnumber claims is not an advantage; their full-credibility standards differ too."),
        ("Advantage: classical credibility has a simple, transparent full-credibility standard. Disadvantage: its Poisson-frequency and constant-severity assumptions can miss overdispersion or heterogeneous losses.",
         "Give properties specific to classical credibility, not generic benefits of averaging."),
        ("Use a state-wide indication as the complement. It has more volume and may be stable, but its mix and loss level may differ from the specific coverage or segment being priced, so adjust for those differences.",
         "A proposed complement needs both a credibility rationale and a comparability limitation."),
    ],
    5: [
        ("For a non-catastrophe selection, exclude the 2014 flood and the $1 million 2016 forest-fire claim. Losses limited to $500,000 total $336.40125 million; excess above $500,000 totals $7.2115 million. Excess loss factor = 1 + 7.2115/336.40125 = 1.02144. Including catastrophes with a consistent treatment was also accepted.",
         "Split each ground-up claim above $500,000 into a $500,000 limited portion and an excess portion; report a factor, not an excess ratio."),
        ("Large claims and catastrophes make annual results volatile. Cap experience at the threshold and add a longer-term expected excess load to smooth rate indications.",
         "The factor addresses volatility in ratemaking, not missing or censored data."),
    ],
    6: [
        ("By December 31, 2014, claims-made coverage includes accident years 2010–2014 reported by then: 85 + 82 + 85 + 44 + 15 = 311 claims.",
         "Read the cumulative triangle at the 2014 calendar-year diagonal and respect the 2010 retroactive date."),
        ("Occurrence coverage beginning January 1, 2015 covers accident years 2015–2017 regardless of report date. Reported counts to date are 77 + 65 + 30 = 172.",
         "Use occurrence year, not claims-made reporting year, for the new coverage."),
        ("The gap is claims occurring in 2010–2014 but reported after the claims-made policy ends. Develop latest 2017 counts using the step factors: 89 + 85 + 112 + 70/0.95 + 83/0.90 ≈ 451.91 ultimate old-year claims. Subtract 311 covered by 2014: about 141 claims fall in the gap.",
         "Develop the latest cumulative count for immature years, then subtract old-year claims reported before the switch."),
        ("Buy an extended reporting period endorsement (tail coverage) on the claims-made policy to cover later reports of events during its covered years.",
         "Name the specific coverage mechanism; a vague 'gap endorsement' is insufficient."),
    ],
    7: [("Fixed cost per exposure is 1,000 × [0.75(3,648)/84,115 + 0.75(4,368)/87,476 + 0.25(315)/20,217] = $73.87. The variable expense ratio is 0.25(3,648)/80,948 + 0.25(4,368)/82,583 + 0.75(315)/18,498 + 1,868/18,498 = 13.82%. Fee = $73.87/(1 - 0.1382 - 0.05) ≈ $91.00 per exposure.",
         "Use the correct written/earned and state/countrywide base for each expense, then gross up for variable expense and profit." )],
    8: [
        ("Select reported 12–24 = 2,452.5/2,100 = 1.16786, 24–36 = 1,250/1,167.5 = 1.07066, and 36–ultimate = 1.067; age-to-ultimate factors are 1.067, 1.14240, and 1.33416 for 2015–2017. On-level earned premiums are approximately 1,976, 2,157.61, and 2,110.14 ($000), using 0%, 12.5%, and 87.5% at the new rate. Trend losses and premium to a common 2017 level; trended reported losses divided by used-up trended premium gives a Cape Cod ratio near 66.67%. De-trend that ratio for each year and add expected unreported losses to actual reported losses. One consistent selection gives ultimates about 1,334, 1,466, and 1,402 ($000).",
         "On-level premium and trend both sides consistently; use reported rather than developed losses in the Cape Cod ratio, and do not add trended reported losses to historical ultimates."),
        ("Trend the three historical ultimates and on-level premiums to the January 2019 policy-effective cohort's average earned date, January 2020. With 2% loss and 3% premium trend, trended totals are about 4,504 and 6,922 ($000). Indication = [1.07 × 4,504 + 0.10 × 6,922]/[(1 - 0.30 - 0.05) × 6,922] - 1 ≈ 22.5%. Different internally consistent Cape Cod and trend selections may vary.",
         "The indication must use the ultimates from part a, future trend periods, ULAE, and fixed/variable expense and profit provisions."),
    ],
    9: [("Use cumulative renewal and discount all cash flows. For age 21, weights are 1, 0.75/1.05, and 0.75²/1.05²; discounted premium ≈ $2,676.73 and profit ≈ $62.96, or 2.35%. For age 65, weights are 1, 0.95/1.05, and 0.95²/1.05²; discounted premium ≈ $2,451.02 and profit ≈ $109.29, or 4.46%. The 65-year-old has the larger percentage return.",
         "Apply renewal probability cumulatively in year three and divide discounted profit by discounted premium." )],
    10: [
        ("Objectivity/verifiability: vehicle color is observable, but repainting, wraps, and inconsistent color categories require clear rules and verification. Administrative cost: the insurer must collect color at application, update systems and policy changes, and audit data, which may outweigh its value.",
         "Give two distinct operational criteria and support each with a color-specific explanation."),
        ("Controllability: drivers can change color, but repainting may be costly, so a surcharge is not easily avoided. Social acceptability: a BI premium difference by color may seem arbitrary to insureds without an understandable link to loss risk, reducing perceived fairness.",
         "Explain the impact on insureds or society; legal and statistical criteria alone do not answer the social question."),
    ],
    11: [
        ("Among the 20% who shop, 1,000 high-risk customers move from A ($160) to B ($150), and 1,000 low-risk customers move from B ($150) to A ($140). A now has 4,000 high and 6,000 low risks: 4,000(160-170) + 6,000(140-130) = +$20,000. B has 6,000 high and 4,000 low: 6,000(150-170) + 4,000(150-130) = -$40,000.",
         "Move shoppers by risk class and compare each company's charged rate with true expected cost."),
        ("B could introduce risk-based rating that charges high risks more and low risks less as one coordinated correction. It could also tighten underwriting for high-risk submissions or use targeted risk-control measures to lower their expected loss.",
         "Give two distinct ways to improve B's loss position; expense cuts and investment income do not fit the stated assumptions."),
    ],
    12: [
        ("Pure premiums for A, B, C are $56.58, $100.97, and $47.95 per exposure, giving raw relativities 1.000, 1.785, and 0.847. Square-root credibility is 0.921, 0.635, and 0.765. Blend raw and current relativities to 1.000, 1.534, and 1.071, then apply a revenue-neutral off-balance factor of about 1.1201. Class rate changes are approximately +12.0% for A, +56.2% for B, and -33.3% for C.",
         "Use losses per earned exposure, normalize to A, blend with current relativities, and off-balance to preserve total premium."),
        ("The classes may have different mixes of other rating variables. A univariate pure premium comparison then attributes those other variables' costs to this class variable, so indicated factors can differ from true factors.",
         "Identify distributional bias across other rating variables, rather than competitive or regulatory selection."),
        ("For each class, divide losses by exposures adjusted for the exposure-weighted average relativity of the other rating variables before deriving this variable's relativities.",
         "Describe an adjustment to the univariate pure premium method itself."),
    ],
    13: [("Select a low-density factor near 0.9, relying more on the competitor (0.9) and industry (0.4) than the internal GLM (1.5): low-density exposure is sparse and the GLM error band is very wide. Select a high-density factor near 1.3: the internal GLM is more credible there, its narrower band supports a factor above 1.0, and it lies between competitor 1.1 and industry 1.6. Medium remains 1.0.",
          "Support both selections with GLM diagnostics and both external references, especially low-density credibility." )],
    14: [
        ("Required insurance is $250,000 × 90% = $225,000, so the coinsurance fraction is $200,000/$225,000 = 8/9. For a $50,000 loss, indemnity is $44,444.44 and penalty is $5,555.56. For $220,000, indemnity is $195,555.56 and penalty relative to the $200,000 policy limit is $4,444.44. For $250,000, indemnity is capped at $200,000 and the coinsurance penalty is $0; the remaining $50,000 is uninsured above the limit.",
         "Cap indemnity at the insured value and separate coinsurance penalty from loss above the policy limit."),
        ("An insured can face a coinsurance penalty on a partial loss and can lack funds to rebuild after a large loss. For the insurer, underinsurance can make premiums inadequate because smaller losses remain covered while the insured value, on which premium is based, is too low relative to exposure.",
         "Tie both problems to the mechanics of partial losses, policy limits, or coinsurance."),
    ],
    15: [
        ("Expected losses = $6,305,000/$100 × 3.25 = $204,912.50; expected primary and excess are $40,982.50 and $163,930. With actual primary $30,000, actual excess $260,000, ballast $40,000, and weight 0.30, mod = (30,000 + 40,000 + 0.30×260,000)/(40,982.50 + 40,000 + 0.30×163,930) ≈ 1.137.",
         "Use payroll and the expected loss rate to split expected primary and excess; do not use the expected loss ratio here."),
        ("Basic premium = $800,000 × [0.25 - 0.70(1.10 - 1) + (0.40 - 0.05)×0.70×1.10] = $359,600. Converted limited losses = $200,000×1.10 = $220,000. Preliminary retrospective premium = ($359,600 + $220,000)×1.05 = $608,580; this is between the $560,000 minimum and $1,040,000 maximum, so final premium is $608,580.",
         "Use the supplied limited losses, calculate basic premium with the net insurance charge and LCF adjustment, then check the minimum and maximum."),
    ],
    16: [
        ("Accident year 2015 cumulative paid claims at 12, 24, and 36 months (December 31 of 2015, 2016, and 2017) are $600, $1,350, and $2,025. Exclude claim A, whose accident occurred in 2014.",
         "Development ages run from the accident year's December 31, not each claim's accident date."),
        ("Cumulative reported = cumulative paid plus ending case outstanding. At 12 months: B $1,000 + C $750 = $1,750. At 24 months: B $1,200 + C $750 + D $325 = $2,275. At 36 months: B $1,200 + C $650 + D $625 = $2,475.",
         "Carry forward claim B's $200 case reserve at 36 months even without a 2017 transaction."),
    ],
    17: [
        ("Select volume-weighted reported factors about 1.2511 (12–24), 1.1025 (24–36), 1.0498 (36–48), and 1.000 (48–ultimate). BF ultimate = reported + 70%×earned premium×(1 - 1/CDF). The 2013–2017 ultimates ($000) are about 7,000, 8,400, 10,998, 13,413, and 16,564; dividing by earned premium gives ultimate claim ratios 70.0%, 70.0%, 73.3%, 74.5%, and 75.3%.",
         "Report ultimate claim ratios, not only ultimate dollars, and divide by earned premium."),
        ("The most recent ultimate ratios rise above the assumed 70% and reach about 75%. If the underlying claims ratio is deteriorating, BF's fixed 70% expected ratio understates the immature years' expected unreported claims; update that prior assumption.",
         "Base the critique on the observed rising ratios, not an unsupported claim of case-reserve changes."),
    ],
    18: [("Develop counts separately from severity. Volume-weighted count factors are about 1.0393, 1.00385, and 1.000, giving 2017 ultimate count ≈ 1,132. Develop historical reported severity (claims divided by counts) with factors about 1.1506, 1.0843, and 1.0423. Trend the developed 2014–2016 severities at 5% to 2017 and select about $4.63 thousand, then apply the 30% tort-reform reduction. Ultimate 2017 claims ≈ 1,132×4.63×0.70 = $3.67 million.",
          "The severity development pattern must be calculated independently of reported-claim factors, then trended and reduced for tort reform." )],
    19: [
        ("Paid development is distorted upward by an early shock payment, since it multiplies unusually high paid-to-date. Reported BF also incorporates the shock through reported-to-date, but does not multiply it by the full reported development factor; review whether its expected-loss prior includes the shock.",
         "Discuss the specified paid development and reported BF techniques, not paid BF."),
        ("Paid development uses claim payments and can still be applied if rate-change history is poor. Reported BF needs a credible expected-claims prior; without reliable on-level premium, obtain an independent prior or it may be unsuitable.",
         "Unreliable rate changes affect the BF prior, not the paid development pattern directly."),
        ("Earlier settlements accelerate paid emergence, so unadjusted historical paid factors applied to the newer fast-paying data tend to overstate ultimate. Reported BF is less affected if reported claims and its prior are stable.",
         "Distinguish changed payment timing from reported development."),
        ("Unexpectedly high severity makes historic paid development factors applied to current paid amounts less representative and may understate ultimate. Reported BF also understates unreported claims if its expected-loss prior uses too-low severity trend; revisit that prior.",
         "Explain how higher severity affects current experience and the expected-claims assumption."),
    ],
    20: [
        ("Stronger case reserves raise reported claims and case outstanding without an equivalent change in ultimate cost. Historical reported development factors applied to the stronger diagonal overstate ultimate; case-outstanding development likewise overstates because the current outstanding balance is inflated relative to its historical pattern.",
         "Both unadjusted techniques are biased high; identify how the stronger case reserve enters each."),
        ("A paid claims development technique is unaffected directly because paid claims do not include case outstanding.",
         "Specify paid, not an ambiguous development or BF technique."),
    ],
    21: [("Use the 2016 calendar-year diagonal as the reference for closure rates: about 80.03%, 84.87%, 89.79%, and 96.97% at ages 12, 24, 36, and 48. Restate open counts from each year's reported counts using these rates, and restate average case amounts from the 2016 diagonal with 12% severity trend. Add adjusted paid claims to adjusted open count × adjusted average case. This yields an adjusted 2017 12-month reported amount near $297,920. Volume-weighted adjusted reported age factors are about 2.196, 1.121, and 1.056. Assuming no reported development after 48 months gives 2017 ultimate loss + ALAE ≈ $774,321; a justified other tail assumption changes the result.",
          "Adjust both settlement rate and case adequacy using the 2016 diagonal. The given 1.2 tail is for adjusted paid claims, not automatically for adjusted reported claims." )],
    22: [
        ("Select gross reported age factors 2.0, 1.5, 1.2, then apply the 1.1 tail. Gross ultimates ($000) for 2014–2017 are 792, 950.4, 1,188, and 1,386. After the 20% quota share and the stop-loss cap, insurer net ultimates are 633.6, 750, 900, and 1,000 ($000), respectively.",
         "Apply quota share before aggregate stop loss and use cumulative age-to-ultimate factors including the tail."),
        ("For 2017, gross IBNR is 1,386 - 350 = 1,036 ($000). Current net reported is 0.8×350 = 280, and net ultimate is 1,000, so net IBNR is 720. Ceded IBNR = 1,036 - 720 = 316 ($000).",
         "Subtract net IBNR from gross IBNR; ceded ultimate alone is not ceded IBNR."),
        ("Stop loss usually makes the insurer's net tail factor smaller than gross because late development above the attachment is ceded, though it need not be 1.0. A constant quota share scales claims at every maturity equally, so net and gross tail factors are equal.",
         "State the relationship from the insurer's net perspective for each treaty."),
    ],
    23: [
        ("Paid ULAE/paid claims ratios are 40%, 20%, 15%, and 15% for 2014–2017. Treat 2014 as an outlier and select 15%. IBNYR is 40%×$100,000 = $40,000; IBNER is $60,000. Classical unpaid ULAE = 15%×[$40,000 + 50%×($60,000 + $150,000 case)] = $21,750.",
         "Use paid claims as the ratio denominator, justify excluding the 2014 outlier, and apply half the ratio to case plus IBNER."),
        ("Kittel uses paid ULAE divided by the average of paid and incurred claims. The ratios for 2015–2017 are 13.33%, 12%, and 12%; select 12%. With the same IBNYR, IBNER, and case reserves, unpaid ULAE = 12%×[$40,000 + 50%×($60,000 + $150,000)] = $17,400.",
         "The Kittel denominator is the average of paid and incurred claims, not either one alone."),
        ("A major shift in claim settlement or case-reserve practices can change paid ULAE relative to paid claims and distort the classical ratio. Kittel's average paid-and-incurred denominator dampens that timing distortion by recognizing both payments and reported amounts.",
         "Identify a specific changing ULAE/claims relationship and explain the denominator's corrective role."),
    ],
    24: [
        ("Actual June-to-December reported emergence ($000) for 2014–2017 is -38, 548, 882, and 868. Use the selected ultimate and semiannual factors to calculate expected emergence = ultimate×(1/CDF after - 1/CDF before): approximately 93, 550, 863, and 603. Actual minus expected is about -131, -2, +19, and +265 ($000), respectively.",
         "Compare incremental emergence over the six months; the factors are semiannual, not annual."),
        ("A shift toward longer-tailed claims can delay reports relative to the selected pattern. Stronger earlier case estimates followed by reductions or faster closures can also lower later reported emergence.",
         "Give two causal explanations with the direction of the deviation."),
        ("A shift toward shorter-tailed claims can accelerate reporting. Unexpected large claims or case-reserve strengthening can raise reported claims beyond the expected emergence.",
         "Explain why the circumstances increase reported, not merely paid, claims."),
    ],
    25: [
        ("Paid ALAE age factors are 1,920/320 = 6, 3,840/1,920 = 2, and 4,224/3,840 = 1.1, with no tail. Ultimate 2017 ALAE = 340×6×2×1.1 = $4,488 thousand.",
         "Develop the ALAE triangle itself, not paid claims or an ALAE-to-claims ratio."),
        ("Selected additive increases in the paid ALAE/paid claims ratio are +0.04 (12–24), +0.02 (24–36), and 0 (36–48). Ultimate ratios for 2014–2017 are 0.10, 0.10, 0.10, and 0.14. Apply them to selected ultimate claims of 42,240, 36,960, 43,560, and 66,440 ($000): ultimate ALAE = 4,224, 3,696, 4,356, and 9,301.6 ($000).",
         "Use additive ratio development for every accident year and multiply by selected ultimate claims."),
        ("Select $4,488 thousand from direct paid ALAE development for 2017. The ratio method produces $9,301.6 thousand largely because the selected 2017 ultimate claims value is much higher than the reported-development estimate; its sensitivity to that claim selection weakens the result. Other justified selections are possible.",
         "Explain a strength of the selected method or a data-specific weakness of the other; the report did not require one unique selection."),
    ],
    26: [
        ("Increasing case-reserve adequacy raises reported claims and inflates unadjusted reported development ($6,500) and reported BF ($6,100). Reported Berquist-Sherman adjusts the case-reserve change and returns about $4,900, in line with the paid-based estimates of $4,800–$4,850.",
         "The paid estimates' agreement supports a case-reserve change rather than a payment-pattern change."),
        ("First, inspect average case outstanding by maturity across calendar-year diagonals; an abrupt increase points to stronger case reserves. Second, inspect reported-to-paid claim ratios or the case-outstanding share of reported claims by maturity; a jump while paid patterns stay stable supports the same diagnosis.",
         "Describe how each diagnostic would reveal increasing case adequacy, not merely name it."),
        ("Compare cumulative closed claim counts to reported claim counts by development age (claim-disposal rates). If disposal rates are stable, the paid Berquist-Sherman settlement-rate adjustment changes little, explaining why paid development and paid Berquist-Sherman estimates are close.",
         "Use closed/reported for disposal rate and connect stable rates to the similar paid estimates."),
    ],
}


def main():
    workbook = load_workbook(BytesIO(WORKBOOK.read_bytes()), data_only=True)
    point_grid = workbook["Point Grid"]
    pages, reported_points = report_pages()
    questions = []
    for number in range(1, 27):
        sheet = workbook[str(number)]
        markers = [(cell.row, cell.value.strip()[0].lower()) for row in sheet for cell in row
                   if cell.column == 2 and isinstance(cell.value, str) and re.fullmatch(r"[a-d]\.\s*", cell.value)]
        if not markers:
            markers = [(SINGLE[number], "a")]
        parts = []
        for index, (start, letter) in enumerate(markers):
            stop = markers[index + 1][0] if index + 1 < len(markers) else sheet.max_row + 1
            prompt = " ".join(row_text(sheet, row) for row in range(start + (number not in SINGLE), stop)).strip()
            prompt = re.sub(r"\s+", " ", prompt)
            points = point_grid.cell(number + 2, index + 3).value
            assert isinstance(points, (int, float)) and points > 0, (number, letter, points)
            solution, insight = ANSWERS[number][index]
            parts.append({"id": letter, "points": points, "prompt": prompt, "solution": solution, "insight": insight})
        assert len(parts) == len(ANSWERS[number]), number
        total = sum(part["points"] for part in parts)
        assert abs(total - point_grid.cell(number + 2, 2).value) < 1e-9, number
        assert abs(total - reported_points[number]) < 1e-9, (number, total, reported_points[number])
        question = {
            "id": f"spring-2018-{number}", "number": number, "exam": "Spring 2018",
            "chapterIds": CHAPTERS[number], "points": total,
            "solutionPages": pages[number], "sourceBlocks": source_blocks(sheet, number, markers[0][0]),
            "parts": parts,
        }
        if number == 13:
            question["figure"] = {
                "src": "assets/exam-graphs/spring-2018-q13.png",
                "title": "Population density relativities - internal GLM results",
                "alt": "Internal GLM factors for low, medium and high population density with upper and lower standard-error bounds and exposure bars. Low density has far fewer exposures and a much wider error band.",
            }
        questions.append(question)
    assert sum(question["points"] for question in questions) == 55.5
    OUTPUT.write_text("// Regular Spring 2018 CBT exam; rebuild with scripts/build_spring2018.py.\n"
                      "window.SPRING_2018_QUESTIONS = " + json.dumps(questions, ensure_ascii=False, indent=2) + ";\n")
    print(f"Wrote {len(questions)} questions, {sum(len(q['parts']) for q in questions)} scored parts, 55.5 points.")


if __name__ == "__main__":
    main()
