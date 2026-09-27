"""Rebuild Fall 2018 only (Python: openpyxl, pypdf, Pillow; system: pdftoppm).

All question pages (physical PDF 4–27) were visually checked. Excel supplies
cell values and prompts; the PDF controls wording, units, headings, corrections,
the graph and scored parts. Answers summarize an accepted report method.
No existing exam data or browser storage is modified.
"""
import json
import re
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader
from build_spring2019 import displayed

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / 'past_exams/admissions_studytools_exam5_f18-5.pdf'
WORKBOOK = ROOT / 'past_exams_excel/CBT_Exam_5_F.18_v2.xlsx'

CHAPTERS = {
    1:['ratemaking-4'], 2:['ratemaking-5'], 3:['ratemaking-5','ratemaking-9'],
    4:['ratemaking-16'], 5:['ratemaking-6'], 6:['ratemaking-7'],
    7:['ratemaking-6','ratemaking-8','reserving-11'], 8:['ratemaking-13'],
    9:['ratemaking-10'], 10:['ratemaking-11'], 11:['ratemaking-15'],
    12:['ratemaking-9'], 13:['ratemaking-11'], 14:['reserving-1'],
    15:['reserving-1','reserving-14'], 16:['reserving-7','reserving-8','reserving-9','reserving-15'],
    17:['reserving-11'], 18:['reserving-10'],
    19:['reserving-7','reserving-8','reserving-10','reserving-15'],
    20:['reserving-6','reserving-13'], 21:['reserving-14'],
    22:['reserving-11','reserving-17'], 23:['reserving-16'],
    24:['reserving-7','reserving-15'],
}
SINGLE = {1:12, 5:38, 11:22, 14:3, 19:9, 21:29}

# (first heading row, last data row, data rows, columns, title, complete headings).
# Headings are from the printed PDF, never joined mechanically from Excel cells.
TABLES = {
 2:[(5,14,[9,10,12,13,14],[2,3,4,5,6,7],'Premium Transactions',
     ['Policy','Original Effective Date','Original Expiration Date','Transaction Effective Date','Full-Term Premium ($)','Notes']),
    (16,21,range(18,22),[2,3,4,5],'Loss Transactions',['Policy','Accident Date','Payment Date','Loss Payment ($)'])],
 3:[(5,10,range(7,11),list(range(2,9)),'',['Effective Date','Overall Average Rate Change','Rate Per Exposure ($)','Class Factor X','Class Factor Y','Class Factor Z','Expense Fee ($)'])],
 4:[(5,9,range(7,10),[2,3,4,5],'Loss Costs by Report Year Lag ($)',['Report Year','0','1','2'])],
 5:[(5,9,range(7,10),[2,5],'',['Accident Year','Incurred Loss and ALAE as of December 31, 2017 ($000s)']),
    (11,29,range(14,30),[2,3,4,5],'',['Year Ending Quarter','Frequency','Severity ($)','Pure Premium ($)']),
    (11,18,range(14,19),[7,8,9,10],'Annual Exponential Trends',['# of Points','Frequency','Severity','Pure Premium'])],
 7:[(5,11,range(9,12),[2,3,4,5],'Cumulative Reported Claim Counts as of (months)',['Accident Year','12','24','36']),
    (5,11,range(9,12),[7,8,9,10],'Cumulative Reported Loss + ALAE ($) as of (months)',['Accident Year','12','24','36']),
    (13,15,range(13,16),[2,3],'',['Value','Item']),
    (17,21,range(19,22),[2,3],'',['Calendar Year','Earned Exposures']),
    (23,29,range(23,30),[2,3],'',['Value','Item'])],
 8:[(5,7,range(5,8),[2,3],'',['Value','Item'])],
 11:[(5,14,range(5,15),[2,3],'',['Value','Item'])],
 12:[(7,10,range(8,11),[2,3],'',['Variable 1 Segment','Rating Factor 1']),
     (12,15,range(13,16),[2,3],'',['Variable 2 Segment','Rating Factor 2']),
     (17,22,range(20,23),[2,3,4,5],'Earned Exposures',['Variable 2','A','B','C']),
     (24,29,range(27,30),[2,3,4,5],'Reported Loss & ALAE ($)',['Variable 2','A','B','C'])],
 13:[(5,12,range(6,13),[2,3,4,5],'',['Size of Loss','Loss Distribution'])],
 15:[(5,17,range(7,18),[2,3,4,5,7],'',['Claim ID','Accident Date','Transaction Date','Gross Amount Paid on Transaction Date ($)','Gross Ending Case Outstanding ($)']),
     (19,20,range(19,21),[2,3],'',['Value','Item'])],
 16:[(8,13,range(10,14),[2,3,4,5,6],'Book A — Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
     (8,13,range(10,14),[8,9],'Book A',['Calendar Year','Earned Premium ($000)']),
     (15,16,range(15,17),[2,3],'',['Value','Item']),
     (20,25,range(22,26),[2,3,4,5,6],'Book B — Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
     (20,25,range(22,26),[8,9],'Book B',['Calendar Year','Earned Premium ($000)'])],
 17:[(5,11,range(8,12),[2,3,4],'',['Calendar Year','Earned Premium ($000)','On-Level Adjustment']),
     (5,9,range(8,10),[6,7,8],'',['Accident Year','Ultimate Claim Counts','Ultimate Severity ($)']),
     (13,15,range(13,16),[2,3],'',['Value','Item'])],
 18:[(4,10,range(7,11),[2,3,4,5,6],'Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
     (4,10,range(7,11),[8,9],'',['Calendar Year','Earned Premium ($000)']),
     (12,12,[12],[2,3],'',['Value','Item'])],
 20:[(5,10,range(7,11),[2,3,4,5,6],'Case Outstanding ($000) as of (months)',['Accident Year','12','24','36','48']),
     (12,17,range(14,18),[2,3,4,5,6],'Cumulative Paid Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
     (19,24,range(21,25),[2,3,4,5,6],'Open Claim Counts as of (months)',['Accident Year','12','24','36','48']),
     (26,27,range(26,28),[2,3],'',['Value','Item'])],
 21:[(5,10,range(7,11),[2,3,4,5,6],'Cumulative Received Salvage and Subrogation (S&S) ($000) as of (months)',['Accident Year','12','24','36','48']),
     (12,17,range(14,18),[2,3,4,5,6],'Cumulative Paid Claims Gross of S&S ($000) as of (months)',['Accident Year','12','24','36','48']),
     (19,25,range(22,26),[2,3],'',['Accident Year','Selected Ultimate Claims Gross of S&S ($000)'])],
 22:[(5,10,range(7,11),[2,3,4],'',['Calendar Year','Paid Claims ($)','Paid ULAE ($)']),
     (12,17,range(14,18),[2,3,4,5,6,7],'',['Accident Year','Paid Claims ($)','Reported Claim Counts','Closed Claim Counts','Ultimate Claim Counts','Ultimate Claims ($)']),
     (19,20,range(19,21),[2,3],'',['Value','Item'])],
 23:[(5,9,range(7,10),[2,3,4,5],'Cumulative Paid Claims Only ($) as of (months)',['Accident Year','12','24','36']),
     (11,16,range(13,17),[2,3],'',['Accident Year','Calendar Year 2017 Paid Claims Only ($)']),
     (18,23,range(20,24),[2,3,4,5,6],'Cumulative Paid ALAE ($) as of (months)',['Accident Year','12','24','36','48'])],
 24:[(5,12,range(8,13),[2,3,4],'',['Age (Month)','Cumulative Paid Development Factors to Ultimate','Cumulative Reported Development Factors to Ultimate']),
     (14,17,range(15,18),[2,3],'Accident year 2017 as of March 31, 2018:',['Value','Item']),
     (19,21,range(20,22),[2,3],'Accident year 2017 as of May 31, 2018:',['Value','Item'])],
}

# Explicit PDF corrections to the workbook's values, dates, wording and units.
CELLS = {
 (2,'G10'):'Additional Premium for Endorsement',
 (5,'D20'):'19,200', (5,'D26'):'20,445', (5,'D27'):'20,882',
 (5,'B26'):'March 31, 2017',
 (7,'B13'):'$98,000',(7,'B14'):'$318,000',(7,'B15'):'3%',
 (7,'B23'):'5%',(7,'B24'):'3%',(7,'B25'):'$21',(7,'B26'):'15%',
 (7,'B27'):'10%',(7,'B28'):'2%',(7,'B29'):'$950',
 (8,'B5'):'$500',(8,'B6'):'$600',(8,'B7'):'$510',
 (8,'C7'):"All competitors' average premium",
 (11,'B5'):'$500,000',(11,'B8'):'$2,000,000',(11,'B9'):'$100,000',
 (11,'C14'):'Additional risk margin as a % of excess losses',
 (15,'B19'):'60%',(15,'B20'):'$1,500',
 (16,'B15'):'75%',(17,'B13'):'-1.3%',(17,'B14'):'6.0%',(17,'B15'):'15%',
 (20,'B26'):'7.5%',(22,'B19'):'60%',(22,'B20'):'40%',
 (21,'B3'):'Given the following as of December 31, 2017:',
 (21,'B29'):'Estimate ultimate salvage and subrogation for accident year 2017 using a ratio approach.',
 (22,'B3'):'Given the following information as of December 31, 2017:',
 (22,'B28'):'Estimate unpaid ULAE as of December 31, 2017.',
 (23,'B3'):'Given the following information as of December 31, 2017:',
 (23,'B25'):'• There is no development beyond 48 months.',
 (24,'B33'):'Describe a situation in which the actuary would revise the March 31, 2018 estimate of ultimate claims',
 (24,'B37'):'Describe a situation in which the actuary would not revise the March 31, 2018 estimate of ultimate',
}

def value(n, cell):
    if (n, cell.coordinate) in CELLS:
        return CELLS[n, cell.coordinate]
    if isinstance(cell.value, datetime):
        return f'{cell.value:%B} {cell.value.day}, {cell.value.year}'
    result = displayed(cell)
    if isinstance(cell.value, (int,float)) and cell.value >= 1000 and cell.value == int(cell.value) and '%' not in cell.number_format:
        result = f'{cell.value:,.0f}' if cell.column != 2 else str(int(cell.value))
    return re.sub(r'\s+([.,:])', r'\1', result)

def lines(sheet, n, first, last):
    return ' '.join(value(n,c) for row in sheet.iter_rows(min_row=first,max_row=last) for c in row[1:] if c.value is not None and c.data_type != 'f').strip()

def blocks(sheet, n, stop):
    result = []
    specs = TABLES.get(n, [])
    covered = {r for start,end,*_ in specs for r in range(start,end+1)}
    previous = None
    for r in range(3,stop):
        for start,end,rows,cols,title,headers in specs:
            if start != r:
                continue
            data = [[value(n,sheet.cell(row,col)) for col in cols] for row in rows]
            if headers[0] in ('Accident Year','Calendar Year','Report Year'):
                for row in data: row[0] = row[0].replace(',','')
            if n == 13:
                # Printed loss intervals and percentages; Excel spreads intervals over three cells.
                data = [['X ≤ $400,000','50.0%'],['$400,000 < X ≤ $550,000','25.0%'],['$550,000 < X ≤ $700,000','10.0%'],['$700,000 < X ≤ $850,000','10.0%'],['$850,000 < X ≤ $1,000,000','2.5%'],['$1,000,000 < X ≤ $1,500,000','2.5%'],['Total:','100.0%']]
            if n == 15 and start == 5:
                for i,row in enumerate(data):
                    if not row[0]: row[:2] = data[i-1][:2]
            b = dict(type='table',title=title,headers=headers,rows=data)
            if n == 12 and start in (17,24):
                b['groups'] = [dict(label='',span=1),dict(label='Variable 1',span=3)]
            result.append(b)
        if r in covered or (n == 16 and r in (6,18)):
            previous = None
            continue
        text = lines(sheet,n,r,r)
        if not text:
            previous = None
        elif previous == r-1 and not text.startswith(('•','i.','ii.','iii.')):
            result[-1]['text'] += ' ' + text
            previous = r
        else:
            result.append(dict(type='line',text=text))
            previous = r
    return result

def report_pages():
    pdf = PdfReader(PDF)
    starts = []
    for p,page in enumerate(pdf.pages,1):
        match = re.search(r'QUESTION\s*:?\s*(\d+)\s+TOTAL POINT VALUE',page.extract_text() or '',re.I)
        if match: starts.append((int(match[1]),p))
    assert [n for n,p in starts] == list(range(1,25))
    return {n:list(range(p,starts[i+1][1] if i+1<len(starts) else len(pdf.pages)+1)) for i,(n,p) in enumerate(starts)}

# Each tuple is a readable sample solution and a concise report-grounded insight.
ANSWERS = {
 1:[("Assess proportionality to expected loss, practicality, and historical precedent. Boat value is not proportional to liability damage to others, but does relate to physical-damage repair or replacement severity. Value can be difficult to define and verify and may be manipulated; this affects both coverages. Changing the base also requires systems/data changes and may cause premium swings for both. Retain boat-years: more boats produce more claims, boat counts are objective and readily verified, and no conversion cost is needed.","Evaluate all three criteria for both coverages and make a supported recommendation. Other justified recommendations were accepted.")],
 2:[("2017 calendar year earned premium = 800 × 6/12 + 400 × 3/12 + 1,000 × 9/12 + 500 × 3/12 = $1,375.","The $400 endorsement premium is full-term, so prorate it. The 2018 cancellation does not affect 2017 calendar-year earnings."),
    ("Only Policies B and C belong to policy year 2017. Earned premium = 1,000 + 500 × 6/12 = $1,250 after C's cancellation. Losses = 500 + 750 = $1,250. Loss ratio = 1,250 / 1,250 = 100%.","Include the cancellation and use policy-year membership, not calendar-year transactions.")],
 3:[("2017 average rate level = 0.75 × 1 + 0.25 × 1.10 = 1.025. Current rate level = 1.10 × 1.01 = 1.111. On-level factor = 1.111 / 1.025 = 1.0839.","Use the semiannual-policy parallelogram weights and each applicable overall rate level."),
    ("Original premium = 1,000 × 0.85 + 120 = $970. Current premium = 1,175 × 0.75 + 132 = $1,013.25. On-level factor = 1,013.25 / 970 = 1.0446.","Use Class Y and include the expense fee at both rate levels."),
    ("An aggregate parallelogram adjustment is inappropriate because the overall average rate change does not capture the different class-factor changes on October 1, 2017. Applying the method separately by class using each class's own rate impact can be appropriate.","Explain the difference between overall and class-specific rate changes.")],
 4:[("$1,000 = 500 + 300 + 200, covering all report-year 2015 loss costs.","An answer of $500 required explicitly assuming a first-year claims-made policy."),
    ("$1,050 = 500 + 330 + 220, following accident year 2015 through its reporting lags.","Occurrence coverage follows the accident year across report years."),
    ("Occurrence pricing is more affected because it is exposed to both reporting and settlement lag. Claims-made coverage removes the reporting lag, leaving a shorter period over which an unexpected trend increase can affect costs.","Explain why the lag difference matters, rather than merely calling claims-made shorter-tailed."),
    ("Claims-made coverage has no pure IBNR (IBNYR) for claims not yet reported under the policy; it retains IBNER from settlement development. Removing reporting uncertainty reduces reserve inadequacy risk.","Distinguish IBNYR from IBNER; do not simply say there is no IBNR."),
    ("The period between premium collection and claim payment is shorter, so funds are invested for less time and earn less investment income.","Compare premium collection to payment, not accident date to payment.")],
 5:[("Select a stable 5% severity trend. Use 6% frequency trend to July 1, 2017 and −7% thereafter to reflect the underwriting change. The average future accident date is April 1, 2020 (two years of semiannual policies beginning January 1, 2019). Trend July 1, 2015 to July 1, 2017 for 2 years and then to April 1, 2020 for 2.75 years: 15,000 × (1.06 × 1.05)^2 × (0.93 × 1.05)^2.75 = 17,405.25 ($000s). No further development is needed at 36 months.","Use two-step trending, justify both selections, and calculate the projection date for semiannual policies and a two-year rate period.")],
 6:[("Separate fixed and variable expenses. Premium-based: divide each portion by written or earned premium, matching when the expense is incurred, to obtain separate expense ratios. Exposure-based: divide fixed expenses by written or earned exposures (or policy count), producing dollars per exposure; variable expenses remain a ratio to premium. Project the selected provisions to the future period as appropriate.","Explicitly distinguish fixed from variable costs and identify the fixed-expense denominator for each method."),
    ("Premium-based: rate changes can alter the fixed-expense ratio even when fixed costs have not changed. Exposure-based: economies of scale in a changing book can make historical fixed expenses per exposure unrepresentative of future costs.","Describe a specific distortion for each approach, not merely data availability concerns.")],
 7:[("Net reinsurance cost = 318,000 − 98,000 = $220,000. Project 2017 exposures two years to the 2019 reinsurance term: 17,000 × 1.03^2 = 18,035.3. Cost per exposure = 220,000 / 18,035.3 ≈ $12.20.","Use the latest year's exposures and the reinsurance contract's period, not the rate revision period."),
    ("Reported severities (12, 24, 36 months) are 2015: 15,000, 18,000, 19,800; 2016: 14,500, 17,400; 2017: 15,500. Select severity factors 1.20 and 1.10, with no tail. Count factors are 0.95 and approximately 0.98. Ultimate loss and ALAE: 2015 = $8,850,600; 2016 = 9,256,800 × 1.10 × 0.98 = $9,978,830; 2017 = 9,145,000 × 1.32 × 0.931 = $11,238,473.","Develop counts and severities separately; a direct chain-ladder answer does not demonstrate the requested frequency-severity technique. Rounding produces accepted variations."),
    ("Trend to April 1, 2020. Annual pure premiums are 8,850,600 / 14,000 × 1.03^4.75 ≈ $727; 9,978,830 / 15,000 × 1.03^3.75 ≈ $743; and 11,238,473 / 17,000 × 1.03^2.75 ≈ $717. Their equal-weight average is about $729 before ULAE, or about $765.7 including the 5% ULAE provision.","Give equal weight to each year's pure premium. Do not trend the exposure denominators. ULAE may be included here or in part d, but not omitted or counted twice."),
    ("Indicated premium ≈ (729 × 1.05 + 21 + 12.20) / (1 − 0.15 − 0.10 − 0.02) = $1,094. Indicated change ≈ 1,094 / 950 − 1 = 15.2% (about 15.16% with the report's rounding).","Include net reinsurance cost, ULAE, and the contingency provision in their proper places.")],
 8:[("Retention may fall as policyholders move to cheaper competitors. Profit per retained risk should increase as premium moves toward the indicated level.","State two distinct consequences; adverse selection or a worsening loss ratio was not credited as an unsupported consequence."),
    ("Price: a large increase encourages customers to shop and choose a cheaper competitor. Service: satisfied customers are more likely to renew because they value the insurer's service and claims handling.","Explain how each factor affects renewal rather than merely naming it."),
    ("Reduce operating expenses; tighten underwriting to favor more profitable risks; introduce loss-mitigation programs to reduce claims.","Give three non-pricing actions and specify the direction of any change in mix or coverage."),
    ("Different risk and coverage mixes make average premiums incomparable. Compare quotes for the same risk profile and coverage, or re-rate the company's book under competitors' filed rating plans.","The proposed solution must address the specific comparability issue identified.")],
 9:[("Do not include number of vehicles on this evidence. It fails the main-effect test: nearly all error ranges contain 1.00. It fails the consistency test: results vary materially across years outside the first few vehicle counts. The 10% chi-square percentage exceeds a 5% significance threshold, so the statistical test also fails.","State a decision and correctly apply the main-effect, consistency and statistical tests to the supplied graphs."),
    ("Premiums must be on-leveled; there is no standard default distribution for loss ratios; and a rate change can make a fitted loss-ratio model obsolete.","Give challenges specific to loss-ratio modeling rather than generic GLM difficulties."),
    ("Univariate analysis may be preferable for a simple rating plan because it is easier to explain and more transparent to stakeholders.","Connect the reason explicitly to the comparison with GLM analysis.")],
 10:[("Two broad territories ignore important within-territory differences such as urban versus rural risks. Weather alone also misses geographic variation in theft, fire and other perils, leaving heterogeneous risks grouped together.","Provide two disadvantages of the existing approach."),
     ("1. Define basic geographic units such as ZIP codes or counties; consider granularity and stability of boundaries. 2. Estimate systematic geographic risk with a GLM using demographic/physical variables; control for correlation with other rating factors. 3. Estimate and spatially smooth residual geographic risk; choose distance or adjacency weights suited to the peril. 4. Cluster the units into rating territories; consider the clustering method and desired exposure balance.","Describe all four steps, with a consideration for each, including final clustering.")],
 11:[("Loss below deductible = 2,000,000 × 90% = $1,800,000; excess loss = $200,000; all ALAE = $200,000. Processing cost = 3% × 1,800,000 = $54,000; credit risk = $18,000; excess risk margin = 7% × 200,000 = $14,000. Premium = (200,000 + 200,000 + 100,000 + 54,000 + 18,000 + 14,000) / (1 − 0.12 − 0.04) = $697,619.","Include all ALAE and fixed expenses; apply processing and credit costs to losses below the deductible and the risk margin to excess losses. The report's second sample has an arithmetic typo; its formula also yields $697,619.")],
 12:[("Adjust exposures using Variable 2 factors. A: 800 × .75 + 300 × .95 + 500 = 1,385; B: 1,500 × .75 + 750 × .95 + 500 = 2,337.5; C: 600 × .75 + 500 × .95 + 1,500 = 2,425. Loss totals are $795,000, $4,735,000 and $2,500,000. Adjusted pure premiums are $574.01, $2,025.67 and $1,030.93. Rebase to C: A = 0.5568, B = 1.9649, C = 1.0000.","Use exposures adjusted for the other variable and rebase to C."),
     ("Using rounded indicated factors 0.56, 1.96 and 1.00, selected factors are 0.73, 1.98 and 1.00. Preserve total premium: new base = 1,000 × (1,385 × .90 + 2,337.5 × 2 + 2,425) / (1,385 × .73 + 2,337.5 × 1.98 + 2,425) ≈ $1,035. Precise unrounded relativities give a slightly different accepted result.","Use selected factors and adjusted exposure weights to offset the classification change."),
     ("Affordability: insurance should remain affordable. Controllability: insureds should be able to influence their class and obtain a lower rate through their actions.","Give social criteria rather than legal or operational criteria.")],
 13:[("At $1,500,000 coverage, mean severity = .5(200,000) + .25(475,000) + .1(625,000) + .1(775,000) + .025(925,000) + .025(1,250,000) = $413,125. Rate per $1,000 = 413,125 × .02 / 1,500 = $5.51. At $800,000, cap all losses at the limit. The $700,000–$850,000 layer has mean payment (2/3)750,000 + (1/3)800,000 = $766,666.67. Total capped severity is $397,916.67, giving 397,916.67 × .02 / 800 = $9.95 per $1,000.","Use layer midpoints, cap the straddling layer correctly, retain higher-layer probability, and include frequency."),
     ("Insured: a total or near-total loss may not be fully covered, leaving funds needed to rebuild. Insurer: rates based on full insurance-to-value underprice underinsured homes because expected loss does not decrease proportionally with the amount insured.","Explain the directional problem from both perspectives."),
     ("Coinsurance apportionment = min[1,000,000 / (1,500,000 × .80), 1] = 5/6. (i) $800,000 loss: indemnity = $666,666.67 and penalty = $133,333.33. (ii) $1,200,000 loss: indemnity = min(1,200,000 × 5/6, 1,000,000) = $1,000,000. Penalty = min(loss, limit) − indemnity = $0.","The $200,000 above the policy limit in case ii is uninsured loss, not a coinsurance penalty.")],
 14:[("Accurate unpaid claims support pricing, underwriting, reinsurance and strategy decisions. Overstated reserves can prompt unnecessary rate increases or tighter underwriting; understated reserves can delay needed corrective action.","Describe consequences for management with sufficient detail, not just a brief label."),
     ("Investors rely on financial statements to assess strength, profitability and dividends. Incorrect reserves can make the insurer appear stronger or weaker than it is and distort investment decisions.","Explain how reserve accuracy affects investors' decisions."),
     ("Regulators use unpaid claims estimates to assess solvency and rate adequacy. Understated reserves can conceal financial weakness and delay intervention until insolvency is difficult to prevent.","Describe the regulatory impact rather than saying regulators determine the reserve level.")],
 15:[("2015 gross reported = (1,000 + 550) for A + (300 + 1,050) for B = $2,900.","Reported claims equal paid claims plus the change in case outstanding; do not apply reinsurance."),
     ("2016 net paid = (500 + 600 + 450) × .40 + 1,200 + 400 = $2,220. C and D remain below their individual paid-loss retentions.","Retain 40%, not 60%, for 2015 claims and apply the $1,500 retention per 2016 claim."),
     ("2016 gross reported: A = 500 + 225 − 550 = 175; B = 600 + 450 + 150 − 1,050 = 150; C = 1,200 + 575 = 1,775; D = 400 + 900 = 1,300. Total = $3,400.","Include the change in case outstanding for earlier claims and do not apply reinsurance."),
     ("A: .40 × (725 − 225) = $200. B: $0. C was already above the $1,500 reported retention at year-end 2016, so its net reported change is $0. D: min(400 + 800 + 625, 1,500) − (400 + 900) = $200. Total = $400.","Apply quota share to A and compare capped cumulative reported amounts for C and D.")],
 16:[("36-to-48 factor = 123,700 / 112,500 ≈ 1.10. Including the 1.06 tail, 36-to-ultimate ≈ 1.166. Ultimate = 111,100 × 1.166 ≈ 129,543 ($000). Using unrounded factors gives about 129,491 ($000).","Include the tail and use adjacent development ages; rounding variations were accepted."),
     ("Select the average 24-to-36 factor: [(112,500 / 92,000) + (111,100 / 92,600)] / 2 ≈ 1.2114. The 24-to-ultimate factor is about 1.412. BF ultimate = 94,400 + 182,800 × .75 × (1 − 1/1.412) ≈ 134,400 ($000).","Use the cumulative unreported proportion and the supplied 75% expected claims ratio."),
     ("2017 expected ultimate claims = 184,200 × .75 = 138,150 ($000).","Use the given expected claims ratio rather than estimating another one."),
     ("Use Book A's development pattern to develop Book B's 2015 reported claims. Book B is small and volatile; Book A is more credible and writes the same state and line. Explicitly assume the two books have comparable development patterns.","Recommend a specific approach and justify borrowing Book A's experience.")],
 17:[("Use on-level premium as the exposure proxy. Trend historical frequencies to 2017: 2,200 × .987^3 / (127,500 × .71) = .02336 and 1,970 × .987^2 / (117,600 × .66) = .02472 claims per $1,000 premium. Select approximately .024. Trend severities and apply the legislative factor: 32,600 × 1.06^3 × .85 ≈ $33,003 and 35,300 × 1.06^2 × .85 ≈ $33,714; select about $33,358. For 2016 reverse the 2017 adjustments: frequency ≈ .024 / .987 × .85 = .02067 per $1,000 historical premium; severity ≈ 33,358 / (1.06 × .85) = $37,024. Ultimate 2016 ≈ 64,300 × .0207 × 37,024 / 1,000 = 49,300 ($000). Ultimate 2017 is about 47,200 ($000), with the report showing 47,247 using its rounded count of 1,416.","Apply frequency and severity trends separately, on-level premium consistently, and the 15% legislative saving only to 2017. The report's displayed 2016 frequency formula has a typographical error; .85 belongs in the numerator."),
     ("Claim-count definitions must be consistent across years. Claim mix must remain sufficiently homogeneous for historical frequency and severity patterns to represent future experience.","State technique-specific assumptions, not generic data-quality requirements.")],
 18:[("Using simple-average factors, select approximately 1.9963, 1.4932, 1.3889 and 1.3. CDFs at ages 12, 24, 36 and 48 are 5.3822, 2.6961, 1.8056 and 1.3. Used-up premium = 5,300/1.3 + 7,200/1.8056 + 7,800/2.6961 + 8,500/5.3822 ≈ 12,536.8. Expected claims ratio = (2,500 + 2,300 + 1,900 + 1,100)/12,536.8 ≈ .6222. Ultimate 2017 = 1,100 + 8,500 × .6222 × (1 − 1/5.3822) ≈ 5,406 ($000).","Estimate the ratio from total reported claims divided by used-up premium, then add expected unreported claims to reported claims."),
     ("Accurate pricing need not change the underlying expected claims ratio, but the new mix appears to report faster. Older, higher development factors understate used-up premium, overstate the Cape Cod expected claims ratio and unreported percentage, and therefore overstate 2017 ultimate claims.","Address changed reporting patterns and the direction of bias, even when both classes are priced accurately.")],
 19:[("(i) Development: old factors are too high for the faster reporting pattern, so IBNR is overstated; higher reported loss from younger drivers amplifies the dollar effect. (ii) Expected claims: if each class is priced adequately, its premium reflects its loss cost, so the expected claims ratio remains appropriate and reporting speed does not bias the method. If younger drivers are underpriced, the mix shift can understate expected claims and IBNR. (iii) Cape Cod: excessive development factors understate used-up premium and inflate the estimated claims ratio and unreported proportion, overstating IBNR, generally less than the development technique.","Address both reporting speed and business mix for all three techniques. State the premium-adequacy assumption for the expected claims technique.")],
 20:[("Average case outstanding ($000) by development age: 2014 = 140, 110, 130, 120; 2015 = 135, 150, 144; 2016 = 180, 152; 2017 = 180. Changes down several columns differ from the 7.5% severity trend and generally indicate strengthening case adequacy.","Examine multiple ages and compare changes with severity trend; a single age is insufficient."),
     ("Restate each age's prior average case outstanding from the latest diagonal using 7.5% annual severity trend, multiply by historical open counts, then add paid claims. Adjusted reported claims ($000) are approximately: 2014 = 62,961, 82,762, 165,767, 285,600; 2015 = 62,758, 79,800, 161,000; 2016 = 66,443, 88,400; 2017 = 72,600. Select factors about 1.305, 2.0105, 1.723 and 1.05. Ultimate 2017 = 72,600 × 1.305 × 2.0105 × 1.723 × 1.05 ≈ 344,608 ($000).","Restate to the latest diagonal's adequacy level, detrend historical average case amounts correctly, and include the 1.05 tail."),
     ("The adjusted estimate is lower: applying historical unadjusted factors to strengthened current reported claims would overstate ultimate claims.","Compare the direction of the adjusted result with the unadjusted result."),
     ("The result is highly sensitive to the judgmental severity trend selection; a wrong trend can materially distort the reserve estimate.","Describe a limitation of the adjustment, not just a situation where it is unsuitable.")],
 21:[("Form cumulative S&S / cumulative paid gross claims ratios at each age. Select approximate multiplicative ratio-development factors 1.217, 1.0285 and 1.014, with no tail. Ultimate ratios for 2014–2016 are about .435, .408 and .4224; their average is about .4218. Treat the higher developed 2017 ratio (about .4823) as random fluctuation and select .4218. Ultimate 2017 S&S = .4218 × 16,400 = 6,917.52 ($000). Retaining the higher 2017 ratio with a justified trend selection was also accepted.","Develop S&S-to-paid ratios, consider earlier years' ultimate ratios, and justify the final selection. Do not use ultimate claims as the denominator of the historical ratios.")],
 22:[("Use average severity on unsettled claims: (ultimate claims − paid claims)/(ultimate count − closed count), multiplied by ultimate count − reported count. IBNYR: 2014 = $0; 2015 = (21,400 − 18,000)/(330 − 270) × 30 = $1,700; 2016 = 9,500/140 × 55 ≈ $3,732; 2017 = 16,800/255 × 135 ≈ $8,894. Total ≈ $14,326.","Estimate unreported claim dollars, not just counts or total unpaid claims. Other reasonable severity assumptions were accepted."),
     ("Select a 5% ULAE-to-paid-claims ratio from the calendar-year experience. Total unpaid claims = 3,400 + 9,500 + 16,800 = $29,700. Using part a, case outstanding plus IBNER ≈ 29,700 − 14,326 = $15,374. Unpaid ULAE = .05 × [14,326 + .40 × 15,374] ≈ $1,024.","Apply 100% of ULAE effort to IBNYR and 40% to other unpaid claims; avoid double counting IBNYR.")],
 23:[("The new diagonal is 2014: 626,900 + 75,200 = $702,100; 2015: 453,600 + 158,800 = $612,400; 2016: 170,000 + 289,000 = $459,000; 2017: $172,000.","Add each calendar-year payment to its matching accident year's prior diagonal."),
     ("Construct paid ALAE / paid claims ratios. Approximate selected additive age-to-age changes are .024, .007 and .014; their sum is .045. The 2017 ultimate ratio is .033 + .045 = .078. Paid-claim factors are about 2.70, 1.35 and 1.12, yielding ultimate claims 172,000 × 2.70 × 1.35 × 1.12 = $702,172.80. Ultimate ALAE = 702,172.80 × .078 = $54,769.48.","Use additive changes for the ALAE ratio and apply the ultimate ratio to ultimate claims, not current paid claims."),
     ("An incorrect ultimate claims estimate propagates into the ALAE estimate because the latter is a ratio of claims.","Identify a disadvantage of the ratio method.")],
 24:[("At 15 months the reported proportion is 1/1.46 = .68493; at 18 months it is 1/1.38 = .72464. Interpolate to 17 months: .68493 + (2/3)(.72464 − .68493) = .71139. Expected incremental reported = (3,300 − 2,400) × (.71139 − .68493)/(1 − .68493) ≈ $75.61. Expected cumulative ≈ $2,475.61; actual $2,750 is about $274 higher. The report gives $2,475.66 using rounded proportions.","Interpolate percentages, not development factors; apply emergence to unreported claims and then add reported-to-date."),
     ("At 15 months the paid proportion is .50; at 18 months it is 1/1.65 = .60606. At 17 months it is .50 + (2/3)(.60606 − .50) = .570707. Expected incremental paid = (3,300 − 1,820) × (.570707 − .50)/.50 = $209.29. Expected cumulative = $2,029.29; actual $2,050 is only $20.71 higher.","Use the unpaid amount and compare cumulative amounts after interpolation."),
     ("Increase the ultimate if the higher reported amount comes from a large unpaid claim expected to develop beyond the current IBNR provision.","Identify a change in ultimate cost, rather than a change only in reporting or settlement timing."),
     ("Do not revise the ultimate if the excess reported emergence comes from stronger case reserves while paid emergence remains close to expected. This changes the reported pattern without necessarily changing ultimate cost.","An organizational timing or adequacy change can explain the discrepancy; an actual large loss would still affect this accident year's ultimate.")],
}

def main():
    workbook = load_workbook(WORKBOOK,data_only=True)
    pages = report_pages()
    questions = []
    for n in range(1,25):
        sheet = workbook[str(n)]
        markers = [(c.row,c.value.strip()[1]) for row in sheet for c in row if c.column==1 and isinstance(c.value,str) and re.fullmatch(r'\([a-g]\)',c.value.strip())]
        if not markers: markers = [(SINGLE[n],'a')]
        parts = []
        for i,(start,letter) in enumerate(markers):
            end = markers[i+1][0]-1 if i+1<len(markers) else sheet.max_row
            parts.append(dict(id=letter,points=workbook['Point Grid'].cell(n+7,3+i).value,prompt=lines(sheet,n,start,end)))
        if n == 14:
            parts = [dict(id=letter,points=.5,prompt=f'For {stakeholder}, describe the importance of having accurate unpaid claim estimates.') for letter,stakeholder in zip('abc',['Internal Management','Investors','Regulators'])]
        assert len(parts) == len(ANSWERS[n]), n
        for part,(solution,insight) in zip(parts,ANSWERS[n]):
            part.update(solution=solution,insight=insight)
        total = sum(p['points'] for p in parts)
        assert total == workbook['Point Grid'].cell(n+7,2).value, n
        q = dict(id=f'fall-2018-{n}',number=n,exam='Fall 2018',chapterIds=CHAPTERS[n],points=total,questionPage=n+3,solutionPages=pages[n],sourceBlocks=blocks(sheet,n,markers[0][0]),parts=parts)
        if n == 9:
            q['figure'] = dict(src='assets/exam-graphs/fall-2018-q9.png',title='Number of vehicles GLM output',alt='Two original GLM graphs: exposures and indicated relativity with upper and lower standard error bounds by number of vehicles; exposures and relativities for 2015, 2016 and 2017. Vehicle counts range from 1 to 7+, with a 1.0 reference line.')
        if n == 14:
            q['sourceBlocks'] = [dict(type='line',text='For each of the following stakeholders, describe the importance of having accurate unpaid claim estimates: i. Internal Management; ii. Investors; iii. Regulators.')]
        questions.append(q)
    assert sum(q['points'] for q in questions) == 55
    assert sum(len(q['parts']) for q in questions) == 64
    (ROOT/'question_data/fall2018_questions.js').write_text('// Fall 2018: PDF-checked transcription; regenerate with scripts/build_fall2018.py.\nwindow.FALL_2018_QUESTIONS = '+json.dumps(questions,ensure_ascii=False,indent=2)+';\n')
    # Direct original-PDF crop, including both plots, axes, and legends.
    with tempfile.TemporaryDirectory() as directory:
        prefix = str(Path(directory)/'q9')
        subprocess.run(['pdftoppm','-f','12','-l','12','-scale-to','2400','-singlefile','-png',str(PDF),prefix],check=True)
        im = Image.open(prefix+'.png')
        w,h = im.size
        im.crop((int(w*78/1160),int(h*225/1500),int(w*1086/1160),int(h*1028/1500))).save(ROOT/'assets/exam-graphs/fall-2018-q9.png')
    print(f'Wrote {len(questions)} questions, {sum(len(q["parts"]) for q in questions)} scored parts, 55 points.')

if __name__ == '__main__':
    main()
