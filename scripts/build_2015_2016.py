"""Build the Spring 2016 and Fall 2015 question sets from official PDFs.

The CBT workbooks provide question structure and point grids. Printed PDFs
control the table headings, wording, corrected values, sample answers, and
examiner guidance. The Spring 2016 question booklet is scanned.
"""
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from build_spring2019 import displayed

ROOT = Path(__file__).resolve().parents[1]
CONFIG = {
  'spring-2016': {
    'workbook':'CBT_Exam_5_S.16_v01.xlsx', 'pdf':'admissions_studytools_exam5_sp16-5.pdf',
    'var':'SPRING_2016_QUESTIONS', 'count':25, 'part_count':63, 'points':57.75,
    'report_start':[32,34,37,40,42,44,46,49,53,55,57,60,62,66,70,72,74,76,79,81,83,86,90,93,96],
  },
  'fall-2015': {
    'workbook':'CBT_Exam_5_F.15_v01.xlsx', 'pdf':'admissions_studytools_exam5_f15-5.pdf',
    'var':'FALL_2015_QUESTIONS', 'count':25, 'part_count':53, 'points':55.75,
    'report_start':[34,36,39,41,44,47,49,51,53,57,59,62,65,66,68,70,72,74,77,80,82,84,86,86,92],
  },
}
CHAPTERS = {
  'spring-2016': [
    ['ratemaking-5','ratemaking-6'],['ratemaking-4','ratemaking-5'],['ratemaking-4','ratemaking-9'],
    ['ratemaking-6'],['ratemaking-16'],['ratemaking-8','ratemaking-12'],['ratemaking-7'],
    ['ratemaking-8','reserving-7'],['ratemaking-15'],['ratemaking-9','ratemaking-10'],
    ['ratemaking-11'],['ratemaking-9'],['ratemaking-9'],['reserving-7'],['reserving-1'],
    ['reserving-7','reserving-8','reserving-9'],['reserving-12'],['reserving-9','reserving-10'],
    ['reserving-11'],['reserving-13'],['reserving-6','reserving-13'],
    ['reserving-16','reserving-17'],['reserving-16'],['reserving-15'],['reserving-15']],
  'fall-2015': [
    ['ratemaking-4','ratemaking-5'],['ratemaking-4'],['ratemaking-12'],['ratemaking-6'],
    ['ratemaking-6'],['ratemaking-7'],['ratemaking-8','ratemaking-6'],['ratemaking-7'],
    ['ratemaking-9'],['ratemaking-10'],['ratemaking-9'],['ratemaking-11'],['ratemaking-15'],
    ['reserving-1'],['reserving-1'],['reserving-6'],['reserving-7','reserving-8'],
    ['reserving-9'],['reserving-15'],['reserving-11'],['reserving-13'],['reserving-7'],
    ['reserving-14'],['reserving-16'],['reserving-15']],
}
# (first covered row, last covered row, data rows, workbook columns,
#  complete printed title, complete printed column headings)
T = {
 'spring-2016': {
  1:[(5,8,range(6,9),[2,3,4],'',['Accident Year','Earned Premium ($000)','Ultimate Losses ($000)'])],
  2:[(5,10,range(6,11),[2,3,4,5,6,7,8],'',['Policy','Original Effective Date','Original Expiration Date','Transaction Effective Date','Territory','Full-Term Written Premium ($)','Notes'])],
  4:[(5,16,range(6,17),[2,3,4,5,6,7,8,9],'',['Claim','Policy Effective Date','Accident Date','Report Date','Transaction Date','Claim Status','Loss Payment ($)','Case Reserve Change ($)'])],
  5:[(7,15,range(8,16),[2,3,4],'',['Accident Date','Report Date','Claim Amount ($)'])],
  6:[(6,11,range(6,12),[2,4],'',['Item','Value'])],
  7:[(6,11,range(8,12),[2,3,4,5,6],'Expense Ratios by Calendar Year',['Expense','2013','2014','2015','Percent Fixed'])],
  8:[(7,10,range(9,11),[2,3],'Rate Change History',['Effective Date','Change']),
     (12,14,range(13,15),[2,3,4,5],'',['Measure','2013','2014','2015']),
     (16,21,range(19,22),[2,3,4,5],'Cumulative Reported Loss and ALAE ($) as of (months)',['Accident Year','12','24','36']),
     (23,28,range(26,29),[2,3,4,5],'Cumulative Reported Loss and ALAE Excluding Catastrophes ($) as of (months)',['Accident Year','12','24','36'])],
  9:[(18,21,range(19,22),[2,3,4],'',['Loss Limit ($000)','Loss Elimination Ratio','Excess Ratio'])],
  10:[(6,9,range(7,10),[2,3,4,5,6],'',['Class','Premium at Current Rate Level ($)','Reported Loss and ALAE ($)','Number of Claims','Current Relativity'])],
  11:[(6,11,range(8,12),[2,3,4,5,6,7,8],'Loss Distribution by Policy Limit',['Size of Loss','$100,000 Limit — Claims','$100,000 Limit — Losses ($000)','$250,000 Limit — Claims','$250,000 Limit — Losses ($000)','$500,000 Limit — Claims','$500,000 Limit — Losses ($000)']),
      (24,29,range(25,30),[3,4],'',['Limit of Liability ($)','Increased Limits Factor'])],
  12:[(6,10,range(8,11),[2,3,4,5],'Building Type — Rating Factor',['Building Type','Current','Proposed','Exposures']),
      (12,16,range(14,17),[2,3,4,5],'Years Since Claim — Discount',['Years Since Claim','Current','Proposed','Exposures'])],
  13:[(8,11,range(9,12),[2,3,4,5,6],'',['Territory','Exposures','Indemnity Loss and ALAE ($)','Medical Loss and ALAE ($)','Workers Compensation Total Current Relativity'])],
  14:[(6,12,range(9,13),[2,3,4,5,6],'Company Cumulative Paid Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (14,19,range(17,20),[2,3,4,5],'Industry Paid Claims Age-to-Age Factors',['Accident Year','12–24','24–36','36–48'])],
  15:[(6,10,range(8,11),[2,3,4,5],'Personal Auto: Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36']),
      (12,16,range(14,17),[2,3,4,5],'Commercial Auto: Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36'])],
  16:[(6,10,range(7,11),[2,3,4,5],'',['Accident Year','Cumulative Paid Claims','Paid Development Ultimate Claims','Paid Development Age-to-Ultimate Factor'])],
  17:[(6,9,range(7,10),[2,3,4,5],'',['Accident Year','Reported Claims ($000)','Paid Claims ($000)','Reported Claim Count Development Factor to Ultimate']),
      (11,14,range(12,15),[2,3,4],'',['Age (months)','Case Outstanding to Previous Case Outstanding','Incremental Paid Claims to Previous Case Outstanding'])],
  20:[(6,11,range(8,12),[2,3,4,5,6,7],'Cumulative Closed Claim Counts as of (months)',['Accident Year','12','24','36','48','Estimated Ultimate Claim Count']),
      (13,15,[15],[2,3,4,5,6],'',['Selected Disposal Rate','12','24','36','48']),
      (17,22,range(19,23),[2,3,4,5,6],'Cumulative Paid Claims ($) as of (months)',['Accident Year','12','24','36','48']),
      (24,28,range(26,29),[2,3,4,5,6],'Parameter (a; b) for Two-Point Exponential Fit as of (months)',['Accident Year','12','24','36','48'])],
  21:[(6,10,range(8,11),[2,3,4,5],'Cumulative Paid Claims ($) as of (months)',['Accident Year','12','24','36']),
      (12,16,range(14,17),[2,3,4,5],'Cumulative Reported Claims ($) as of (months)',['Accident Year','12','24','36']),
      (18,22,range(20,23),[2,3,4,5],'Open Claim Counts as of (months)',['Accident Year','12','24','36'])],
  22:[(6,12,range(8,13),[2,3,4,5,6,7],'Cumulative Paid Claims ($) as of (months)',['Accident Year','12','24','36','48','60']),
      (6,12,range(8,13),[9,10,11,12,13,14],'Cumulative Paid ALAE ($) as of (months)',['Accident Year','12','24','36','48','60']),
      (15,20,range(17,21),[2,3,4,5,6],'Paid Claims Development Factors',['Accident Year','12–24','24–36','36–48','48–60']),
      (15,20,range(17,21),[9,10,11,12,13],'Paid ALAE Development Factors',['Accident Year','12–24','24–36','36–48','48–60']),
      (23,28,range(24,29),[2,3,4,5,6],'',['Calendar Year','Earned Premium ($)','Paid Claims ($)','Paid ALAE ($)','Paid ULAE ($)'])],
  23:[(6,11,range(8,12),[2,3,4,5,6],'Cumulative Paid Claims ($000) as of (months)',['Accident Year','6','18','30','42']),
      (13,18,range(15,19),[2,3,4,5,6],'Cumulative Paid ALAE ($000) as of (months)',['Accident Year','6','18','30','42'])],
  24:[(6,11,range(8,12),[2,3,4,5,6],'Cumulative Reported Claims ($) as of (months)',['Accident Year','12','24','36','48']),
      (14,18,range(15,19),[2,3,4],'',['Accident Year','Selected Ultimate Claims ($)','Reported Claims ($) as of December 31, 2015'])],
  25:[(6,10,range(8,11),[2,3,4,5,7,8],'Ultimate Claims Estimates ($000)',['Accident Year','Paid Development','Reported Development','Case Outstanding Development','Paid Bornhuetter–Ferguson','Reported Bornhuetter–Ferguson'])],
 },
 'fall-2015': {
  1:[(6,10,range(7,11),[2,3],'',['Effective Date','Number of Autos Written on Effective Date'])],
  4:[(12,18,range(13,19),[2,3,4],'Annual Trend Fits',['Number of Points','Frequency Exponential Fit','Severity Exponential Fit'])],
  5:[(6,8,range(7,9),[2,3],'',['Effective Date','Direct Impact of Benefit Change'])],
  6:[(6,13,range(7,14),[2,3],'Calendar Year 2014',['Item','Value'])],
  7:[(6,14,range(8,15),[2,3,4,5,6],'Non-Catastrophe Data',['Calendar/Accident Year','Earned Exposures','Amount of Insurance Years ($000)','Indicated Ultimate Frequency Trended to 2014','Indicated Ultimate Loss and ALAE Severity ($) Trended to 2014'])],
  9:[(6,10,range(7,11),[2,3],'Relativities',['Rating Class','Relativity']),
     (6,9,range(8,10),[5,6,7],'2014 Earned Exposures',['Smoke Detector','Territory A','Territory B']),
     (13,16,range(15,17),[2,3,4],'Accident Year 2014 Incurred Loss and ALAE',['Smoke Detector','Territory A','Territory B'])],
  11:[(15,17,range(16,18),[2,3,4],'Amount of Insurance',['Class','Current Relativity','Indicated Relativity']),
      (19,21,range(20,22),[2,3,4],'Territory',['Class','Current Relativity','Indicated Relativity']),
      (23,26,range(25,27),[2,3,4],'In-Force Exposure Distribution',['Amount of Insurance','Territory 1','Territory 2'])],
  13:[(6,11,range(6,12),[2,3],'',['Parameter','Value']),
      (13,15,range(14,16),[2,3],'',['Evaluation at Age','Limited Reported Losses ($)'])],
  14:[(6,9,range(7,10),[2,3,4],'',['Underwriting Year','State A Earned Premium ($000)','State B Earned Premium ($000)']),
      (11,14,range(13,15),[2,3,4,5,6],'Reported Cumulative Development Factors as of (months)',['State','12','24','36','48'])],
  16:[(5,10,range(6,11),[2,3],'Claim #1',['Date','Event']),
      (12,16,range(13,17),[2,3],'Claim #2',['Date','Event']),
      (18,21,range(19,22),[2,3],'Claim #3',['Date','Event']),
      (23,25,range(24,26),[2,3],'Claim #4',['Date','Event'])],
  17:[(6,11,range(8,12),[2,3,4,5,6],'Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (13,17,range(14,18),[2,3],'',['Calendar Year','On-Level Earned Premium ($000)'])],
  18:[(6,10,range(7,11),[2,3],'',['Calendar/Accident Year','Earned Premium ($000)']),
      (12,17,range(14,18),[2,3,4,5,6],'Reported Claims ($000) as of (months)',['Calendar/Accident Year','12','24','36','48'])],
  19:[(7,13,range(9,14),[2,3,4,5,6,7],'Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48','60']),
      (15,23,list(range(17,21))+[22,23],[2,3,4,5,6,7],'Reported Claim Age-to-Age Factors',['Accident Year','12–24','24–36','36–48','48–60','60–Ultimate']),
      (26,32,range(28,33),[2,3,4,5,6,7],'Paid Claims ($000) as of (months)',['Accident Year','12','24','36','48','60']),
      (34,42,list(range(36,40))+[41,42],[2,3,4,5,6,7],'Paid Claim Age-to-Age Factors',['Accident Year','12–24','24–36','36–48','48–60','60–Ultimate']),
      (45,51,range(47,52),[2,3,4,5,6,7],'Ratio of Paid Claims to Reported Claims as of (months)',['Accident Year','12','24','36','48','60']),
      (53,60,range(56,61),[2,3,4,5,6,7],'Projected Ultimate Claims ($000)',['Accident Year','Earned Premium','Reported Development','Paid Development','Reported Bornhuetter–Ferguson','Paid Bornhuetter–Ferguson']),
      (78,83,range(79,84),[3,4],'Selected Ultimate Claims',['Accident Year','Selected Ultimate Claims ($000)'])],
  21:[(6,10,range(8,11),[2,3,4,5],'Cumulative Reported Claim Counts',['Accident Year','12','24','36']),
      (12,16,range(14,17),[2,3,4,5],'Cumulative Closed Claim Counts',['Accident Year','12','24','36']),
      (28,32,range(30,33),[2,3,4,5],'Unadjusted Paid Claims Severity on Closed Claims',['Accident Year','12','24','36'])],
  23:[(6,11,range(9,12),[2,3,4,5,6,7],'Net of Excess of Loss Reinsurance — Claims as of December 31, 2014',['Policy Year','Gross Ultimate Claims Estimate ($000)','Net Ultimate Claims Estimate ($000)','Reported ($000)','Paid ($000)','Stop Loss Limit ($000)'])],
  24:[(6,11,range(8,12),[2,3,4,5,6],'Paid Claims Only ($000) as of (months)',['Accident Year','12','24','36','48']),
      (13,18,range(15,19),[2,3,4,5,6],'Paid ALAE ($000) as of (months)',['Accident Year','12','24','36','48']),
      (21,25,range(22,26),[2,3],'',['Accident Year','Selected Ultimate Claims Only ($000)'])],
  25:[(6,10,range(9,11),[2,3,4,5,6,7],'Claims as of December 31, 2014 and Development Technique Ultimate Claims',['Accident Year','Reported Claims ($000)','Paid Claims ($000)','Reported Development Ultimate ($000)','Paid Development Ultimate ($000)','Selected Ultimate Claims ($000)'])],
 }
}

# The workbook is a transcription aid. These cells differ from the printed question pages.
FIXES = {
 ('spring-2016',2,'B3'):'policies',
 ('spring-2016',3,'B10'):'The company is also considering keeping its current exposure base as car-years but including hours driven in its risk classification system. Briefly discuss the appropriateness of adding this risk characteristic to the company’s risk classification system using three considerations from the Actuarial Standard of Practice No. 12: Risk Classification (for All Practice Areas).',
 ('spring-2016',4,'F9'):'March 1, 2014',
 ('spring-2016',7,'C10'):'1.5%',
 ('spring-2016',9,'B12'):'• Risk margin = 10% of excess losses.',
 ('fall-2015',7,'B27'):'• Target underwriting profit provision = 6%.',
 ('fall-2015',16,'B21'):'January 30, 2014',
}
for row,year in zip(range(8,12),range(2011,2015)):
    FIXES['spring-2016',24,f'B{row}']=str(year)
for row,year in zip(range(15,19),range(2011,2015)):
    FIXES['spring-2016',24,f'B{row}']=str(year)
FIXES['spring-2016',24,'C16']='1,000'

def val(exam,n,cell):
    if (exam,n,cell.coordinate) in FIXES:return FIXES[exam,n,cell.coordinate]
    if cell.value is None:return ''
    if isinstance(cell.value,datetime):return f'{cell.value:%B} {cell.value.day}, {cell.value.year}'
    if isinstance(cell.value,int) and 2000<=cell.value<=2030:return str(cell.value)
    result=displayed(cell)
    if isinstance(cell.value,(int,float)) and cell.value>=1000 and cell.value==int(cell.value) and '%' not in cell.number_format:
        result=f'{cell.value:,.0f}'
    return re.sub(r'\s+',' ',result).strip()

def row_text(exam,n,sh,row):
    return ' '.join(val(exam,n,c) for c in sh[row][1:17] if c.value is not None and c.data_type!='f').strip()

def blocks(exam,n,sh,first_part):
    specs=T[exam].get(n,[])
    covered={row for first,last,*_ in specs for row in range(first,last+1)}
    result=[]
    for row in range(3,sh.max_row+1):
        for first,last,rows,cols,title,headers in specs:
            if row==first:
                result.append({'type':'table','title':title,'headers':headers,
                    'rows':[[val(exam,n,sh.cell(r,c)) for c in cols] for r in rows]})
        if row in covered or row>=first_part:continue
        line=row_text(exam,n,sh,row)
        if line:result.append({'type':'line','text':line})
    return result,covered

def clean(s):
    s=re.sub(r'(?m)^\s*EXAM 5 (?:SPRING 2016 |FALL 2015 )?SAMPLE ANSWERS AND EXAMINER.S REPORT\s*$','',s)
    s=s.replace('','•').replace('','→').replace('','→').replace('','←')
    return re.sub(r'\n{3,}','\n\n',s).strip()

def report_parts(s):
    markers=list(re.finditer(r'(?im)^\s*Part\s+([a-z])\s*:?(?:\s*\.?\d+(?:\.\d+)?\s+point(?:\(s\)|s)?)?\s*$',s))
    if not markers:return {'a':s}
    return {m.group(1).lower():s[m.end():markers[i+1].start() if i+1<len(markers) else None] for i,m in enumerate(markers)}

def report(exam,n):
    raw=(ROOT/'tmp'/exam.replace('-','')/f'report-q{n:02}.txt').read_text()
    sample,insight=re.split(r'(?im)^\s*EXAMINER.S REPORT\s*$',raw,maxsplit=1)
    sample=re.split(r'(?im)^\s*SAMPLE ANSWERS?\s*$',sample,maxsplit=1)[-1]
    answers=report_parts(sample)
    notes=report_parts(insight)
    for key,s in list(answers.items()):
        s=clean(s)
        samples=re.split(r'(?im)^\s*Sample (?:Answer |Response )?\d+\s*:?(?:\s*\([^)]*\))?\s*$',s)
        answers[key]=next((clean(item) for item in samples if len(clean(item))>=30), '')
    for key,s in list(notes.items()):
        s=clean(s)
        # Use the examiner's expectation, omitting generic performance notes and error lists.
        m=re.search(r'(?i)Candidates (?:were expected to|should have|needed to)[^.]*\.',s)
        if m:s=m.group(0)
        else:
            s=re.split(r'(?im)^\s*Common (?:errors|mistakes|responses)\b',s,maxsplit=1)[0]
            s=re.sub(r'\s+',' ',s).strip()
            s=re.split(r'(?<=[.!?])\s+',s)[0]
        s=re.sub(r'\s+',' ',s).strip()
        s=re.sub(r'(?i)^Candidates (?:were expected to|should have|needed to)\s+','',s)
        if len(s)>280:
            s=s[:280].rsplit(' ',1)[0].rstrip(' ,;:')+'.'
        if s and not s.endswith(('.','!','?')):s+='.'
        notes[key]=s[0].upper()+s[1:] if s else ''
    return answers,notes

def question_page(exam,n):
    if exam=='spring-2016':return n+3 if n<=22 else n+4
    if n<=10:return n+4
    if n<=19:return n+5
    return n+6

SOLUTION_FIXES={
 ('spring-2016',13,'a'):'Trend each loss type from July 1, 2015 to January 1, 2018 (2.5 years), then develop separately. Indemnity = observed losses × 2.50 × 1.03^2.5; medical = observed losses × 1.50 × 1.06^2.5. The total trended ultimate losses are about $8,853,931 for A, $8,942,832 for B, and $12,502,185 for C. Divide by exposures to get pure premiums of $3,541.57, $2,555.09, and $2,778.26. Relative to base territory C, the indicated relativities are 1.275, 0.919, and 1.000.',
 ('spring-2016',13,'b'):'The current exposure-weighted average relativity is about 1.0143; the indicated average is about 1.0385. Use an off-balance factor of 1.0143/1.0385 ≈ 0.977 so total premium is unchanged. The territorial changes are (1.275/1.20) × 0.977 − 1 ≈ +3.7% for A, (0.919/0.90) × 0.977 − 1 ≈ −0.2% for B, and (1.000/1.000) × 0.977 − 1 ≈ −2.3% for C.',
 ('spring-2016',14,'b'):'Select industry paid age-to-age factors 2.01, 1.255, and 1.10 using the averages at each maturity. Accident year 2015 ultimate claims = $1,000,000 × 2.01 × 1.255 × 1.10 = approximately $2,774,805.',
 ('spring-2016',14,'a'):'One accepted selection uses the company paid age-to-age factors 1.939, 1.667, and 1.256, based on all-year weighted averages and the available 36-to-48 observation. Accident year 2015 ultimate claims = $1,000,000 × 1.939 × 1.667 × 1.256 ≈ $4,060,000. Other supported factor selections received credit because the company pattern is volatile.',
 ('spring-2016',14,'c'):'Use about $2.775 million from the industry factors. The company age-to-age factors fluctuate substantially, and its 12-month estimate is sensitive to a highly leveraged development factor. An estimate that blends company and industry results with a sound justification is also reasonable.',
 ('fall-2015',5,'a'):'For fourth accident quarter 2014, the fractions of losses at benefit levels 1.000, 1.065, and 1.065 × 1.043 are 0.25, 0.625, and 0.125. The direct benefit adjustment is (1.065 × 1.043)/(1.000 × 0.25 + 1.065 × 0.625 + 1.065 × 1.043 × 0.125) = 1.0534.',
 ('fall-2015',5,'b'):'For first policy quarter 2014, the fractions of losses at benefit levels 1.000, 1.065, and 1.065 × 1.043 are 0.03125, 0.59375, and 0.375. The direct benefit adjustment is (1.065 × 1.043)/(1.000 × 0.03125 + 1.065 × 0.59375 + 1.065 × 1.043 × 0.375) = 1.0284.',
 ('fall-2015',23,'a'):'Apply the stop loss limit of $1.5 million to each policy year’s net ultimate and reported claims. Policy-year IBNR is $25,000 for 2012, $0 for 2013, and $100,000 for 2014, totaling $125,000 net of all reinsurance.',
 ('fall-2015',23,'b'):'Apply the stop loss limit of $1.5 million to each policy year’s net ultimate and paid claims. Unpaid claims are $275,000 for 2012, $300,000 for 2013, and $500,000 for 2014, totaling $1,075,000 net of all reinsurance.',
}
INSIGHT_FIXES={
 ('spring-2016',5,'a'):'Apply the coverage triggers to identify which claims belong to each claims-made and occurrence policy.',
 ('spring-2016',11,'b'):'Explain why trend affects excess layers more strongly than basic limits and raises the increased limits factor.',
 ('spring-2016',11,'c'):'Select and support a credibility complement for the $250,000 excess of $250,000 layer.',
 ('spring-2016',20,'a'):'Adjust closed counts using the selected disposal rates, build the adjusted paid claims triangle, select development factors, and calculate the revised ultimate.',
 ('fall-2015',4,'a'):'Recognize the frequency trend change, choose and justify frequency and severity trends, determine the correct trend period, and calculate the pure premium trend factor.',
 ('fall-2015',5,'a'):'Weight the three benefit levels by their share of fourth accident quarter 2014 losses, then on-level to the current combined benefit level.',
 ('fall-2015',5,'b'):'Weight the three benefit levels for first policy quarter 2014, then on-level to the current combined benefit level.',
 ('fall-2015',5,'c'):'Explain that unadjusted pure premium trend includes the one-time benefit changes, and restate losses to a common benefit level before selecting trend.',
 ('fall-2015',6,'a'):'Calculate the underwriting expense ratio from expenses and premium, then use the combined ratio and LAE-to-loss ratio to find the operating expense ratio.',
 ('fall-2015',10,'a'):'Use the charts to assess stability over time, credibility, GLM uncertainty, and agreement with the loss ratio indications.',
 ('fall-2015',12,'b'):'Combine the uniform and point-mass portions of limited average severity, then apply frequency to find the pure premium.',
 ('fall-2015',16,'a'):'Place all four claims in the correct report-year reported and accident-year paid triangles, net of recoveries.',
 ('fall-2015',18,'a'):'Calculate and justify selected development factors, combine them into a cumulative factor, and apply the Bornhuetter–Ferguson formula.',
 ('fall-2015',18,'b'):'A cumulative factor below one implies more than 100% reported under Bornhuetter–Ferguson. Explain whether to cap the factor or use another method.',
 ('fall-2015',25,'a'):'Treat accident year 2013 as 24 months mature at December 31, 2014, and derive the paid and reported 24-to-ultimate factors.',
 ('fall-2015',25,'b'):'Use selected ultimate claims with the development factors from part a to derive expected paid and reported claims at 12 and 24 months.',
}

def main(exam):
    cfg=CONFIG[exam]
    book=load_workbook(ROOT/'past_exams_excel'/cfg['workbook'],data_only=False)
    grid=book['Point Grid']
    pdf=PdfReader(ROOT/'past_exams'/cfg['pdf'])
    questions=[]
    for n in range(1,cfg['count']+1):
        sh=book[str(n)]
        markers=[(c.row,c.value.strip()[1].lower()) for row in sh for c in row
            if c.column==1 and isinstance(c.value,str) and re.fullmatch(r'\([a-z]\)',c.value.strip(),re.I)]
        assert markers,(exam,n,'no parts')
        source,covered=blocks(exam,n,sh,markers[0][0])
        answers,notes=report(exam,n)
        parts=[]
        for i,(start,letter) in enumerate(markers):
            end=markers[i+1][0] if i+1<len(markers) else sh.max_row+1
            points=grid.cell(n+7,ord(letter)-ord('a')+3).value
            assert isinstance(points,(int,float)),(exam,n,letter,'points',points)
            prompt=' '.join(row_text(exam,n,sh,row) for row in range(start,end) if row not in covered and row_text(exam,n,sh,row))
            answer=SOLUTION_FIXES.get((exam,n,letter),answers.get(letter,''))
            insight=INSIGHT_FIXES.get((exam,n,letter),notes.get(letter,''))
            if not answer and len(markers)==1:answer=answers.get('a','')
            if not insight and len(markers)==1:insight=notes.get('a','')
            assert prompt and answer and insight,(exam,n,letter,'missing prompt/answer/insight',bool(prompt),bool(answer),bool(insight))
            parts.append({'id':letter,'points':points,'prompt':prompt,'solution':answer,'insight':insight})
        first=cfg['report_start'][n-1]
        last=cfg['report_start'][n] if n<cfg['count'] else len(pdf.pages)+1
        solution_pages=list(range(first,max(first+1,last)))
        q={'id':f'{exam}-{n}','number':n,'exam':exam.replace('-',' ').title(),
           'chapterIds':CHAPTERS[exam][n-1], 'points':sum(p['points'] for p in parts),
           'questionPage':question_page(exam,n),'solutionPages':solution_pages,
           'sourceBlocks':source,'parts':parts}
        if exam=='fall-2015' and n==10:
            q['figure']={'src':'assets/exam-graphs/fall-2015-q10.png','title':'Model diagnostic graphs',
                'alt':'Two model diagnostic charts provided with the question'}
        questions.append(q)
    assert len(questions)==cfg['count']
    assert sum(len(q['parts']) for q in questions)==cfg['part_count']
    assert abs(sum(q['points'] for q in questions)-cfg['points'])<1e-9
    path=ROOT/'question_data'/(exam.replace('-','')+'_questions.js')
    path.write_text(f'// Transcribed from the official {exam.replace("-"," ").title()} exam PDF; point grid checked against the CBT workbook.\nwindow.{cfg["var"]} = '+json.dumps(questions,ensure_ascii=False,indent=2)+';\n')
    print(exam,len(questions),sum(len(q['parts']) for q in questions),sum(q['points'] for q in questions),path)

if __name__=='__main__':
    for exam in CONFIG:main(exam)
