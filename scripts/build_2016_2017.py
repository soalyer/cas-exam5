"""Transcribe the Spring 2017 and Fall 2016 CBT sheets against the official PDFs.

The table specifications below use complete headings from the printed exams.
The workbook supplies cell values, question prompts, and the scored-part grid.
"""
import json
import re
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from build_spring2019 import displayed

ROOT = Path(__file__).resolve().parents[1]
CONFIG = {
 'spring-2017': ('CBT_Exam_5_S.17_v01.xlsx','SPRING_2017_QUESTIONS',26,62,57.5,
   [34,35,38,40,44,46,48,50,53,55,58,60,62,67,69,71,74,77,79,83,84,86,88,91,94,95]),
 'fall-2016': ('CBT_Exam_5_F.16_v01.xlsx','FALL_2016_QUESTIONS',27,64,56,
   [33,36,38,40,42,44,46,47,50,51,52,55,57,60,62,64,67,69,72,74,76,78,82,86,89,92,95]),
}
CHAPTERS = {
 'spring-2017': [
 ['ratemaking-4'],['ratemaking-5'],['ratemaking-6'],['ratemaking-7'],['ratemaking-8'],
 ['ratemaking-9'],['ratemaking-12'],['ratemaking-10'],['ratemaking-12'],['ratemaking-9'],
 ['ratemaking-11'],['ratemaking-15'],['ratemaking-8','reserving-7'],['reserving-6'],['reserving-6','reserving-13'],
 ['reserving-11'],['reserving-11'],['reserving-8','reserving-9'],['reserving-7','reserving-12'],['reserving-7'],
 ['reserving-15'],['reserving-13'],['reserving-7','reserving-10'],['reserving-17'],['reserving-16'],['reserving-15']],
 'fall-2016': [
 ['ratemaking-4'],['ratemaking-5'],['ratemaking-16'],['ratemaking-6'],['ratemaking-6','reserving-14'],
 ['ratemaking-7'],['ratemaking-7'],['ratemaking-8','reserving-7'],['ratemaking-7','ratemaking-8'],['ratemaking-9'],
 ['ratemaking-11'],['ratemaking-10'],['ratemaking-9','ratemaking-12'],['ratemaking-11'],['ratemaking-15'],
 ['reserving-6'],['reserving-1'],['reserving-8','reserving-9'],['reserving-7','reserving-12'],['reserving-11'],
 ['reserving-9','reserving-10'],['reserving-11'],['reserving-9','reserving-13'],['reserving-14'],['reserving-7','reserving-16'],
 ['reserving-17'],['reserving-15']]
}
# Each table: first covered row, final data row, data rows, workbook columns,
# printed table title, and complete headings from the official question page.
T = {
 'spring-2017': {
  2:[(5,8,range(6,9),[2,3],'',['Rate Change Effective Date','Overall Average Rate Change'])],
  3:[(5,10,range(6,11),[2,3,4,5,6],'',['Claim Number','Accident Date','Transaction Date','Incremental Payment ($)','Ending Case Reserves ($)'])],
  4:[(5,12,range(6,13),[2,3],'($000)',['Item','Amount'])],
  6:[(5,7,range(6,8),[2,3,4,5,6],'',['Class','Exposures','Current Rate ($)','True Expected Cost ($)','Proposed Rate ($)'])],
  7:[(5,8,range(6,9),[2,3,4,5],'',['Class','Earned Exposures','Reported Loss and ALAE ($)','Current Relativity'])],
  8:[(25,26,[26],[2,3,4,5,6],'',['Number of Occupants','1–2','3–4','5–8','>8'])],
  9:[(5,17,range(6,18),[2,3,4,5,6],'',['State','Class','Exposures','Losses ($)','Current Pure Premium ($)'])],
  10:[(5,9,range(7,10),[2,3],'Territory Factors',['Territory','Factor']),
      (21,24,range(22,25),[2,3,4],'Calendar-Accident Year 2016',['Territory','Earned Exposures','Ultimate Loss Cost ($)'])],
  11:[(5,11,range(6,12),[2,3,4],'',['Size of Loss ($000)','Loss Distribution','Average Reported Loss ($000)'])],
  13:[(5,7,range(6,8),[2,3],'',['Calendar Year','Earned Premium ($000)']),
      (5,9,range(7,10),[5,6],'Rate Change History',['Effective Date','Average Rate Change']),
      (11,15,range(13,16),[2,3,4,5],'Reported Loss and ALAE ($000) Capped at $100,000 as of (months)',['Accident Year','12','24','36']),
      (17,24,range(20,25),[2,3,4],'Excess Loss and ALAE ($000) History — Trended Reported Loss and ALAE',['Accident Year','Unlimited','Excess of $100,000'])],
  14:[(5,9,range(7,10),[2,3,4,5],'Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36']),
      (11,15,range(13,16),[2,3,4,5],'Incremental Reported Claim Counts as of (months)',['Accident Year','12','24','36']),
      (17,21,range(19,22),[2,3,4,5],'Cumulative Reported Claims ($000) by Report Year as of (months)',['Report Year','12','24','36'])],
  15:[(5,9,range(7,10),[2,3,4,5],'Cumulative Paid Claims ($000) as of (months)',['Accident Year','12','24','36']),
      (11,15,range(13,16),[2,3,4,5],'Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36']),
      (17,21,range(19,22),[2,3,4,5],'Cumulative Reported Claim Counts as of (months)',['Accident Year','12','24','36']),
      (23,27,range(25,28),[2,3,4,5],'Open Claim Counts as of (months)',['Accident Year','12','24','36'])],
  16:[(5,9,range(6,10),[2,3,4,5],'',['Accident Year','Selected Ultimate Claim Counts','Earned Premium ($000)','Premium On-Level Adjustment Factor to 2016'])],
  17:[(5,9,range(7,10),[2,3,4,5],'Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36']),
      (11,15,range(13,16),[2,3,4,5],'Cumulative Reported Claim Counts as of (months)',['Accident Year','12','24','36'])],
  18:[(5,10,range(7,11),[2,3,4,5,6],'Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (5,9,range(7,10),[8,9,10,11],'Reported Claims Age-to-Age Factors',['Accident Year','12–24','24–36','36–48']),
      (13,17,range(14,18),[2,3],'',['Accident Year','Earned Premium'])],
  19:[(5,10,range(7,11),[2,3,4,5,6],'Cumulative Paid Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (12,17,range(14,18),[2,3,4,5,6],'Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (19,24,range(21,25),[2,3,4,5,6],'Case Outstanding ($000) as of (months)',['Accident Year','12','24','36','48'])],
  20:[(5,10,range(7,11),[2,3,4,5,6],'Cumulative Paid Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (12,17,range(14,18),[2,3,4,5,6],'Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48'])],
  21:[(5,10,range(7,11),[2,3,4,5,6],'Ultimate Claim Estimates ($000) as of December 31, 2015 and 2016',['Accident Year','2015 Paid Development','2015 Reported Development','2016 Paid Development','2016 Reported Development'])],
  22:[(5,10,range(7,11),[2,3,4,5,6],'Cumulative Closed Claim Counts as of (months)',['Accident Year','12','24','34','48']),
      (12,17,range(14,18),[2,3,4,5,6],'Cumulative Reported Claim Counts as of (months)',['Accident Year','12','24','36','48'])],
  23:[(5,10,range(7,11),[2,3,4,5,6,7],'Cumulative Paid Claims as of (months)',['Accident Year','Earned Premium','12','24','36','48'])],
  25:[(6,8,[8],[2,4],'Estimated Ultimate ALAE Without Adjustment for Large Claim',['Paid ALAE Development Technique','Paid ALAE to Paid Claims Only Ratio']),
      (10,12,range(11,13),[2,3,4,5,6],'',['Claims','Paid Claims Only','Paid ALAE','Ultimate Claims Only','Ultimate ALAE'])],
 },
 'fall-2016': {
  1:[(5,11,range(6,12),[2,3,4],'',['Effective Date','Expiration Date','Number of Policies'])],
  2:[(5,8,range(6,9),[2,3],'',['Effective Date','Rate Change'])],
  3:[(5,11,range(7,12),[2,3,4,5,6,7],'Reported Loss ($) by Report Year Lag',['Report Year','0','1','2','3','4'])],
  4:[(5,10,range(6,11),[2,3,4],'',['Accident Year','Frequency','Severity ($)']),
      (12,15,range(13,16),[2,3],'',['Accident Year','Ultimate Losses ($000)'])],
  5:[(7,10,range(8,11),[2,3,4,5],'',['Accident Year','Earned Exposures','Direct Ultimate Losses ($000)','Claim Counts']),
      (12,19,range(14,20),[2,3,4],'Ultimate Value of Direct Claims Excess of $500,000',['Accident Year','Claim','Direct Ultimate Loss of Individual Claims ($000)']),
      (21,24,range(22,25),[2,3],'',['Accident Year','Retention ($000)'])],
  6:[(5,11,range(7,12),[2,3,4,5],'Calendar Year ($000)',['Expense','2013','2014','2015']),
      (13,16,range(15,17),[2,3,4,5],'Calendar Year ($000)',['Premium','2013','2014','2015'])],
  11:[(5,9,range(6,10),[2,3,4],'',['Claim Type','Number of Claims','Loss Amount of Each Claim ($)'])],
  8:[(15,21,range(15,22),[2,4],'',['Item','Value']),
     (23,26,range(25,27),[2,3],'Rate Change History',['Effective Date','Change']),
     (28,29,[29],[2,3,4,5,6],'',['Measure','2012','2013','2014','2015']),
     (31,35,range(33,36),[2,3,4,5],'Cumulative Reported Loss ($000) as of (months)',['Accident Year','12','24','36'])],
  13:[(5,9,range(6,10),[2,3,4,5],'',['Territory','Number of Exposures','Trended and Ultimate Incurred Losses and ALAE ($)','Current Territorial Relativity'])],
  16:[(5,11,range(7,12),[2,3,4,5,6,7,8,9],'Claim Transactions by Calendar Year',['Claim ID','Accident Date','2013 Payments','2013 Ending Case Outstanding','2014 Payments','2014 Ending Case Outstanding','2015 Payments','2015 Ending Case Outstanding'])],
  18:[(5,9,range(6,10),[2,3,4,5,6],'Unpaid Claims Estimate',['Accident Year','On-Level Earned Premium','Paid Claims ($000)','Paid Bornhuetter–Ferguson ($000)','Paid Development ($000)'])],
  19:[(5,10,range(7,11),[2,3,4,5,6],'Cumulative Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (5,10,range(7,11),[8,9,10,11,12],'Cumulative Paid Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (12,18,range(14,19),[2,3,4],'Industry Benchmark Claims Development Factors',['Age to Age','Reported','Paid'])],
  21:[(5,9,range(6,10),[2,3,4,5],'',['Accident Year','Reported Claims ($000)','On-Level Earned Premium ($000)','Reported Development Ultimate ($000)'])],
  22:[(6,11,range(8,12),[2,3,4,5,6],'Closed Claim Counts as of (months)',['Accident Year','12','24','36','48']),
      (6,11,range(8,12),[8,9,10,11,12],'Paid Claims ($000) as of (months)',['Accident Year','12','24','36','48']),
      (13,18,range(15,19),[2,3,4,5,6],'Reported Claim Counts as of (months)',['Accident Year','12','24','36','48']),
      (13,18,range(15,19),[8,9,10,11,12],'Reported Claims ($000) as of (months)',['Accident Year','12','24','36','48'])],
  23:[(5,10,range(7,11),[2,3,4,5,6],'Case Outstanding ($) as of (months)',['Accident Year','12','24','36','48']),
      (12,17,range(14,18),[2,3,4,5,6],'Paid Claims ($) as of (months)',['Accident Year','12','24','36','48']),
      (19,24,range(21,25),[2,3,4,5,6],'Open Claim Counts as of (months)',['Accident Year','12','24','36','48']),
      (26,28,range(27,29),[2,3],'',['Calendar Year','Written Premium'])],
  24:[(5,10,range(7,11),[2,3,4,5],'Paid Claims ($) as of (months)',['Accident Year','12','24','36']),
      (12,17,range(14,18),[2,3,4,5],'Received Salvage and Subrogation ($) as of (months)',['Accident Year','12','24','36'])],
  25:[(5,9,range(7,10),[2,3,4,5],'Reported Claims Only ($) as of (months)',['Accident Year','12','24','36']),
      (11,15,range(13,16),[2,3,4,5],'Reported ALAE ($) as of (months)',['Accident Year','12','24','36'])],
  26:[(5,9,range(6,10),[2,3,4],'',['Accident Year','Exposures','Ultimate Claims ($)']),
      (11,15,range(12,16),[2,3,4,5],'',['Calendar Year','Reported Claims ($)','Paid Claims ($)','Paid ULAE ($)'])],
  27:[(5,8,range(7,9),[2,3,4,6],'Claim Count Estimates as of December 31, 2014 and 2015',['Accident Year','2014 Selected Ultimate Claim Counts','2014 Reported Claim Counts','2015 Reported Claim Counts']),
      (10,13,range(11,14),[2,3],'',['Maturity (months)','Cumulative Percent of Claim Counts Reported'])],
 }
}
FIXES = {
 ('spring-2017',2,'C5'):'Overall Average Rate Change',
 ('spring-2017',22,'E6'):'34',
 ('fall-2016',3,'B28'):'Describe how a switch from occurrence to claims-made coverage could affect the target underwriting profit provision.',
 ('fall-2016',16,'D8'):'20',
 ('fall-2016',14,'B16'):'Calculate the coinsurance apportionment ratio, assuming the property is valued at $425,000 instead of $500,000.',
 ('fall-2016',17,'B5'):'• Actuary A relies only on internal data, aggregated across all lines of business.',
 ('fall-2016',18,'B21'):'After constructing these estimates, the actuary learns of a change in the claims department in 2014 that has led to slower claims payments. Discuss whether the unpaid claims estimate from each technique below would be overstated or understated when calculated without making any adjustments to recognize the slower claims payments:',
 ('fall-2016',6,'B3'):'The following information is available for a single-state, mono-line insurer:',
 ('fall-2016',6,'B9'):'Commissions/Brokerage',
 ('spring-2017',15,'B34'):'ii. Case reserve adequacy',
}

def val(exam,n,c):
    if (exam,n,c.coordinate) in FIXES:return FIXES[exam,n,c.coordinate]
    if c.value is None:return ''
    if isinstance(c.value,datetime):return f'{c.value:%B} {c.value.day}, {c.value.year}'
    if isinstance(c.value,int) and 2000<=c.value<=2030:return str(c.value)
    result=displayed(c)
    if isinstance(c.value,(int,float)) and c.value>=1000 and c.value==int(c.value) and '%' not in c.number_format:
        result=f'{c.value:,.0f}'
    return re.sub(r'\s+',' ',result).strip()

def row_text(exam,n,sh,row):
    return ' '.join(val(exam,n,c) for c in sh[row][1:17] if c.value is not None and c.data_type!='f').strip()

def blocks(exam,n,sh,stop):
    specs=T[exam].get(n,[])
    covered={r for first,last,*_ in specs for r in range(first,last+1)}
    out=[];last_line=-10
    for row in range(3,stop):
        for first,last,rows,cols,title,headers in specs:
            if row==first:
                if exam=='spring-2017' and n==8:continue # Figure contains the printed table.
                out.append({'type':'table','title':title,'headers':headers,
                    'rows':[[val(exam,n,sh.cell(r,c)) for c in cols] for r in rows]})
        if row in covered:continue
        if exam=='spring-2017' and n==8 and row==47:continue # Included in figure.
        s=row_text(exam,n,sh,row)
        if not s:continue
        if row==last_line+1 and out and out[-1]['type']=='line' and not out[-1]['text'].endswith(('.',':','?','!')) and not s.startswith(('•','i.','ii.','iii.')):
            out[-1]['text']+=' '+s
        else:out.append({'type':'line','text':s})
        last_line=row
    return out

def report_part_text(text):
    markers=list(re.finditer(r'(?m)^Part\s+([a-z])\s*:?\s*(?:(?:\d+(?:\.\d+)?|\.\d+)\s+points?)?\s*$',text,re.I))
    if not markers:return {'a':text}
    return {m.group(1).lower():text[m.end():markers[i+1].start() if i+1<len(markers) else None] for i,m in enumerate(markers)}

def clean(s):
    s=re.sub(r'(?m)^\s*EXAM 5 (?:SPRING 2017|FALL 2016) SAMPLE ANSWERS AND EXAMINER.S REPORT\s*$','',s)
    s=re.sub(r'\n{3,}','\n\n',s)
    return s.strip()

def report(exam,n):
    s=(ROOT/'tmp'/exam.replace('-','')/f'report-q{n:02}.txt').read_text()
    sample,insight=re.split(r'(?m)^EXAMINER.S REPORT\s*$',s,maxsplit=1)
    sample=re.split(r'(?m)^SAMPLE ANSWERS?\s*$',sample,maxsplit=1)[-1]
    answers=report_part_text(sample)
    notes=report_part_text(insight)
    for key,s in list(answers.items()):
        s=clean(s)
        s=re.sub(r'(?m)^Sample (?:Answer )?1:?\s*$','',s)
        s=re.split(r'(?m)^Sample (?:Answer )?2:?\s*$',s,maxsplit=1)[0]
        s=s.replace('','→').replace('','→').replace('','←')
        answers[key]=clean(s)
    for key,s in list(notes.items()):
        s=clean(s)
        s=re.split(r'(?im)^\s*Common (?:errors|mistakes|responses)\b',s,maxsplit=1)[0]
        s=re.sub(r'(?i)\bCandidates were expected to\b','',s)
        s=re.sub(r'(?i)\bCandidates should have\b','',s)
        s=re.sub(r'\s+',' ',s).strip(' .:;')
        sentences=re.split(r'(?<=[.!?])\s+',s)
        s=sentences[0].strip()
        if len(s)>230:
            prefix=s[:230]
            cut=max(prefix.rfind(','),prefix.rfind(';'))
            s=prefix[:cut] if cut>100 else prefix.rsplit(' ',1)[0]
        s=s.rstrip(' ,;:')
        if not s.endswith(('.','!','?')):s+='.'
        notes[key]=s[0].upper()+s[1:] if s else ''
    return answers,notes

SOLUTION_FIXES={
 ('spring-2017',4,'a'):'Commissions and brokerage are 2,250/15,000 = 15% of written premium because they are incurred when policies are written. General expenses are 360/12,000 = 3% of earned premium because they support policies throughout the coverage period.',
 ('spring-2017',4,'b'):'Other acquisition costs are 750/15,000 = 5% of written premium; taxes, licenses, and fees are 300/15,000 = 2%. With commissions at 15%, general expenses at 3%, and the underwriting profit provision at −5%, the permissible loss and LAE ratio is 1 − (15% + 3% + 5% + 2%) − (−5%) = 80%. The report prints 3% for taxes in one line, but its 25% total uses the correct 2%.',
 ('spring-2017',4,'d'):'The actual loss and LAE ratio is 10,000/12,000 = 83.33%, above the 80% permissible ratio. The company did not meet its underwriting profit expectation.',
 ('spring-2017',23,'b'):'Adjust pre-decision paid claims for the 20% severity increase and on-level earned premium for the January 2014 rate change. One accepted Cape Cod calculation uses adjusted paid claims of 4,766 and used-up on-level earned premium of 7,666, giving a 62.2% expected loss ratio. Accident-year 2016 ultimate is 619 + 3,215 × 62.2% × (1 − 1/(2.25 × 1.25 × 1.10)) ≈ 1,972.',
 ('fall-2016',23,'a'):'Detrend average case outstanding by 10% per year to a common severity level, multiply by open counts, and add paid claims to build the adjusted reported triangle. Selected 12-to-ultimate factor = 1.985. Calendar-year 2015 earned premium = (34,500 + 37,500)/2 = 36,000. Bornhuetter–Ferguson IBNR = 0.65 × 36,000 × (1 − 1/1.985) ≈ 11,612. Add the 2015 case outstanding of 6,230: unpaid claims ≈ $17,842.',
}
INSIGHT_FIXES={
 ('fall-2016',23,'a'):'Trend average case outstanding, rebuild reported claims, and use the adjusted pattern in Bornhuetter–Ferguson. Derive earned premium and add case outstanding to IBNR.',
 ('fall-2016',18,'c'):'Slower payments leave expected-claims ultimate unchanged and correctly raise its unpaid estimate. Paid Bornhuetter–Ferguson and paid development understate unpaid claims if old payment patterns are used.',
 ('spring-2017',8,'a'):'Support the rating decision with the GLM relativity pattern, confidence intervals, year-to-year stability, and overall chi-square result.',
 ('spring-2017',12,'b'):'Use schedule rating for individual risk features that matter prospectively but are not captured well by past experience.',
 ('spring-2017',14,'a'):'Develop accident-year and report-year losses separately, then use their relationship to isolate incurred but not yet reported claims.',
 ('fall-2016',13,'b'):'Calculate each territorial change, then offset the base rate so the overall premium remains unchanged.',
 ('fall-2016',20,'a'):'Explain which conditions favor a disposal-rate frequency-severity method and why; naming a condition alone is insufficient.',
 ('fall-2016',25,'b'):'Compare ALAE and claims development factors, their ratio, and ALAE size before deciding whether to develop them together.',
 ('fall-2016',27,'a'):'Calculate expected and actual 2015 claim-count emergence for both accident years and compare their levels.',
 ('fall-2016',27,'b'):'Name a limitation of actual-versus-expected emergence and describe an alternative calculation that addresses it.',
}

def question_page(exam,n):
    if exam=='fall-2016':return n+3
    if n<=8:return n+3
    if n<=13:return n+4
    return n+5

def main(exam):
    filename,var,count,parts_expected,points_expected,starts=CONFIG[exam]
    book=load_workbook(ROOT/'past_exams_excel'/filename,data_only=False)
    grid=book['Point Grid']
    questions=[]
    for n in range(1,count+1):
        sh=book[str(n)]
        markers=[(c.row,c.value.strip()[1].lower()) for row in sh for c in row if c.column==1 and isinstance(c.value,str) and re.fullmatch(r'\([a-z]\)',c.value.strip(),re.I)]
        assert markers,(exam,n,'no parts')
        answers,notes=report(exam,n)
        parts=[]
        for i,(start,letter) in enumerate(markers):
            end=markers[i+1][0] if i+1<len(markers) else sh.max_row+1
            points=grid.cell(n+7,ord(letter)-ord('a')+3).value
            assert isinstance(points,(int,float)),(exam,n,letter,'points',points)
            prompt=' '.join(row_text(exam,n,sh,row) for row in range(start,end) if row_text(exam,n,sh,row))
            if exam=='fall-2016' and n in (10,17):
                prompt=row_text(exam,n,sh,6 if n==10 else 8)
            assert prompt,(exam,n,letter,'empty prompt')
            answer=SOLUTION_FIXES.get((exam,n,letter),answers.get(letter,''))
            insight=INSIGHT_FIXES.get((exam,n,letter),notes.get(letter,''))
            if not answer and len(markers)==1:answer=answers.get('a','')
            if not insight and len(markers)==1:insight=notes.get('a','')
            assert answer and insight,(exam,n,letter,'missing report answer or insight')
            parts.append({'id':letter,'points':points,'prompt':prompt,'solution':answer,'insight':insight})
        start_page=starts[n-1]
        end_page=starts[n] if n<count else 97
        q={'id':f'{exam}-{n}','number':n,'exam':exam.replace('-', ' ').title(),'chapterIds':CHAPTERS[exam][n-1],
           'points':sum(p['points'] for p in parts),'questionPage':question_page(exam,n),
           'solutionPages':list(range(start_page,end_page)),'sourceBlocks':blocks(exam,n,sh,markers[0][0]),'parts':parts}
        if exam=='spring-2017' and n==22:q['notice']='The printed question labels the third closed-count column “34” months; the examiner report treats that column as 36 months.'
        if exam=='spring-2017' and n in (8,26):q['figure']={'src':f'assets/exam-graphs/spring-2017-q{n}.png','title':('GLM diagnostic graphs' if n==8 else 'Ultimate claim ratio by evaluation year'),'alt':('GLM output for number of occupants' if n==8 else 'Four ultimate claim ratio estimates by evaluation year')}
        if exam=='fall-2016' and n==12:q['figure']={'src':'assets/exam-graphs/fall-2016-q12.png','title':'GLM output','alt':'Generalized linear model diagnostic graph'}
        questions.append(q)
    assert len(questions)==count
    assert sum(len(q['parts']) for q in questions)==parts_expected
    assert abs(sum(q['points'] for q in questions)-points_expected)<1e-9
    out=ROOT/'question_data'/(exam.replace('-','')+'_questions.js')
    out.write_text(f'// Transcribed from the official {exam.replace("-", " ").title()} exam PDF; point grid checked against the CBT workbook.\nwindow.{var} = '+json.dumps(questions,ensure_ascii=False,indent=2)+';\n')
    print(exam,len(questions),sum(len(q['parts']) for q in questions),sum(q['points'] for q in questions),out)

if __name__=='__main__':
    for exam in CONFIG:main(exam)
